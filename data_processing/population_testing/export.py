from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import (
    Font,
    PatternFill,
)
from openpyxl.utils import (
    get_column_letter,
)

from ..shared.files import ensure_dir
from .assignment import ADDITIONAL_ACCOUNTS



SOURCE_OUTPUT_COLUMNS = [
    "Term",
    "Student Last Name",
    "Student First Name",
    "CWID",
    "Primary Student Level Desc",
    "Primary Program",
    "Primary 1st Major Desc",
    "Registered Credits",
    "Student Residency Desc",
]


def build_bucket_summary(
    df: pd.DataFrame,
) -> pd.DataFrame:
    """
    Build one summary row for every population bucket.
    """
    rows: list[dict[str, object]] = []

    for bucket_id, group in df.groupby(
        "BUCKET_ID",
        sort=True,
    ):
        selected_count = int(
            group[
                "SELECTED_FOR_TESTING"
            ].sum()
        )

        test_count = int(
            group[
                "TESTING_ENVIRONMENT"
            ]
            .eq("TEST")
            .sum()
        )

        prod_count = int(
            group[
                "TESTING_ENVIRONMENT"
            ]
            .eq("PROD")
            .sum()
        )

        rows.append(
            {
                "Bucket ID": bucket_id,
                "Worksheet": (
                    group["BUCKET_SHEET"].iloc[0]
                ),
                "Bucket Description": (
                    group["BUCKET_NAME"].iloc[0]
                ),
                "Population Count": len(group),
                "Selected for Testing": (
                    selected_count
                ),
                "Selection Rate": (
                    selected_count / len(group)
                ),
                "Assigned Count": int(
                    group["ASSIGNED_TO"]
                    .ne("")
                    .sum()
                ),
                "TEST Count": test_count,
                "PROD Count": prod_count,
                "Test All": bool(
                    group[
                        "IS_GRAD_ONLINE"
                    ].iloc[0]
                ),
            }
        )

    return pd.DataFrame(rows)


def _bucket_output(
    df: pd.DataFrame,
) -> pd.DataFrame:
    """
    Return only the requested source columns.
    """
    return df.loc[
        :,
        SOURCE_OUTPUT_COLUMNS,
    ].copy()


def _sample_output(
    df: pd.DataFrame,
) -> pd.DataFrame:
    """
    Return requested source columns plus the staff
    assignment and environment.
    """
    output_columns = (
        SOURCE_OUTPUT_COLUMNS
        + [
            "ASSIGNED_TO",
            "TESTING_ENVIRONMENT",
        ]
    )

    output = df.loc[
        :,
        output_columns,
    ].copy()

    return output.rename(
        columns={
            "ASSIGNED_TO": "Assigned Staff",
            "TESTING_ENVIRONMENT": "Environment",
        }
    )


def _staff_output(
    df: pd.DataFrame,
    *,
    staff_name: str,
) -> pd.DataFrame:
    """
    Return the requested source columns and assigned
    environment for an individual staff worksheet.

    Regular staff worksheets also receive a blank
    Comments column.
    """
    output_columns = (
        SOURCE_OUTPUT_COLUMNS
        + [
            "TESTING_ENVIRONMENT",
        ]
    )

    output = df.loc[
        :,
        output_columns,
    ].copy()

    output = output.rename(
        columns={
            "TESTING_ENVIRONMENT": "Environment",
        }
    )

    if staff_name != ADDITIONAL_ACCOUNTS:
        output["Comments"] = ""

    return output


def _format_worksheet(ws) -> None:
    """
    Apply basic output formatting.
    """
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    for column_cells in ws.iter_cols():
        max_length = max(
            (
                len(str(cell.value))
                for cell in column_cells
                if cell.value is not None
            ),
            default=0,
        )

        width = min(
            max(
                max_length + 2,
                10,
            ),
            45,
        )

        column_letter = get_column_letter(
            column_cells[0].column
        )

        ws.column_dimensions[
            column_letter
        ].width = width


def export_population_workbook(
    df: pd.DataFrame,
    *,
    validation_df: pd.DataFrame,
    output_file: Path,
    staff_names: tuple[str, ...],
) -> Path:
    """
    Export summary, validation, bucket, sample, and
    staff-assignment worksheets.
    """
    ensure_dir(
        output_file.parent
    )

    summary_df = build_bucket_summary(
        df
    )

    sample_df = df[
        df["SELECTED_FOR_TESTING"]
    ].copy()

    with pd.ExcelWriter(
        output_file,
        engine="openpyxl",
    ) as writer:
        summary_df.to_excel(
            writer,
            sheet_name="Summary",
            index=False,
        )

        validation_df.to_excel(
            writer,
            sheet_name="Validation",
            index=False,
        )

        _sample_output(
            sample_df
        ).to_excel(
            writer,
            sheet_name="Testing Sample",
            index=False,
        )

        for _, summary_row in (
            summary_df.iterrows()
        ):
            bucket_df = df[
                df["BUCKET_ID"]
                == summary_row["Bucket ID"]
            ]

            _bucket_output(
                bucket_df
            ).to_excel(
                writer,
                sheet_name=summary_row[
                    "Worksheet"
                ],
                index=False,
            )

        for number, staff_name in enumerate(
            staff_names,
            start=1,
        ):
            staff_df = sample_df[
                sample_df["ASSIGNED_TO"]
                == staff_name
            ]

            sheet_name = (
                f"A{number:02d} {staff_name}"
            )[:31]

            sheet_name = "".join(
                "-"
                if character in "\\/*?:[]"
                else character
                for character in sheet_name
            )

            _staff_output(
                staff_df,
                staff_name=staff_name,
            ).to_excel(
                writer,
                sheet_name=sheet_name,
                index=False,
            )

        for worksheet in (
            writer.book.worksheets
        ):
            _format_worksheet(
                worksheet
            )

        summary_ws = writer.book[
            "Summary"
        ]

        for cell in summary_ws["F"][1:]:
            cell.number_format = "0.0%"

    return output_file