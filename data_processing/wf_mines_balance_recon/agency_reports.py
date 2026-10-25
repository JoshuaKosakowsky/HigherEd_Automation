"""Agency-specific source parsing; no student matching or balance adjustment."""
from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter

from shared.cancellation import CancellationToken
from .source_workbooks import open_source_workbook


COLLECTION_HEADERS = (
    'Your Accnt #', 'Name', 'Placement', 'Orig. Balance', 'Curr. Balance', 'Collected', 'Remarks',
)
RETURN_HEADERS = (
    'Your Accnt #', 'Name', 'Consumer #', 'Orig. Balance', 'Curr. Balance',
    'Last Pay Date', 'Placement Date', 'Return Date', 'Reason',
)
COLLECTION_OUTPUT_HEADERS = (*COLLECTION_HEADERS, 'Creditor', 'Type', 'Source File')
RETURN_OUTPUT_HEADERS = (*RETURN_HEADERS, 'Creditor', 'Source File')
_FILENAME = re.compile(
    r'Report\s*-\s*(AccountStatusReport_NoBreakout|ReturnedAccounts)\s*-\s*'
    r'([0-9]{5})-([0-9]{4}-[0-9]{2}-[0-9]{2})', re.IGNORECASE,
)


@dataclass(frozen=True)
class AgencyReport:
    report_type: str
    creditor: str
    report_date: date

    @property
    def headers(self) -> tuple[str, ...]:
        return COLLECTION_HEADERS if self.report_type == 'collections' else RETURN_HEADERS

    @property
    def output_tab(self) -> str:
        return 'W&F Collections' if self.report_type == 'collections' else 'W&F Returned'


@dataclass(frozen=True)
class ParsedAgencySheet:
    sheet_name: str
    row_count: int
    notes: str


@dataclass(frozen=True)
class AgencyImport:
    report: AgencyReport
    frame: pd.DataFrame
    sheets: tuple[ParsedAgencySheet, ...]


def identify_report(path: Path) -> AgencyReport | None:
    """Recognize report names, allowing any five-digit creditor, not a whitelist."""
    match = _FILENAME.fullmatch(path.stem)
    if match is None:
        if re.search(r'AccountStatusReport_NoBreakout|ReturnedAccounts', path.stem, re.IGNORECASE):
            raise ValueError(
                f'{path.name}: expected Report - report type - 5-digit creditor-YYYY-MM-DD.'
            )
        return None
    try:
        report_date = date.fromisoformat(match[3])
    except ValueError:
        raise ValueError(f'{path.name}: invalid report date in filename.') from None
    kind = 'collections' if match[1].casefold() == 'accountstatusreport_nobreakout' else 'returned'
    return AgencyReport(kind, match[2], report_date)


def _header(value: object) -> str:
    return ' '.join(str(value or '').split()).casefold()


def _blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _amount(value: object) -> Decimal | None:
    if isinstance(value, (bool, date, datetime)) or _blank(value):
        return None
    text = str(value).strip().replace('$', '').replace(',', '')
    if text.startswith('(') and text.endswith(')'):
        text = '-' + text[1:-1]
    try:
        amount = Decimal(text)
        return amount if amount.is_finite() else None
    except InvalidOperation:
        return None


def _fee_type(sheet) -> str:
    text = ' '.join(str(cell.value or '') for cell in sheet[4])
    labels = set(re.findall(r'\b(?:WITH\s+FEES?|NO\s+FEES?)\b', text, re.IGNORECASE))
    kinds = {'Fees' if label.upper().startswith('WITH') else 'No Fees' for label in labels}
    return kinds.pop() if len(kinds) == 1 else 'unspecified'


def _verified_footer(cells, headers, records, source_rows) -> bool:
    """Recognize unlabeled trailing summaries by checking counts/aggregates.

    Unlabeled totals immediately follow details. A blank account/name alone is
    not evidence: require matching count and >=2 monetary totals, or explicit
    SUM formulas over the full detail range. Uncertain rows are retained.
    """
    if not records:
        return False
    monetary = [index for index, header in enumerate(headers)
                if header in {'Orig. Balance', 'Curr. Balance', 'Collected'}]
    descriptive = [index for index, header in enumerate(headers)
                   if index != 0 and index not in monetary]
    if any(not _blank(cells[index].value) for index in descriptive):
        return False
    count = _amount(cells[0].value)
    count_matches = count == len(records)
    # Literal, numeric account IDs can resemble counts; validate the totals too.
    if not _blank(cells[0].value) and not count_matches:
        return False
    matches = 0
    formula_matches = 0
    for index in monetary:
        cell = cells[index]
        if _blank(cell.value):
            continue
        if cell.data_type == 'f':
            letter = get_column_letter(cell.column)
            pattern = rf'=SUM\(\$?{letter}\$?7:\$?{letter}\$?{source_rows[-1]}\)'
            if source_rows == list(range(7, source_rows[-1] + 1)) and re.fullmatch(
                pattern, str(cell.value).replace(' ', ''), re.IGNORECASE,
            ):
                matches += 1
                formula_matches += 1
                continue
            return False
        values = [_amount(row[index]) for row in records]
        total = _amount(cell.value)
        if total is None or any(value is None for value in values):
            return False
        if total.quantize(Decimal('.01')) != sum(values).quantize(Decimal('.01')):
            return False
        matches += 1
    return matches >= 2 and (count_matches or formula_matches == matches)


def read_agency_report(
    path: Path, report: AgencyReport, *, cancellation: CancellationToken | None = None,
) -> AgencyImport:
    book = open_source_workbook(path)
    cached = None
    rows = []
    sources = []
    try:
        for sheet in book:
            if cancellation:
                cancellation.check()
            locations = {}
            for cell in sheet[6]:
                key = _header(cell.value)
                if key in {_header(header) for header in report.headers}:
                    if key in locations:
                        raise ValueError(f'{path.name} / {sheet.title}: duplicate row-6 report header.')
                    locations[key] = cell.column
            if len(locations) != len(report.headers):
                if locations:
                    raise ValueError(f'{path.name} / {sheet.title}: row 6 does not match the expected report headers.')
                continue
            if any(not _blank(cell.value) and _header(cell.value) not in locations for cell in sheet[6]):
                raise ValueError(f'{path.name} / {sheet.title}: unexpected row-6 report headers; check the format.')
            columns = [locations[_header(header)] for header in report.headers]
            fee_type = _fee_type(sheet) if report.report_type == 'collections' else None
            records = []
            source_rows = []
            excluded = 0
            uncertain = []
            for number in range(7, sheet.max_row + 1):
                if cancellation and number % 500 == 0:
                    cancellation.check()
                cells = [sheet.cell(number, column) for column in columns]
                if all(_blank(cell.value) for cell in cells):
                    continue
                # If account data appears after a verified summary below, fail
                # the run rather than silently dropping any intervening rows.
                if _verified_footer(cells, report.headers, records, source_rows):
                    excluded += 1
                    continue
                values = []
                for header, cell in zip(report.headers, cells):
                    value = cell.value
                    if cell.data_type == 'f':
                        if cached is None:
                            cached = open_source_workbook(path, data_only=True)
                        value = cached[sheet.title][cell.coordinate].value
                        if value is None:
                            raise ValueError(
                                f'{path.name} / {sheet.title}: a detail formula has no saved value. '
                                'Recalculate and save in Excel, or use a values-only copy.'
                            )
                    if header in {'Your Accnt #', 'Consumer #'} and isinstance(value, (int, float)) and not isinstance(value, bool):
                        if value == int(value):
                            value = (f'{int(value):0{len(cell.number_format)}d}'
                                     if re.fullmatch(r'0+', cell.number_format) else str(int(value)))
                    if header in {'Orig. Balance', 'Curr. Balance', 'Collected'} and isinstance(value, str):
                        amount = _amount(value)
                        if amount is not None:
                            value = amount
                    values.append(value)
                if _blank(values[report.headers.index('Name')]):
                    uncertain.append(number)
                if excluded:
                    raise ValueError(f'{path.name} / {sheet.title}: account data follows a summary row; check the source layout.')
                records.append(values)
                source_rows.append(number)
            note = f'{excluded} verified total row(s) excluded.'
            if fee_type == 'unspecified':
                note += ' Fee classification unspecified (row 4).'
            if not excluded:
                note += ' Confirm final rows: an unlabeled footer was not identified.'
            if uncertain:
                note += ' Rows without a name retained for review: ' + ', '.join(map(str, uncertain[:20]))
                note += (f' (and {len(uncertain) - 20} more).' if len(uncertain) > 20 else '.')
            sources.append(ParsedAgencySheet(sheet.title, len(records), note))
            for values in records:
                rows.append([*values, report.creditor, *([fee_type] if fee_type else []), path.name])
        if not sources:
            raise ValueError(f'{path.name}: no matching report table found with headers on row 6.')
        output_headers = COLLECTION_OUTPUT_HEADERS if report.report_type == 'collections' else RETURN_OUTPUT_HEADERS
        return AgencyImport(report, pd.DataFrame(rows, columns=output_headers), tuple(sources))
    finally:
        book.close()
        if cached is not None:
            cached.close()
