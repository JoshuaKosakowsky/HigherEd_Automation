"""One source interface for CSV, legacy Excel and modern Excel uploads."""
from __future__ import annotations

import csv
import io
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter


SOURCE_SUFFIXES = frozenset({'.csv', '.xls', '.xlsx', '.xlsm'})


def _csv_workbook(path: Path) -> Workbook:
    content = path.read_bytes()
    encoding = 'utf-16' if content.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig'
    try:
        text = content.decode(encoding)
    except UnicodeDecodeError:
        text = content.decode('cp1252')
    book = Workbook()
    sheet = book.active
    sheet.title = 'CSV'
    for number, row in enumerate(csv.reader(io.StringIO(text, newline='')), 1):
        for column, value in enumerate(row, 1):
            if value:
                cell = sheet.cell(number, column, value)
                # CSV has no cell types; preserve identifiers and literal text.
                cell.data_type = 's'
    return book


def _xls_workbook(path: Path) -> Workbook:
    try:
        import xlrd
    except ImportError:
        raise ValueError('XLS support requires updated dependencies. Run setup.ps1, then restart the app.') from None
    result = Workbook()
    result.remove(result.active)
    source = None
    try:
        source = xlrd.open_workbook(str(path), formatting_info=True, on_demand=True, logfile=io.StringIO())
        for original in source.sheets():
            sheet = result.create_sheet(original.name)
            for row in range(original.nrows):
                for column in range(original.row_len(row)):
                    cell = original.cell(row, column)
                    if cell.ctype == xlrd.XL_CELL_EMPTY:
                        continue
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, source.datemode)
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        value = xlrd.error_text_from_code.get(value, '#VALUE!')
                    elif cell.ctype == xlrd.XL_CELL_BLANK:
                        value = None
                    target = sheet.cell(row + 1, column + 1, value)
                    if cell.ctype == xlrd.XL_CELL_TEXT:
                        target.data_type = 's'
                    xf = source.xf_list[cell.xf_index]
                    number_format = source.format_map.get(xf.format_key)
                    if number_format is not None:
                        target.number_format = number_format.format_str
            for row, info in original.rowinfo_map.items():
                sheet.row_dimensions[row + 1].height = info.height / 20
                sheet.row_dimensions[row + 1].hidden = bool(info.hidden)
            for column, info in original.colinfo_map.items():
                dimension = sheet.column_dimensions[get_column_letter(column + 1)]
                dimension.width = info.width / 256
                dimension.hidden = bool(info.hidden)
            for first_row, last_row, first_col, last_col in original.merged_cells:
                sheet.merge_cells(start_row=first_row + 1, end_row=last_row,
                                  start_column=first_col + 1, end_column=last_col)
        return result
    except (xlrd.XLRDError, ValueError, OverflowError, OSError):
        result.close()
        raise ValueError(f'{path.name}: unable to read legacy XLS data. Check that the workbook is valid and not password-protected.') from None
    finally:
        if source is not None:
            source.release_resources()


def open_source_workbook(path: Path, *, data_only: bool = False, read_only: bool = False) -> Workbook:
    """Read directly; never convert or rewrite the employee's source file.

    CSV cells remain text. XLS supplies saved values, including formula caches,
    plus Excel date/number formats. Modern files keep the existing reader.
    """
    suffix = path.suffix.lower()
    if suffix == '.csv':
        return _csv_workbook(path)
    if suffix == '.xls':
        return _xls_workbook(path)
    if suffix in {'.xlsx', '.xlsm'}:
        return load_workbook(path, data_only=data_only, read_only=read_only, keep_links=False)
    raise ValueError('Source uploads must be CSV, XLS, XLSX, or XLSM files.')
