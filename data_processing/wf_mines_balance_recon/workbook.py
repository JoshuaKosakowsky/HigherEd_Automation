"""Read original reports and collect source snapshots and reconcile account totals."""
from __future__ import annotations

import re
from copy import copy
from datetime import date, datetime
from math import ceil
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, PatternFill

from data_processing.shared.workbook_output import save_new_workbook, validate_output
from data_processing.shared.xlsx_output_format import (
    apply_default_font, format_table_columns, style_header_row,
)
from shared.cancellation import CancellationToken
from shared.insights.report_results import checked_report_result, prepare_excel_dates
from .agency_reports import identify_report, read_agency_report
from .reconciliation import add_principal_amount, reconcile_accounts
from .source_workbooks import open_source_workbook


HOLD_HEADERS = (
    'CWID', 'Last Name', 'First Name', 'CO Hold', 'TO Hold',
    'CO From Date', 'CO To Date', 'Current Account Balance',
    'WOFF', 'BDRC', 'Total Balance',
)
ACTIVITY_HEADERS = (
    "'ID'", "'Name'", "'Detail Code'", "'Description'", "'Amount'", "'Balance'",
    "'Term'", "'Aid Year'", "'Feed Document'", "'Feed Date'",
    'detail_code_description', 'Transaction User',
)
REPORT_HEADERS = {'students': HOLD_HEADERS, 'holds': HOLD_HEADERS, 'activity': ACTIVITY_HEADERS}


def _header_key(value: object) -> str:
    return str(value or '').strip().strip("'\"").strip().casefold()


def prepare_report(frame: pd.DataFrame, report: str) -> pd.DataFrame:
    """Validate source identity without filtering or reclassifying source rows."""
    headers = [_header_key(value) for value in frame.columns]
    required = {_header_key(value) for value in REPORT_HEADERS[report]}
    if not required.issubset(headers) or len(headers) != len(set(headers)):
        raise ValueError(f'The {report} report is missing required headers or has duplicate headers.')
    if 'Report Row Count' in frame.columns:
        frame = checked_report_result(frame, 'Report Row Count')
    canonical = {_header_key(value): value for value in REPORT_HEADERS[report]}
    frame = frame.rename(columns={value: canonical.get(_header_key(value), value) for value in frame.columns})
    if report in {'students', 'holds'}:
        frame = frame.loc[:, list(REPORT_HEADERS[report])]
    return prepare_excel_dates(frame)


def read_insights_file(path: Path, report: str) -> pd.DataFrame:
    """Locate exactly one source table, allowing the usual report title rows."""
    required = {_header_key(value) for value in REPORT_HEADERS[report]}
    candidates = []

    def inspect(rows) -> None:
        iterator = iter(rows)
        for number, row in enumerate(iterator, 1):
            if number > 30:
                break
            if not required.issubset({_header_key(value) for value in row}):
                continue
            headers = [str(value) if value is not None else '' for value in row]
            records = []
            for values in iterator:
                if not any(value is not None and value != '' for value in values):
                    continue
                if len(values) > len(headers):
                    raise ValueError(f'The {report} source has rows wider than its headers.')
                records.append([*values, *([None] * (len(headers) - len(values)))])
            candidates.append(prepare_report(pd.DataFrame(records, columns=headers), report))
            break

    book = open_source_workbook(path, read_only=True)
    try:
        for sheet in book:
            if hasattr(sheet, 'reset_dimensions'):
                sheet.reset_dimensions()
            def rows():
                for row in sheet:
                    if any(cell.data_type == 'f' for cell in row):
                        raise ValueError('Upload original Insights values, not a workbook with formulas.')
                    yield [cell.value for cell in row]
            inspect(rows())
    finally:
        book.close()
    if len(candidates) != 1:
        raise ValueError(f'Upload exactly one matching {report} source table; found {len(candidates)}.')
    return candidates[0]


def _sheet_name(name: str, used: set[str]) -> str:
    base = re.sub(r'[\\/*?:\[\]\x00-\x1f]', '_', name).strip().strip("'") or 'WF Source'
    candidate = base[:31]
    index = 2
    while candidate.casefold() in used:
        suffix = f'_{index}'
        candidate = base[:31 - len(suffix)] + suffix
        index += 1
    used.add(candidate.casefold())
    return candidate


def _report_date(value: object) -> datetime | date | None:
    """Read Excel and CSV dates without imposing pandas' year-2262 limit."""
    if pd.isna(value):
        return None
    if isinstance(value, (datetime, date)):
        return value
    if isinstance(value, str) and value.strip():
        try:
            return datetime.fromisoformat(value.strip())
        except ValueError:
            parsed = pd.to_datetime(value, errors='coerce')
            return None if pd.isna(parsed) else parsed
    return None


def _ordered_report(frame: pd.DataFrame, date_header: str) -> pd.DataFrame:
    """Sort actual dates descending, retaining missing/unreadable dates last.

    Sort before removing times so same-day records retain chronological order.
    Unreadable source text stays intact for review rather than becoming blank.
    """
    column = next(header for header in frame.columns if _header_key(header) == _header_key(date_header))
    dates = [_report_date(value) for value in frame[column]]

    def key(index):
        value = dates[index]
        if value is None:
            return (0, 0, 0, 0, 0)
        time = ((value.hour, value.minute, value.second, value.microsecond)
                if isinstance(value, datetime) else (0, 0, 0, 0))
        return (value.toordinal(), *time)

    result = frame.iloc[sorted(range(len(frame)), key=key, reverse=True)].copy()
    for header in result.columns:
        if _header_key(header).endswith('date') or _header_key(header) == 'placement':
            def date_only(value):
                parsed = _report_date(value)
                if isinstance(parsed, datetime):
                    return parsed.date()
                return parsed if parsed is not None else value
            result[header] = result[header].map(date_only)
    return result


def _write_frame(book: Workbook, name: str, frame: pd.DataFrame, cancellation, *, index: int | None = None) -> None:
    if len(frame) + 1 > 1_048_576 or len(frame.columns) > 16_384:
        raise ValueError('A source exceeds Excel worksheet limits. No workbook was saved.')
    sheet = book.create_sheet(name, index=index)
    sheet.append(list(frame.columns))
    for number, values in enumerate(frame.itertuples(index=False, name=None), 2):
        if cancellation and number % 500 == 0:
            cancellation.check()
        sheet.append([None if pd.isna(value) else value for value in values])
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
            elif isinstance(cell.value, datetime):
                cell.number_format = 'yyyy-mm-dd hh:mm:ss'
            elif isinstance(cell.value, date):
                cell.number_format = 'yyyy-mm-dd'
    apply_default_font(sheet)
    style_header_row(sheet)
    format_table_columns(sheet, headers=list(frame.columns), currency_headers={
        'Current Account Balance', "'Amount'", "'Balance'", 'Orig. Balance', 'Curr. Balance', 'Collected',
        'WOFF', 'BDRC', 'Total Balance', 'Principal Amount',
        'Banner Total Balance', 'W&F Principal Amount', 'Difference',
    })
    if name == 'Reconciliation':
        # A known zero is a match; absent/unknown totals must stay visibly distinct.
        for column in sheet.iter_cols():
            header = column[0].value
            if header not in {'Banner Total Balance', 'W&F Principal Amount', 'Difference'}:
                continue
            for cell in column[1:]:
                cell.number_format = '$#,##0.00;[Red]($#,##0.00);$0.00'
                if header == 'Difference' and cell.value is None:
                    cell.fill = PatternFill('solid', fgColor='FFF2CC')
                    cell.comment = Comment(
                        'Not calculated: Banner Total Balance or W&F Principal Amount '
                        'is missing or unreadable. The CWID may occur in only one source. '
                        'A blank is not a zero or a confirmed match.',
                        'Mines Bursar Automation',
                    )
    for column in sheet.iter_cols(min_row=2):
        if any(isinstance(cell.value, datetime) for cell in column):
            dimension = sheet.column_dimensions[column[0].column_letter]
            dimension.width = max(dimension.width, 24)
    for column in sheet.iter_cols():
        if column[0].value == 'Source File':
            width = 65
            sheet.column_dimensions[column[0].column_letter].width = width
            for cell in column:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                dimension = sheet.row_dimensions[cell.row]
                dimension.height = max(dimension.height or 15, 18 * ceil(len(str(cell.value or '')) / width))
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions


def create_review_workbook(
    students: pd.DataFrame, wf_files: tuple[Path, ...], output: Path,
    *, app_version: str, submission_percentage: object = 0.25,
    cancellation: CancellationToken | None = None,
) -> int:
    """Create account reconciliation, named agency tabs and other source snapshots."""
    validate_output(output, wf_files)
    book = Workbook()
    book.remove(book.active)
    used = {'students in collections', 'reconciliation', 'w&f collections', 'w&f returned'}
    grouped = {}
    agency_sheets = 0
    try:
        students = _ordered_report(prepare_report(students, 'students'), 'CO From Date')
        _write_frame(book, 'Students in Collections', students, cancellation)
        for path in wf_files:
            if cancellation:
                cancellation.check()
            # Unisa verification files are accepted in the batch but omitted in full.
            if any(creditor in path.name for creditor in ('12007', '12009', '12559', '12561')):
                continue
            report = identify_report(path)
            imported = read_agency_report(path, report, cancellation=cancellation) if report else None
            parsed_sheets = set()
            if imported:
                grouped.setdefault(report.output_tab, []).append(imported.frame)
                for detail in imported.sheets:
                    parsed_sheets.add(detail.sheet_name)
                    agency_sheets += 1
            source = open_source_workbook(path)
            cached = None
            try:
                for original in source:
                    if original.title in parsed_sheets:
                        continue
                    name = _sheet_name(f'WF {path.stem} - {original.title}', used)
                    sheet = book.create_sheet(name)
                    if original.max_row > 1_048_576 or original.max_column > 16_384:
                        raise ValueError('A W&F worksheet exceeds Excel limits.')
                    for row in original:
                        if cancellation and row[0].row % 500 == 0:
                            cancellation.check()
                        for cell in row:
                            if cell.value is None and not cell.has_style:
                                continue
                            value = cell.value
                            if cell.data_type == 'f':
                                if cached is None:
                                    cached = open_source_workbook(path, data_only=True)
                                value = cached[original.title][cell.coordinate].value
                                if value is None:
                                    raise ValueError(
                                        f'{path.name} / {original.title} has a formula without a saved value. '
                                        'Recalculate and save it in Excel, or upload a values-only copy.'
                                    )
                            target = sheet.cell(cell.row, cell.column, value)
                            if isinstance(value, str):
                                target.data_type = 's'
                            if cell.has_style:
                                target.font, target.fill = copy(cell.font), copy(cell.fill)
                                target.border, target.alignment = copy(cell.border), copy(cell.alignment)
                                target.protection = copy(cell.protection)
                            target.number_format = cell.number_format
                    for key, dimension in original.column_dimensions.items():
                        sheet.column_dimensions[key] = copy(dimension)
                        sheet.column_dimensions[key].parent = sheet
                    for key, dimension in original.row_dimensions.items():
                        sheet.row_dimensions[key] = copy(dimension)
                        sheet.row_dimensions[key].parent = sheet
                    for merged in original.merged_cells.ranges:
                        sheet.merge_cells(str(merged))
                    sheet.freeze_panes = original.freeze_panes
                    sheet.auto_filter.ref = original.auto_filter.ref
                    agency_sheets += 1
            finally:
                source.close()
                if cached is not None:
                    cached.close()
        collections = add_principal_amount(
            pd.concat(grouped.get('W&F Collections', []), ignore_index=True)
            if 'W&F Collections' in grouped else pd.DataFrame(), submission_percentage,
        )
        position = 1
        for name in ('W&F Collections', 'W&F Returned'):
            if name in grouped:
                frame = collections if name == 'W&F Collections' else pd.concat(grouped[name], ignore_index=True)
                frame = _ordered_report(frame, 'Placement' if name == 'W&F Collections' else 'Return Date')
                _write_frame(book, name, frame, cancellation, index=position)
                position += 1
        returned = pd.concat(grouped['W&F Returned'], ignore_index=True) if 'W&F Returned' in grouped else pd.DataFrame()
        _write_frame(book, 'Reconciliation', reconcile_accounts(students, collections, returned), cancellation, index=0)
        book.properties.title = 'W&F - Banner Recon'
        book.properties.description = f'App {app_version}. Submission Percentage: {submission_percentage}. Account totals for review.'
        save_new_workbook(book, output, cancellation=cancellation)
        return agency_sheets
    finally:
        book.close()
