"""Review workbook using the application's established Excel output helpers."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from data_processing.shared.xlsx_output_format import apply_default_font, write_table_sheet
from .calculation import TuitionReview
from .preparation import money


def _text(value: object) -> str:
    return "" if value is None or pd.isna(value) else str(value)


def _group_manual_rows(transactions: pd.DataFrame) -> list[dict]:
    """Net source charges/payments by term and code, keeping them distinct.

    This is the teaching workbook's ledger view, not a paid-tuition calculation.
    Offsetting zero groups stay in the transaction source sheet.
    """
    groups = {}
    for row in transactions.to_dict("records"):
        key = (_text(row["term_code"]), _text(row["detail_code"]), _text(row["type_ind"]))
        if key not in groups:
            groups[key] = {**row, "amount_usd": Decimal("0.00"), "descriptions": set()}
        groups[key]["amount_usd"] += money(row["amount_usd"])
        if _text(row.get("description")):
            groups[key]["descriptions"].add(_text(row["description"]))
    result = []
    for key in sorted(groups, key=lambda key: (key[0], key[2], key[1])):
        row = groups[key]
        if row["amount_usd"] != 0:
            result.append({**row, "description": "; ".join(sorted(row["descriptions"]))})
    return result


def _write_manual_final(workbook: Workbook, review: TuitionReview, transactions: pd.DataFrame,
                        manifest: dict, identity: pd.Series) -> None:
    """Follow the reference Final layout without storing a student template."""
    sheet = workbook.create_sheet("Final", 0)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A5"
    grouped = _group_manual_rows(transactions)
    excluded = [row for row in grouped if row["type_ind"] == "C" and row["eligibility_status"] == "excluded"]
    sheet.append(["CWID", "NAME", None, None, f"Calendar Year {manifest['tax_year']}"])
    sheet.append([_cell(identity.get("cwid", "")), _cell(" ".join(_text(identity.get(key)) for key in ("first_name", "last_name"))),
                  None, None, None])
    headings, totals = [], []

    def section(title: str) -> None:
        sheet.append([title])
        sheet.merge_cells(start_row=sheet.max_row, start_column=1, end_row=sheet.max_row, end_column=5)
        headings.append(sheet.max_row)

    def headers(values: list[str]) -> None:
        sheet.append(values)
        headings.append(sheet.max_row)

    def ledger(rows: list[dict]) -> tuple[int, int]:
        first = sheet.max_row + 1
        for row in rows:
            sheet.append([_cell(row["detail_code"]), _cell(row["description"]), _cell(row["term_description"]),
                          _cell(row["amount_usd"]) if row["type_ind"] == "C" else None,
                          _cell(row["amount_usd"]) if row["type_ind"] == "P" else None])
        return first, sheet.max_row

    def sum_formula(column: str, first: int, last: int) -> str | int:
        return f"=SUM({column}{first}:{column}{last})" if last >= first else 0

    def total(label: str, charge: object, payment: object = None) -> int:
        sheet.append([label, None, None, charge, payment])
        sheet.merge_cells(start_row=sheet.max_row, start_column=1, end_row=sheet.max_row, end_column=3)
        totals.append(sheet.max_row)
        return sheet.max_row

    section(f"Total Charges and Payments for {manifest['tax_year']} (USD)")
    headers(["Detail Code", "Description", "Term", "Charge", "Payment"])
    first, last = ledger(grouped)
    charges_row = total("Total Charges / Payments", sum_formula("D", first, last), sum_formula("E", first, last))
    sheet.append([])
    sheet.append([])
    section("Non-Qualified Educational Expenses (USD)")
    headers(["Detail Code", "Description", "Term", "Charge", "Payment"])
    first, last = ledger(excluded)
    excluded_row = total("Total Non-Qualified Educational Expenses", sum_formula("D", first, last))
    sheet.append([])
    sheet.append([])
    retained_row = total("Charges Less Excluded Fees (USD)", f"=D{charges_row}-D{excluded_row}")
    # A zero account balance or a charge subtotal must never replace paid allocations.
    paid_row = total("Automated Eligible Paid Amount (USD)", '=IF(ISNUMBER(Summary!$B$6),Summary!$B$6,"")')
    total("Retained Charges Less Eligible Paid (USD)", f'=IF(ISNUMBER(D{paid_row}),D{retained_row}-D{paid_row},"")')
    sheet.append([])
    section("Eligible Tuition and Fees — Annual Average Conversion")
    headers(["Eligible Paid (USD)", "CAD per USD", "Term", "Expenses in CAD", "Payment Effective Dates (UTC)"])
    conversion_first = sheet.max_row + 1
    rate = manifest["exchange_rate"]
    # Reuse the same published rate input as the existing Summary, never a daily rate.
    for index, session in enumerate(review.sessions.to_dict("records"), 2):
        row_number = sheet.max_row + 1
        dates = set()
        date_notes = set()
        applications = review.applications[
            review.applications.charge_term_code.astype(str).eq(str(session["term_code"]))
            & review.applications.retained_paid_usd.ne(0)
        ]
        for value in applications.payment_effective_date:
            try:
                timestamp = pd.to_datetime(value, utc=True)
                if pd.isna(timestamp):
                    date_notes.add("Missing date")
                else:
                    dates.add(timestamp.date())
            except (ValueError, TypeError, OverflowError):
                date_notes.add("Invalid date")
        sheet.append([
            f"=Sessions!K{index}" if review.paid_usd is not None else None,
            '=IF(ISNUMBER(Summary!$B$10),Summary!$B$10,"")', _cell(session["term_description"]),
            f"=ROUND(A{row_number}*B{row_number},2)" if review.paid_cad is not None else None,
            ", ".join([value.strftime("%m/%d/%Y") for value in sorted(dates)] + sorted(date_notes)),
        ])
    conversion_last = sheet.max_row
    cad_row = total("Review Amount on TL11A (CAD)", '=IF(ISNUMBER(Summary!$B$7),Summary!$B$7,"")')
    # Per-term cents can differ from rounding the final USD total once. Expose it.
    total("Final Total Rounding Difference (CAD)",
          f'=IF(ISNUMBER(D{cad_row}),ROUND(D{cad_row}-SUM(D{conversion_first}:D{conversion_last}),2),"")'
          if conversion_last >= conversion_first else None)
    sheet.append([])
    section(f"Rate method: {_text(rate.get('method'))}; rate year: {_text(rate.get('rate_year'))}")
    section("Amounts pending review — see Checks" if review.paid_usd is None or review.paid_cad is None
            else "Administrator must confirm attendance, qualifying degree/courses and fee policy before certification.")
    section("Net-zero groups remain on CWID. Paid amounts are allocation snapshots; rerun after source changes.")
    section("Source totals use candidate terms, not payment dates; funding allocated from other terms remains on Applications.")
    apply_default_font(sheet)
    rule = Side(style="thin", color="808080")
    for row in sheet.iter_rows():
        sheet.row_dimensions[row[0].row].height = 22
        for cell in row:
            cell.font = Font(name="Aptos", size=11)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.column == 3:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if cell.column in (4, 5):
                cell.number_format = '#,##0.00;[Red](#,##0.00);"-"'
    for row_number in headings:
        sheet.row_dimensions[row_number].height = 30
        for cell in sheet[row_number]:
            cell.font = Font(name="Aptos", size=11, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=rule)
    for row_number in [1, 2] + totals:
        sheet.row_dimensions[row_number].height = 30
        for cell in sheet[row_number]:
            cell.font = Font(name="Aptos", size=11, bold=True)
    sheet.cell(cad_row, 4).border = Border(bottom=Side(style="double", color="000000"))
    for row_number in range(conversion_first, conversion_last + 1):
        sheet.cell(row_number, 1).number_format = '#,##0.00;[Red](#,##0.00);"-"'
        sheet.cell(row_number, 2).number_format = "0.0000"
        sheet.cell(row_number, 2).alignment = Alignment(horizontal="center", vertical="center")
        sheet.cell(row_number, 5).number_format = "General"
        sheet.row_dimensions[row_number].height = 42
    for column, width in zip("ABCDE", (17, 42, 21, 22, 33)):
        sheet.column_dimensions[column].width = width
    for row_number in (1, 2):
        sheet.cell(row_number, 1).alignment = Alignment(horizontal="center", vertical="center")
    # Native formulas calculate when opened in Excel; no authoring engine is a runtime dependency.
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = f"A1:E{sheet.max_row}"


def _write_manual_source(workbook: Workbook, transactions: pd.DataFrame) -> None:
    """Use the teaching source columns where supplied; never invent missing Banner fields."""
    headers = ["Detail Code", "Description", "Term", "Charge", "Payment", "Balance",
               "Effective Date", "Transaction Date", "Trans Number", "Feed Date", "Entry Date",
               "Type", "Source Amount (USD)"]
    rows = []
    for row in transactions.to_dict("records"):
        rows.append([
            _cell(row["detail_code"]), _cell(row.get("description")), _cell(row["term_description"]),
            _cell(row["amount_usd"]) if row["type_ind"] == "C" else None,
            _cell(row["amount_usd"]) if row["type_ind"] == "P" else None,
            _cell(money(row["stored_balance_usd"])), _cell(row.get("effective_date"), "effective_date"),
            _cell(row.get("activity_date"), "activity_date"), _cell(row["tran_number"]),
            _cell(row.get("feed_date"), "feed_date"), _cell(row.get("entry_date"), "entry_date"),
            _cell(row["type_ind"]), _cell(row["amount_usd"]),
        ])
    write_table_sheet(workbook, sheet_name="CWID", table_name="TL11AManualSource", headers=headers, rows=rows,
                      currency_headers={"Charge", "Payment", "Balance", "Source Amount (USD)"})
    sheet = workbook["CWID"]
    sheet.row_dimensions[1].height = 30
    for cell in sheet[1]:
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column in (7, 8, 10, 11):
        for cells in sheet.iter_cols(min_col=column, max_col=column, min_row=2):
            for cell in cells:
                cell.number_format = "mm/dd/yyyy hh:mm"
    sheet.freeze_panes = "D2"


def _cell(value: object, header: str = "") -> object:
    if value is None or pd.isna(value):
        return None
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (date, datetime)):
        return value
    if header.endswith("_date"):
        try:
            timestamp = pd.to_datetime(value, utc=True)
            if not pd.isna(timestamp):
                return timestamp.tz_localize(None).to_pydatetime()
        except (TypeError, ValueError, OverflowError):
            pass  # Keep invalid source dates visible for review.
    if header in {"credit_hours", "billable_hours", "payment_year"}:
        try:
            return float(value)
        except (TypeError, ValueError):
            pass
    # Source text is data, even if it begins with an Excel formula character.
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def export_review(review: TuitionReview, frames: dict[str, pd.DataFrame], manifest: dict, path: Path) -> Path:
    """Save exclusively; leave financial calculations in the tested domain layer."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    rate = manifest["exchange_rate"]
    identity = frames["identity"].iloc[0]
    summary = [
        ["Tax year", manifest["tax_year"]], ["Environment", manifest["environment"]],
        ["CWID", identity.get("cwid", identity.get("student_id", identity.get("id", "")))],
        ["Student", " ".join(str(identity.get(key, "")) for key in ("first_name", "last_name"))],
        ["Paid amount (USD)", review.paid_usd], ["Converted amount (CAD)", review.paid_cad],
        ["Amount status", "Calculated for administrator review" if review.paid_cad is not None else "Review required; see Checks"],
        ["Certificate status", manifest["certificate_status"]],
        ["Annual rate (CAD per USD)", Decimal(rate["cad_per_usd"]) if rate.get("cad_per_usd") else None],
        ["Rate year", rate.get("rate_year")], ["Rate method", rate.get("method")],
        ["Fee policy", "Exclude FEIT/CFEE; retain other charge codes under report owner policy"],
        ["Minimum course duration", "21 consecutive scheduled days, including start/end; Summer included"],
        ["Attendance review", "Complete confirmation columns on Sessions before certifying"],
        ["Rounding", "Total USD × annual rate; round final CAD to cents, half up"],
        ["Source extracted (UTC)", manifest.get("extracted_at_utc")],
        ["Output application version", manifest["app_version"]],
    ]
    write_table_sheet(workbook, sheet_name="Summary", table_name="TL11ASummary",
                      headers=["Item", "Value"], rows=[[_cell(v) for v in row] for row in summary])
    sheet = workbook["Summary"]
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 95
    for row in sheet.iter_rows(min_row=2):
        row[1].alignment = Alignment(vertical="top", wrap_text=True)
        sheet.row_dimensions[row[0].row].height = 32
    for row in (6, 7):
        sheet.cell(row, 2).number_format = '"USD "#,##0.00' if row == 6 else '"CAD "#,##0.00'
        sheet.cell(row, 2).font = Font(name="Aptos", size=14, bold=True, color="1F4E78")
    sheet.cell(10, 2).number_format = "0.0000"

    # Order shows the usable output first, then allocation build and source evidence.
    tables = [
        ("Sessions", review.sessions),
        ("Tuition Charges", review.charges.drop(columns=["pidm", "extract_row_count", "extract_tax_year"], errors="ignore")),
        ("Applications", review.applications.drop(columns=["pidm", "extract_row_count", "extract_tax_year"], errors="ignore")),
        ("Courses", frames["enrollment_prepared"].drop(columns=["pidm", "extract_row_count", "extract_tax_year"], errors="ignore")),
        ("Programs", frames["programs"].drop(columns=["pidm", "extract_row_count", "extract_tax_year"], errors="ignore")),
        ("Rate and Policy", pd.DataFrame([
            {"Item": key, "Value": json_value} for key, json_value in rate.items()
        ] + [{"Item": "Source package", "Value": manifest.get("source_package", "Current Insights extracts")},
             {"Item": "Query/policy audit", "Value": "See accompanying manifest.json and complete raw CSVs"}])),
        ("Reconciliation", review.reconciliation), ("Checks", review.checks),
    ]
    for index, (name, frame) in enumerate(tables, 1):
        headers = list(frame.columns)
        write_table_sheet(workbook, sheet_name=name, table_name=f"TL11ATable{index}", headers=headers,
            rows=([_cell(value, header) for header, value in zip(headers, row)] for row in frame.itertuples(index=False, name=None)),
            currency_headers={col for col in headers if col.endswith("_usd")},
            count_headers={"course_count", "issue_count", "scheduled_duration_days"})
        ws = workbook[name]
        ws.row_dimensions[1].height = 48
        for cell in ws[1]:
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for col, header in enumerate(headers, 1):
            if header in {"session_start", "session_end"} or header.endswith("_date"):
                for cells in ws.iter_cols(min_col=col, max_col=col, min_row=2):
                    for cell in cells:
                        cell.number_format = "mm/dd/yyyy"
            if header in {"detail", "Value"}:
                ws.column_dimensions[ws.cell(1, col).column_letter].width = 110
                for cells in ws.iter_cols(min_col=col, max_col=col, min_row=2):
                    for cell in cells:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
                        ws.row_dimensions[cell.row].height = 48
            if header.endswith("_confirmation"):
                for cells in ws.iter_cols(min_col=col, max_col=col, min_row=2):
                    for cell in cells:
                        cell.font = Font(name="Aptos", size=12, color="0000FF")
    candidate = frames["transactions_prepared"]
    candidate = candidate[candidate.year_term_candidate.eq("Y")]
    _write_manual_source(workbook, candidate)
    _write_manual_final(workbook, review, candidate, manifest, identity)
    workbook.calculation.fullCalcOnLoad = True
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        with path.open("xb") as handle:
            try:
                workbook.save(handle)
            except Exception:
                path.unlink(missing_ok=True)
                raise
    finally:
        workbook.close()
    return path
