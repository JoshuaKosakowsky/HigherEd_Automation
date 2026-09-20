"""Convert a TouchNet transaction-results CSV into the JPMLB workbook."""

from __future__ import annotations

import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SOURCE_FILENAME_PATTERN = re.compile(
    r"^Transaction_Results_\d{2}_\d{2}_\d{4}_\d{2}_\d{2}_\d{2}\.csv$"
)
MINIMUM_SOURCE_COLUMNS = 16
ACCOUNTING_FORMAT = (
    '_($* #,##0.00_);_($* (#,##0.00);_($* "-"??_);_(@_)'
)


class JPMLBTransformationError(RuntimeError):
    """Raised when a JPMLB workbook cannot be created safely."""


@dataclass(frozen=True)
class JPMLBTransformationResult:
    """Summary of a successfully created JPMLB workbook."""

    source_path: Path
    output_path: Path
    data_row_count: int
    source_column_count: int
    output_column_count: int


def _read_source(path: Path) -> pd.DataFrame:
    if not path.is_file():
        raise JPMLBTransformationError(f"Source file was not found: {path}")

    if not SOURCE_FILENAME_PATTERN.fullmatch(path.name):
        raise JPMLBTransformationError(
            "Source filename must match "
            "Transaction_Results_MM_DD_YYYY_HH_MM_SS.csv: "
            f"{path.name}"
        )

    last_decode_error: UnicodeDecodeError | None = None

    for encoding in ("utf-8-sig", "cp1252"):
        try:
            frame = pd.read_csv(
                path,
                header=None,
                dtype=object,
                keep_default_na=False,
                encoding=encoding,
            )
            return frame.fillna("")
        except UnicodeDecodeError as exc:
            last_decode_error = exc
        except pd.errors.EmptyDataError as exc:
            raise JPMLBTransformationError(
                f"Source CSV is empty: {path.name}"
            ) from exc
        except pd.errors.ParserError as exc:
            raise JPMLBTransformationError(
                f"Source CSV could not be parsed: {path.name}: {exc}"
            ) from exc

    raise JPMLBTransformationError(
        f"Source CSV encoding could not be read: {path.name}"
    ) from last_decode_error


def _transform_columns(source: pd.DataFrame) -> pd.DataFrame:
    if len(source.index) < 2:
        raise JPMLBTransformationError(
            "No populated data rows were found beneath the CSV header."
        )

    source_column_count = len(source.columns)
    if source_column_count < MINIMUM_SOURCE_COLUMNS:
        raise JPMLBTransformationError(
            "The CSV does not contain enough columns for the JPMLB layout: "
            f"found {source_column_count}; expected at least "
            f"{MINIMUM_SOURCE_COLUMNS}."
        )

    # Match the VBA macro exactly: delete original columns L:N, then insert
    # new CWID and NAME columns at resulting positions O:P.
    transformed = pd.concat(
        [source.iloc[:, :11], source.iloc[:, 14:]],
        axis=1,
        ignore_index=True,
    )

    while len(transformed.columns) < 14:
        transformed[len(transformed.columns)] = ""

    inserted = pd.DataFrame(
        {
            0: ["CWID", *([""] * (len(transformed.index) - 1))],
            1: ["NAME", *([""] * (len(transformed.index) - 1))],
        }
    )

    return pd.concat(
        [
            transformed.iloc[:, :14],
            inserted,
            transformed.iloc[:, 14:],
        ],
        axis=1,
        ignore_index=True,
    )


def _parse_amount(value: Any, *, row_number: int) -> float | int | str:
    if value is None or value == "":
        return ""

    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value

    text = str(value).strip()
    if not text:
        return ""

    is_parenthesized = text.startswith("(") and text.endswith(")")
    normalized = text.replace("$", "").replace(",", "")
    if is_parenthesized:
        normalized = f"-{normalized[1:-1]}"

    try:
        number = Decimal(normalized)
    except InvalidOperation as exc:
        raise JPMLBTransformationError(
            f"Amount column M contains a nonnumeric value on CSV row "
            f"{row_number}: {text!r}."
        ) from exc

    if number == number.to_integral_value():
        return int(number)
    return float(number)


def _excel_value(value: Any) -> Any:
    if value is None:
        return ""

    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass

    return value


def _set_summary_style(
    worksheet: Any,
    *,
    cell_range: str,
    fill_color: str,
) -> None:
    fill = PatternFill(fill_type="solid", fgColor=fill_color)
    for row in worksheet[cell_range]:
        for cell in row:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(wrap_text=False)


def _autofit_columns(worksheet: Any, last_column: int) -> None:
    for column_number in range(1, last_column + 1):
        column_letter = get_column_letter(column_number)
        maximum_length = 0

        for cell in worksheet[column_letter]:
            if cell.value is None:
                continue
            maximum_length = max(maximum_length, len(str(cell.value)))

        worksheet.column_dimensions[column_letter].width = min(
            max(maximum_length + 2, 2),
            60,
        )


def _build_workbook(
    *,
    source_path: Path,
    transformed: pd.DataFrame,
) -> Workbook:
    workbook = Workbook()
    worksheet = workbook.active
    safe_title = re.sub(r"[\\/*?:\[\]]", "_", source_path.stem)[:31]
    worksheet.title = safe_title or "Transaction Results"

    for row_number, row in enumerate(
        transformed.itertuples(index=False, name=None),
        start=1,
    ):
        for column_number, value in enumerate(row, start=1):
            if column_number == 13 and row_number >= 2:
                cell_value = _parse_amount(value, row_number=row_number)
            else:
                cell_value = _excel_value(value)
            worksheet.cell(row=row_number, column=column_number, value=cell_value)

    last_data_row = len(transformed.index)
    last_data_column = max(len(transformed.columns), 16)
    total_import_row = last_data_row + 1
    total_post_row = total_import_row + 5
    void_row = total_post_row - 2
    scholarship_row = total_post_row + 3
    lockbox_row = scholarship_row + 3
    unclaimed_row = lockbox_row + 3

    worksheet.cell(total_import_row, 11, "Total Import")
    worksheet.cell(
        total_import_row,
        13,
        f"=SUM(M2:M{last_data_row})",
    )
    _set_summary_style(
        worksheet,
        cell_range=f"K{total_import_row}:M{total_import_row}",
        fill_color="FFFF00",
    )

    worksheet.cell(void_row, 10, "Void Above, then post")
    _set_summary_style(
        worksheet,
        cell_range=f"J{void_row}:M{void_row}",
        fill_color="D9D2E9",
    )

    worksheet.cell(total_post_row, 11, "Total Post")
    worksheet.cell(
        total_post_row,
        13,
        f"=SUM(M{total_post_row - 3}:M{total_post_row - 1})",
    )
    _set_summary_style(
        worksheet,
        cell_range=f"K{total_post_row}:M{total_post_row}",
        fill_color="FFFF00",
    )

    worksheet.cell(scholarship_row, 11, "Scholarships: FA")
    worksheet.cell(
        scholarship_row,
        13,
        f"=SUM(M{scholarship_row - 1})",
    )
    _set_summary_style(
        worksheet,
        cell_range=f"K{scholarship_row}:M{scholarship_row}",
        fill_color="DDEBF7",
    )

    worksheet.cell(lockbox_row, 11, "Total Lockbox")
    worksheet.cell(
        lockbox_row,
        13,
        f"=M{total_import_row}+M{total_post_row}+M{scholarship_row}",
    )
    _set_summary_style(
        worksheet,
        cell_range=f"K{lockbox_row}:M{lockbox_row}",
        fill_color="E2EFDA",
    )

    worksheet.cell(unclaimed_row, 11, "Unclaimed")
    worksheet.cell(
        unclaimed_row,
        13,
        f"=SUM(M{unclaimed_row - 1})",
    )
    _set_summary_style(
        worksheet,
        cell_range=f"K{unclaimed_row}:M{unclaimed_row}",
        fill_color="F4B084",
    )

    for cell in worksheet[1][:last_data_column]:
        cell.alignment = Alignment(
            vertical="center",
            wrap_text=True,
        )

    worksheet.auto_filter.ref = (
        f"A1:{get_column_letter(last_data_column)}{last_data_row}"
    )

    _autofit_columns(worksheet, last_data_column)
    fixed_widths = {
        "A": 14,
        "F": 9,
        "G": 9,
        "I": 8,
        "J": 8,
        "K": 6,
        "M": 13,
        "N": 9,
        "O": 10,
    }
    for column_letter, width in fixed_widths.items():
        worksheet.column_dimensions[column_letter].width = width
    worksheet.column_dimensions["E"].hidden = True

    worksheet.column_dimensions["M"].number_format = ACCOUNTING_FORMAT
    worksheet.column_dimensions["O"].number_format = "0"
    worksheet.column_dimensions["P"].number_format = "@"
    for cell in worksheet["M"]:
        cell.number_format = ACCOUNTING_FORMAT
    for cell in worksheet["O"]:
        cell.number_format = "0"
    for cell in worksheet["P"]:
        cell.number_format = "@"

    for row_number in range(total_import_row, unclaimed_row + 1):
        worksheet.cell(row_number, 11).alignment = Alignment(wrap_text=False)

    worksheet.freeze_panes = "A2"
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    return workbook


def _verify_saved_workbook(
    path: Path,
    *,
    last_data_row: int,
) -> None:
    try:
        workbook = load_workbook(path, data_only=False, read_only=False)
    except Exception as exc:
        raise JPMLBTransformationError(
            f"Created workbook could not be reopened: {path}"
        ) from exc

    try:
        worksheet = workbook.active
        total_import_row = last_data_row + 1
        if worksheet.cell(1, 15).value != "CWID":
            raise JPMLBTransformationError(
                "Created workbook failed validation: CWID header is missing."
            )
        if worksheet.cell(1, 16).value != "NAME":
            raise JPMLBTransformationError(
                "Created workbook failed validation: NAME header is missing."
            )
        if worksheet.cell(total_import_row, 13).value != (
            f"=SUM(M2:M{last_data_row})"
        ):
            raise JPMLBTransformationError(
                "Created workbook failed validation: Total Import formula is missing."
            )
        if worksheet.freeze_panes != "A2":
            raise JPMLBTransformationError(
                "Created workbook failed validation: top row is not frozen."
            )
    finally:
        workbook.close()


def transform_jpmlb_csv(
    *,
    source_path: Path,
    output_path: Path,
) -> JPMLBTransformationResult:
    """Create a new JPMLB workbook without overwriting an existing file."""

    source_path = Path(source_path)
    output_path = Path(output_path)

    if output_path.suffix.lower() != ".xlsx":
        raise JPMLBTransformationError(
            f"JPMLB output must be an .xlsx file: {output_path}"
        )
    if output_path.exists():
        raise JPMLBTransformationError(
            f"Output file already exists and was not overwritten: {output_path}"
        )
    if not output_path.parent.is_dir():
        raise JPMLBTransformationError(
            f"Destination folder does not exist: {output_path.parent}"
        )

    source = _read_source(source_path)
    transformed = _transform_columns(source)
    workbook = _build_workbook(
        source_path=source_path,
        transformed=transformed,
    )

    try:
        workbook.save(output_path)
        workbook.close()
        _verify_saved_workbook(
            output_path,
            last_data_row=len(transformed.index),
        )
    except Exception:
        workbook.close()
        if output_path.exists():
            output_path.unlink()
        raise

    return JPMLBTransformationResult(
        source_path=source_path,
        output_path=output_path,
        data_row_count=len(source.index) - 1,
        source_column_count=len(source.columns),
        output_column_count=len(transformed.columns),
    )
