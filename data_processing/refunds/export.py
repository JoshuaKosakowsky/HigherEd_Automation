from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment
import pandas as pd

from data_processing.shared.xlsx_output_format import (
    apply_default_font,
    format_table_columns,
    style_header_row,
)

from .allocation import CARD_PAYMENT_CODES, REPORT_COLUMNS


REFUND_SHEETS = (
    "Transact Refunds", "Check Refunds", "Parent Refunds", "System Refunds",
    "Third Party Reviews", "Refund Holds", "ACH Clearing", "ACH Reviews",
    "Mines Park Reviews", "Manual Reviews",
)
WORKBOOK_COLUMNS = REPORT_COLUMNS[:3] + [
    "tab_delivery", "tab_refund_amount", "tab_review_note",
] + REPORT_COLUMNS[3:]
DELIVERY_SHEETS = {
    "AFRD (Transact)": "Transact Refunds",
    **{f"{code} (Transact)": "Transact Refunds" for code in CARD_PAYMENT_CODES},
    "RFND (CHECK)": "Check Refunds",
    "ARFD (System)": "System Refunds",
    "Refund Hold - Student": "Refund Holds",
    "AFRD (Transact) - May Be Too Old": "ACH Reviews",
    "ACHK Date Review": "ACH Reviews",
}
COMPONENT_AMOUNT = re.compile(r"^(.*?)\s+(\d+\.\d{2})$")
ACH_WAIT = re.compile(r"^ACHK Clearing Wait until \d{2}/\d{2}/\d{4}$")


@dataclass(frozen=True)
class RefundPortion:
    sheet: str
    delivery: str
    amount: Decimal


def _amount(value: object) -> Decimal | None:
    if value is None or pd.isna(value):
        return None
    try:
        amount = Decimal(str(value))
    except InvalidOperation:
        return None
    return amount if amount.is_finite() else None


def _student_portions(delivery: str, amount: Decimal) -> list[RefundPortion]:
    """Read the allocator's delivery display; reject unknown or unreconciled parts.

    Single-method displays omit amounts; mixed displays explicitly include them.
    Never guess a mixed allocation or send an unfamiliar code to a payment tab.
    """
    parts = delivery.split("; ")
    portions = []
    for part in parts:
        label = part
        portion_amount = amount
        if len(parts) > 1:
            match = COMPONENT_AMOUNT.fullmatch(part)
            if match is None:
                raise ValueError("Unrecognized mixed student delivery amounts.")
            label, raw_amount = match.groups()
            label = label.removesuffix(" /")
            portion_amount = Decimal(raw_amount)
        sheet = "ACH Clearing" if ACH_WAIT.fullmatch(label) else DELIVERY_SHEETS.get(label)
        if sheet is None or portion_amount <= 0:
            raise ValueError("Unrecognized student delivery; verify the recipient and method.")
        portions.append(RefundPortion(sheet, label, portion_amount))
    if sum((part.amount for part in portions), Decimal(0)) != amount:
        raise ValueError("Student delivery components do not equal the student refund amount.")
    return portions


def _tab_rows(row: dict[str, object]) -> list[tuple[str, dict[str, object]]]:
    """Partition one account's refund without changing any allocation fields."""
    total = _amount(row["total_refund_amount"])

    def review(sheet: str, note: str) -> list[tuple[str, dict[str, object]]]:
        return [(sheet, {**row, "tab_delivery": "REVIEW REQUIRED",
                         "tab_refund_amount": total, "tab_review_note": note})]

    # Account-level reviews stay together rather than appearing in payment queues.
    if row["third_party_review_required_ind"] == "Y":
        return review("Third Party Reviews", "Resolve third-party ownership before issuing refunds.")
    if "Mines Park Charge - Review" in str(row["review_reasons"]):
        return review("Mines Park Reviews", "Review housing charges before issuing refunds.")

    student = _amount(row["student_refund_amount"])
    parent = _amount(row["parent_refund_amount"])
    if (student is None or parent is None or total is None
            or min(student, parent) < 0 or student + parent != total or total <= 0):
        return review("Manual Reviews", "Recipient amounts are undetermined or do not reconcile.")

    portions = []
    if student > 0:
        try:
            portions.extend(_student_portions(str(row["proposed_student_delivery"]), student))
        except ValueError as error:
            return review("Manual Reviews", str(error))
    if parent > 0:
        if row["proposed_parent_delivery"] != "RFDP":
            return review("Manual Reviews", "Parent refund delivery requires review.")
        portions.append(RefundPortion("Parent Refunds", "RFDP", parent))

    by_sheet: dict[str, list[RefundPortion]] = defaultdict(list)
    for portion in portions:
        by_sheet[portion.sheet].append(portion)
    return [
        (sheet, {**row,
                 "tab_delivery": "; ".join(f"{part.delivery} {part.amount:.2f}" for part in parts),
                 "tab_refund_amount": sum((part.amount for part in parts), Decimal(0)),
                 "tab_review_note": None})
        for sheet, parts in by_sheet.items()
    ]


CURRENCY_COLUMNS = {
    "tab_refund_amount",
    "full_account_balance",
    "total_refund_amount",
    "student_refund_amount",
    "parent_refund_amount",
    "target_term_fdpl_amount",
    "unused_fdpl_amount",
    "unused_non_fdpl_amount",
    "total_unused_payment_amount",
    "unpaid_charge_amount",
    "previous_term_balance_before_current_payments",
    "prior_terms_balance_before_current_payments",
    "title_iv_applied_to_older_fiscal_years",
    "unrestricted_applied_to_older_terms",
    "original_payment_total",
}

COUNT_COLUMNS = {
    "fdpl_row_count",
    "original_payment_row_count",
    "plus_auth_row_count",
    "negative_source_count",
    "ambiguous_source_pool_count",
}


def export_refund_report(report: pd.DataFrame, output_file: Path) -> Path:
    """Write plain worksheets by refund method, keeping account totals as context."""
    missing = [column for column in REPORT_COLUMNS if column not in report.columns]
    if missing:
        raise ValueError(f"Refund report is missing columns: {', '.join(missing)}")

    workbook = Workbook()
    workbook.remove(workbook.active)
    grouped: dict[str, list[dict[str, object]]] = defaultdict(list)
    for row in report[REPORT_COLUMNS].to_dict("records"):
        for sheet, tab_row in _tab_rows(row):
            grouped[sheet].append(tab_row)

    for sheet in REFUND_SHEETS:
        worksheet = workbook.create_sheet(sheet)
        worksheet.freeze_panes = "A2"
        worksheet.append(WORKBOOK_COLUMNS)
        for row in grouped[sheet]:
            worksheet.append([
                None if pd.isna(row[header]) else row[header] for header in WORKBOOK_COLUMNS
            ])
        apply_default_font(worksheet)
        style_header_row(worksheet)
        format_table_columns(
            worksheet, headers=WORKBOOK_COLUMNS,
            currency_headers=CURRENCY_COLUMNS, count_headers=COUNT_COLUMNS,
        )
        for cell in worksheet[1]:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        worksheet.row_dimensions[1].height = 48
        worksheet.column_dimensions["D"].width = 60
        worksheet.auto_filter.ref = worksheet.dimensions
        for header in (
            "deceased_date", "last_ar_activity_date", "account_control_activity_date",
            "ed_activity_date", "plus_auth_activity_date",
        ):
            column_number = WORKBOOK_COLUMNS.index(header) + 1
            for cells in worksheet.iter_cols(min_col=column_number, max_col=column_number, min_row=2):
                for cell in cells:
                    cell.number_format = (
                        "mm/dd/yyyy" if header == "deceased_date" else "mm/dd/yyyy h:mm AM/PM"
                    )

    output_file.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_file)
    return output_file
