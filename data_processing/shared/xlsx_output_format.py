from __future__ import annotations

from collections.abc import Iterable

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import (
    Table,
    TableStyleInfo,
)
from openpyxl.worksheet.worksheet import Worksheet


DEFAULT_FONT_NAME = "Aptos"
DEFAULT_FONT_SIZE = 12

HEADER_FILL = PatternFill(
    "solid",
    fgColor="1F4E78",
)

BODY_FONT = Font(
    name=DEFAULT_FONT_NAME,
    size=DEFAULT_FONT_SIZE,
)

HEADER_FONT = Font(
    name=DEFAULT_FONT_NAME,
    size=DEFAULT_FONT_SIZE,
    color="FFFFFF",
    bold=True,
)

CURRENCY_FORMAT = (
    '$#,##0.00;[Red]($#,##0.00);-'
)

COUNT_FORMAT = "#,##0"


def apply_default_font(
    worksheet: Worksheet,
) -> None:
    """
    Apply the standard font to every populated cell.
    """
    for row in worksheet.iter_rows():
        for cell in row:
            if cell.value is not None:
                cell.font = BODY_FONT


def style_header_row(
    worksheet: Worksheet,
    *,
    row_number: int = 1,
) -> None:
    """
    Apply the standard header formatting.
    """
    for cell in worksheet[row_number]:
        if cell.value is None:
            continue

        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )


def format_table_columns(
    worksheet: Worksheet,
    *,
    headers: list[str],
    currency_headers: set[str] | None = None,
    count_headers: set[str] | None = None,
    max_column_width: int = 38,
) -> None:
    """
    Apply number formats and readable column widths.
    """
    currency_headers = currency_headers or set()
    count_headers = count_headers or set()

    for column_number, header in enumerate(
        headers,
        start=1,
    ):
        column_letter = get_column_letter(
            column_number
        )

        if header in currency_headers:
            for cell in worksheet[
                column_letter
            ][1:]:
                cell.number_format = CURRENCY_FORMAT

        if header in count_headers:
            for cell in worksheet[
                column_letter
            ][1:]:
                cell.number_format = COUNT_FORMAT

        max_length = max(
            (
                len(str(cell.value))
                for cell in worksheet[
                    column_letter
                ]
                if cell.value is not None
            ),
            default=0,
        )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(max_length + 2, 11),
            max_column_width,
        )


def add_standard_table(
    worksheet: Worksheet,
    *,
    table_name: str,
) -> None:
    """
    Add the standard Excel table style.
    """
    if worksheet.max_row <= 1:
        return

    table = Table(
        displayName=table_name,
        ref=worksheet.dimensions,
    )

    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )

    worksheet.add_table(table)


def write_table_sheet(
    workbook: Workbook,
    *,
    sheet_name: str,
    table_name: str,
    headers: list[str],
    rows: Iterable[Iterable[object]],
    currency_headers: set[str] | None = None,
    count_headers: set[str] | None = None,
) -> None:
    """
    Create a consistently formatted Excel table sheet.
    """
    worksheet = workbook.create_sheet(
        title=sheet_name
    )

    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A2"
    worksheet.append(headers)

    for row in rows:
        worksheet.append(list(row))

    apply_default_font(worksheet)
    style_header_row(worksheet)

    format_table_columns(
        worksheet,
        headers=headers,
        currency_headers=currency_headers,
        count_headers=count_headers,
    )

    add_standard_table(
        worksheet,
        table_name=table_name,
    )