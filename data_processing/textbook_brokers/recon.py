"""Build the monthly Textbook Brokers to Banner reconciliation workbook."""

from __future__ import annotations

import csv
import re
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from shared.progress import ProgressReporter

from openpyxl import Workbook, load_workbook
from openpyxl.pivot.cache import CacheDefinition, CacheField, CacheSource, WorksheetSource
from openpyxl.pivot.record import RecordList
from openpyxl.pivot.table import DataField, Location, PivotField, PivotTableStyle, RowColField, TableDefinition
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


MONTH_PATTERN = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")
MONEY = Decimal("0.01")
RECON_SHEETS = {
    "IA": ("IA Charge Report", "IA Pivot", "FRST", "FRST Pivot", "IA Recon"),
    "FA": ("FA Charge Report", "FA Pivot", "BOOK", "BOOK Pivot", "BOOK Recon"),
}
IA_HEADERS = ("Identification", "First Name", "Last Name", "Sub Total", "Tax", "Total",
              "Type", "Payment Method", "Processed By", "Payment Info")
FA_HEADERS = ("Student ID", "First", "Last", "Transaction ID", "Note", "Date", "Year",
              "Term", "Subtotal", "Tax", "Spent", "Payment Info")
BANNER_HEADERS = (
    "'ID'", "'Name'", "'Detail Code'", "'Description'", "'Amount'",
    "'Balance'", "'Term'", "'Aid Year'", "'Period'", "'Transaction Number'",
    "'Transaction Number Paid'", "'Receipt'", "'Source Code'",
    "'Cashier User ID'", "'Cashier Session'", "'Cashier End Date'",
    "'Course Reference Number'", "'Cross Reference ID'",
    "'Cross Reference Source'", "'Cross Reference Number'",
    "'Contract Payment'", "'Feed Indicator'", "'Feed Document'",
    "'Feed Date'", "'Invoice Number'", "'Invoice Number Paid'",
    "'Invoice Statement Date'", "'Effective Date'", "'Bill Date'",
    "'Due Date'", "'Activity Date'",
)
MONEY_FORMAT = '"$"#,##0.00;[Red]("$"#,##0.00)'
BLUE_HEADER = "5B8FC9"
GREEN_HEADER = "85B965"
BLUE_TAB_LIGHT = "BDD7EE"
GREEN_TAB_LIGHT = "C6E7B8"
BLUE_TAB_DARK = "4472A8"
GREEN_TAB_DARK = "609447"


def _key(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").strip("' \t").casefold())


def _id(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _money(value: object, location: str) -> Decimal:
    if value is None or str(value).strip() == "":
        raise ValueError(f"Missing amount at {location}.")
    try:
        result = Decimal(str(value).replace(",", "").replace("$", "").strip())
    except InvalidOperation as error:
        raise ValueError(f"Invalid amount at {location}.") from error
    if not result.is_finite() or result != result.quantize(MONEY):
        raise ValueError(f"Amount must have cents precision at {location}.")
    return result


def _rows(path: Path, sheet_name: str | None = None) -> tuple[list[str], list[list[object]]]:
    if path.suffix.casefold() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as source:
            rows = list(csv.reader(source))
    elif path.suffix.casefold() in {".xlsx", ".xlsm"}:
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet_name is None:
                sheet = book.active
            else:
                matches = [sheet for sheet in book if _key(sheet.title) == _key(sheet_name)]
                if not matches:
                    raise ValueError(f"{path.name} has no {sheet_name} tab.")
                sheet = matches[0]
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        finally:
            book.close()
    else:
        raise ValueError(f"Unsupported input file: {path.name}.")
    for index, row in enumerate(rows):
        if any(value is not None and str(value).strip() for value in row):
            return [str(value or "") for value in row], rows[index + 1 :]
    raise ValueError(f"No headers found in {path.name}.")


def _column(headers: list[str], choices: tuple[str, ...], *, required: bool = True) -> int | None:
    names = {_key(header): index for index, header in enumerate(headers)}
    for choice in choices:
        if _key(choice) in names:
            return names[_key(choice)]
    if required:
        raise ValueError(f"Missing column: {' or '.join(choices)}.")
    return None


def _cell(row: list[object], index: int | None) -> object:
    return row[index] if index is not None and index < len(row) else None


def _broker_rows(path: Path, kind: str) -> tuple[list[tuple[object, ...]], dict[str, tuple[str, str, Decimal]]]:
    headers, rows = _rows(path, RECON_SHEETS[kind][0])
    id_col = _column(headers, ("Customer ID", "Identification", "Student ID", "CWID"))
    alternate_id = _column(headers, ("Student ID",), required=False)
    first_col = _column(headers, ("First Name", "First"))
    last_col = _column(headers, ("Last Name", "Last"))
    subtotal_col = _column(headers, ("Sub Total", "Subtotal"), required=False)
    tax_col = _column(headers, ("Tax",), required=False)
    total_col = _column(headers, ("Total", "Spent"))
    source_rows = []
    totals: dict[str, Decimal] = defaultdict(Decimal)
    names: dict[str, tuple[str, str]] = {}
    for number, row in enumerate(rows, 2):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        primary_id = _id(_cell(row, id_col))
        secondary_id = _id(_cell(row, alternate_id))
        if primary_id and secondary_id and primary_id != secondary_id:
            raise ValueError(f"Conflicting student IDs in {path.name}, {kind} row {number}.")
        cwid = primary_id or secondary_id
        first = str(_cell(row, first_col) or "").strip()
        last = str(_cell(row, last_col) or "").strip()
        if not cwid and not first and not last:
            # Vendor exports and older recon sheets may end with a total row.
            continue
        if not cwid:
            raise ValueError(f"Blank student ID in {path.name}, {kind} row {number}.")
        amount = _money(_cell(row, total_col), f"{path.name}, {kind} row {number}")
        subtotal = _cell(row, subtotal_col)
        tax = _cell(row, tax_col)
        # The IA and FA report layouts differ; preserve their financial fields.
        if kind == "IA":
            source_rows.append((cwid, first, last, subtotal, tax, float(amount), None, None, None, None))
        else:
            source_rows.append((cwid, first, last, None, None, None, None, None,
                                subtotal, tax, float(amount), None))
        totals[cwid] += amount
        if cwid not in names or (not names[cwid][0] and first):
            names[cwid] = first, last
    source_rows.sort(key=lambda row: (str(row[2] or "").casefold(), str(row[1] or "").casefold(), str(row[0])))
    return source_rows, {cwid: (*names[cwid], total) for cwid, total in totals.items()}


def _banner_rows(path: Path | None, code: str, month: str | None = None) -> tuple[list[str], list[tuple[object, ...]], dict[str, tuple[str, Decimal]]]:
    if path is None:
        return [], [], {}
    selected_sheet = None
    if path.suffix.casefold() in {".xlsx", ".xlsm"}:
        book = load_workbook(path, read_only=True)
        try:
            if any(_key(name) == _key(code) for name in book.sheetnames):
                selected_sheet = code
        finally:
            book.close()
    headers, rows = _rows(path, selected_sheet)
    id_col = _column(headers, ("'ID'", "ID", "CWID", "Student ID"))
    name_col = _column(headers, ("'Name'", "Name"))
    amount_col = _column(headers, ("'Amount'", "Amount"))
    code_col = _column(headers, ("'Detail Code'", "Detail Code"), required=False)
    feed_col = _column(headers, ("'Feed Date'", "Feed Date"), required=False)
    totals: dict[str, Decimal] = defaultdict(Decimal)
    names: dict[str, str] = {}
    kept = []
    for number, row in enumerate(rows, 2):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        found_code = str(_cell(row, code_col) or "").strip().upper()
        if code_col is not None and found_code != code:
            raise ValueError(f"Expected {code} only in {path.name}; found {found_code or 'blank'} at row {number}.")
        feed_value = _cell(row, feed_col)
        if month and feed_value is not None and str(feed_value).strip():
            try:
                feed_date = feed_value.date() if isinstance(feed_value, datetime) else (
                    feed_value if isinstance(feed_value, date) else date.fromisoformat(str(feed_value)[:10])
                )
            except ValueError as error:
                raise ValueError(f"Invalid Feed Date in {path.name}, row {number}.") from error
            if feed_date.strftime("%Y-%m") != month:
                raise ValueError(f"Feed Date outside {month} in {path.name}, row {number}.")
        cwid = _id(_cell(row, id_col))
        if not cwid:
            raise ValueError(f"Blank student ID in {path.name}, row {number}.")
        amount = _money(_cell(row, amount_col), f"{path.name}, row {number}")
        name = str(_cell(row, name_col) or "").strip()
        kept.append(tuple(row))
        totals[cwid] += amount
        if cwid not in names or (not names[cwid] and name):
            names[cwid] = name
    kept.sort(key=lambda row: (str(_cell(list(row), name_col) or "").casefold(),
                               _id(_cell(list(row), id_col))))
    return headers, kept, {cwid: (names[cwid], amount) for cwid, amount in totals.items()}


def _replace_data(sheet, rows: list[tuple[object, ...]], width: int) -> None:
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    for row in rows:
        sheet.append(list(row[:width]) + [None] * max(0, width - len(row)))


def _sort_key(cwid: str, broker: dict, banner: dict) -> tuple[str, str]:
    if cwid in broker:
        first, last, _ = broker[cwid]
        return (last or first).casefold(), cwid
    return banner[cwid][0].casefold(), cwid


def _write_recon(sheet, broker: dict, banner: dict) -> None:
    if sheet.max_row > 3:
        sheet.delete_rows(4, sheet.max_row - 3)
    sheet["K3"] = "Banner, not TBB"
    sheet["L3"] = "TBB, not Banner"
    for row_index, cwid in enumerate(sorted(broker.keys() | banner.keys(), key=lambda value: _sort_key(value, broker, banner)), 4):
        if cwid in broker:
            first, last, amount = broker[cwid]
            for col, value in zip("ABCD", (cwid, first, last, float(amount))):
                sheet[f"{col}{row_index}"] = value
        if cwid in banner:
            name, amount = banner[cwid]
            for col, value in zip("GHI", (cwid, name, float(amount))):
                sheet[f"{col}{row_index}"] = value
        sheet[f"K{row_index}"] = f'=IF(AND(G{row_index}<>"",A{row_index}=""),1,0)'
        sheet[f"L{row_index}"] = f'=IF(AND(A{row_index}<>"",G{row_index}=""),1,0)'
        sheet[f"M{row_index}"] = f"=ABS(ROUND(N(D{row_index})-N(I{row_index}),2))"
        for col in "DIM":
            sheet[f"{col}{row_index}"].number_format = MONEY_FORMAT
        if cwid in broker and cwid in banner and broker[cwid][2] == banner[cwid][1]:
            sheet[f"M{row_index}"].fill = PatternFill("solid", fgColor="E2F0D9")
    sheet.freeze_panes = "A4"


def _make_pivot(sheet, source: str, headers: tuple[str, ...], rows: tuple[int, ...],
                amount: int, title: str, cache_id: int, last_row: int) -> None:
    """Build a native refreshable PivotTable without an external workbook."""
    cache = CacheDefinition(
        cacheSource=CacheSource(type="worksheet", worksheetSource=WorksheetSource(
            ref=f"A1:{get_column_letter(len(headers))}{max(2, last_row)}", sheet=source)),
        cacheFields=[CacheField(name=header, numFmtId=0) for header in headers],
        recordCount=0, refreshOnLoad=True, enableRefresh=True, saveData=False,
        missingItemsLimit=0, createdVersion=8, refreshedVersion=8, minRefreshableVersion=3,
    )
    cache.records = RecordList(r=[])
    fields = []
    for index in range(len(headers)):
        fields.append(PivotField(
            axis="axisRow" if index in rows else None,
            dataField=index == amount, showAll=False,
            compact=False, outline=False, defaultSubtotal=False,
            sortType="ascending" if index == rows[-1] else "manual",
        ))
    pivot = TableDefinition(
        name=f"ReconPivot{cache_id}", cacheId=cache_id, dataCaption="Values",
        location=Location(ref=f"A3:{get_column_letter(len(rows) + 1)}5",
                          firstHeaderRow=1, firstDataRow=1, firstDataCol=len(rows)),
        pivotFields=fields, rowFields=[RowColField(x=index) for index in rows],
        dataFields=[DataField(name=title, fld=amount, subtotal="sum")],
        pivotTableStyleInfo=PivotTableStyle(
            name="PivotStyleLight16", showRowHeaders=True, showColHeaders=True,
            showRowStripes=False, showColStripes=False, showLastColumn=True),
        updatedVersion=8, minRefreshableVersion=3, createdVersion=8,
        useAutoFormatting=True, itemPrintTitles=True, indent=0,
        compact=False, compactData=False, applyWidthHeightFormats=True,
    )
    pivot.cache = cache
    sheet.add_pivot(pivot)


def _header(sheet, row: int, labels: tuple[str, ...], fill: str) -> None:
    for index, label in enumerate(labels, 1):
        cell = sheet.cell(row, index, label)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center")


def _make_workbook() -> Workbook:
    workbook = Workbook()
    workbook.active.title = "IA Recon"
    for name in ("IA Pivot", "IA Charge Report", "FRST Pivot", "FRST",
                 "BOOK Recon", "FA Pivot", "FA Charge Report", "BOOK Pivot", "BOOK"):
        workbook.create_sheet(name)
    for kind, (source, source_pivot, banner, banner_pivot, recon) in RECON_SHEETS.items():
        source_headers = IA_HEADERS if kind == "IA" else FA_HEADERS
        _header(workbook[source], 1, source_headers, BLUE_HEADER)
        _header(workbook[banner], 1, BANNER_HEADERS, GREEN_HEADER)
        broker_pivot_headers = tuple(source_headers[index] for index in (0, 1, 2)) + (
            "Sum of Total" if kind == "IA" else "Sum of Spent",)
        _header(workbook[source_pivot], 3, broker_pivot_headers, BLUE_HEADER)
        _header(workbook[banner_pivot], 3, ("'ID'", "'Name'", "Sum of 'Amount'"), GREEN_HEADER)
        recon_sheet = workbook[recon]
        tab_light = BLUE_TAB_LIGHT if kind == "IA" else GREEN_TAB_LIGHT
        tab_dark = BLUE_TAB_DARK if kind == "IA" else GREEN_TAB_DARK
        recon_sheet.sheet_properties.tabColor = tab_dark
        for name in (source, source_pivot, banner, banner_pivot):
            workbook[name].sheet_properties.tabColor = tab_light
        for start, end, label, color in (
            ("A1", "D1", kind, BLUE_HEADER if kind == "IA" else GREEN_HEADER),
            ("G1", "I1", "Banner", GREEN_HEADER),
            ("K1", "N1", "Difference", "7F8C8D"),
        ):
            for row in recon_sheet[f"{start}:{end}"]:
                for cell in row:
                    cell.fill = PatternFill("solid", fgColor=color)
                    cell.alignment = Alignment(horizontal="centerContinuous", vertical="center")
            cell = recon_sheet[start]
            cell.value = label
            cell.font = Font(bold=True, color="FFFFFF", size=14)
        for col, label in {"A": "Student ID", "B": "First", "C": "Last",
                           "D": "Sum of Spent", "G": "ID", "H": "Name",
                           "I": "Sum of 'Amount'", "K": "Banner, not TBB",
                           "L": "TBB, not Banner", "M": "Difference", "N": "Notes"}.items():
            cell = recon_sheet[f"{col}3"]
            cell.value = label
            cell.fill = PatternFill("solid", fgColor=(BLUE_HEADER if kind == "IA" else GREEN_HEADER)
                                    if col in "ABCD" else GREEN_HEADER if col in "GHI" else "7F8C8D")
            cell.font = Font(bold=True, color="FFFFFF")
        for col, width in {"A": 16, "B": 18, "C": 22, "D": 18, "G": 16,
                           "H": 30, "I": 19, "K": 19, "L": 19, "M": 16, "N": 30}.items():
            recon_sheet.column_dimensions[col].width = width
        recon_sheet.row_dimensions[1].height = 26
        recon_sheet.freeze_panes = "A4"
        for name in (source, banner):
            sheet = workbook[name]
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}1"
            for column in range(1, sheet.max_column + 1):
                sheet.column_dimensions[get_column_letter(column)].width = 20
        for name in (source_pivot, banner_pivot):
            sheet = workbook[name]
            sheet.freeze_panes = "A4"
            for column in range(1, sheet.max_column + 1):
                sheet.column_dimensions[get_column_letter(column)].width = 22
    return workbook


def build_recon(
    *, brokers: Path, frst: Path | None, book: Path | None,
    output: Path, month: str | None = None,
    progress_reporter: ProgressReporter | None = None,
) -> dict[str, int]:
    """Create a complete recon workbook and native PivotTables from Python."""
    paths = [brokers, *(path for path in (frst, book) if path)]
    if any(output.resolve() == path.resolve() for path in paths):
        raise ValueError("Output must differ from every input workbook.")
    if output.exists():
        raise ValueError("Output already exists. Choose a new filename.")
    if not frst and not book:
        raise ValueError("Provide at least one Banner extract.")
    if month is not None and not MONTH_PATTERN.fullmatch(month):
        raise ValueError("Recon month must be YYYY-MM.")
    if progress_reporter:
        progress_reporter.report("Creating reconciliation workbook layout")
    workbook = _make_workbook()
    counts = {}
    for kind, banner_path in (("IA", frst), ("FA", book)):
        source_name, source_pivot_name, code, banner_pivot_name, recon_name = RECON_SHEETS[kind]
        if progress_reporter:
            progress_reporter.report(f"Reading {kind} Textbook Brokers charges")
        source_rows, broker_totals = _broker_rows(brokers, kind)
        if progress_reporter:
            progress_reporter.report(f"Reading {code} Banner transactions")
        banner_headers, banner_rows, banner_totals = _banner_rows(banner_path, code, month)
        if progress_reporter:
            progress_reporter.report(f"Matching and writing {kind} reconciliation")
        _replace_data(workbook[source_name], source_rows, 10 if kind == "IA" else 12)
        banner_width = len(BANNER_HEADERS)
        output_headers = BANNER_HEADERS
        output_keys = {_key(header) for header in output_headers}
        source_positions = {_key(header): index for index, header in enumerate(banner_headers)}
        extra_positions = [index for index, header in enumerate(banner_headers)
                           if _key(header) not in output_keys]
        ordered_rows = [
            tuple(_cell(list(row), source_positions.get(_key(header))) for header in output_headers)
            + tuple(_cell(list(row), index) for index in extra_positions)
            for row in banner_rows
        ]
        _replace_data(workbook[code], ordered_rows, banner_width + len(extra_positions))
        for index, header in enumerate((banner_headers[pos] for pos in extra_positions), banner_width + 1):
            workbook[code].cell(1, index, header)
        _write_recon(workbook[recon_name], broker_totals, banner_totals)
        source_headers = IA_HEADERS if kind == "IA" else FA_HEADERS
        _make_pivot(workbook[source_pivot_name], source_name, source_headers, (0, 1, 2),
                    5 if kind == "IA" else 10,
                    "Sum of Total" if kind == "IA" else "Sum of Spent",
                    1 if kind == "IA" else 3, len(source_rows) + 1)
        _make_pivot(workbook[banner_pivot_name], code, BANNER_HEADERS, (0, 1), 4,
                    "Sum of 'Amount'", 2 if kind == "IA" else 4, len(banner_rows) + 1)
        counts[kind] = len(broker_totals.keys() | banner_totals.keys())
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    output.parent.mkdir(parents=True, exist_ok=True)
    created = False
    try:
        if progress_reporter:
            progress_reporter.report("Saving reconciliation workbook")
        with output.open("xb") as stream:
            created = True
            workbook.save(stream)
    except Exception:
        if created:
            output.unlink(missing_ok=True)
        raise
    finally:
        workbook.close()
    return counts
