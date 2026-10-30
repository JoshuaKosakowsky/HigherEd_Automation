"""Plain review worksheets using the repository's existing Excel formatting."""

from __future__ import annotations

import re
from math import ceil
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from data_processing.shared.xlsx_output_format import apply_default_font, format_table_columns
from shared.cancellation import CancellationToken
from data_processing.shared.workbook_output import save_new_workbook, validate_output
from .reconciliation import COMBINED_HEADERS, Reconciliation
from .sources import header_key, parse_date
from .worktags import StudentWorktags, WORKTAG_HEADERS


EXCEL_MAX_ROWS = 1_048_576
# Follow the employee's Journal Lines Data layout; keep the reconciliation's
# internal column order stable for period totals and activity controls.
COMBINED_COLUMN_ORDER = (7, 11, 1, 9, 6, 4, 13, 14, 12, 3, 0, 8, 10, 19)
COMBINED_OUTPUT_HEADERS = [COMBINED_HEADERS[index] for index in COMBINED_COLUMN_ORDER]
COMBINED_OUTPUT_HEADERS[2] = "Recon Period"
COMBINED_OUTPUT_HEADERS[3] = "Journal Number"
DOC_OUTPUT_HEADERS = ["Fiscal Year", "Period", "Feed Document", "Workday Total Amt",
                      "Banner Total Amt", "Difference", "Notes"]
CWID_OUTPUT_HEADERS = ["Fiscal Year", "Period", "CWID", "Banner Total Amt",
                       "Workday Total Amt", "Net Period Activity", "Notes / Explanation",
                       "Project", "Grant", "Operating Account", "Ledger Acct", "Spend Cat"]
PERIOD_OUTPUT_HEADERS = ["Fiscal Year", "Period", "Quarter", "Workday SIS Amt",
                         "Workday Other Amt", "Banner Amt", "Difference", "Notes"]


def combined_output_row(row: list[object]) -> list[object]:
    values = [row[index] for index in COMBINED_COLUMN_ORDER]
    # The caller has selected rows with six- or seven-digit posting references.
    values[3] = str(row[3])
    values[9] = values[3]
    return values


def export_reconciliation(
    review: Reconciliation, output: Path, *, app_version: str,
    banner_source: str, cancellation: CancellationToken | None = None,
    worktags: StudentWorktags | None = None,
) -> None:
    validate_output(output)
    book = Workbook()
    book.remove(book.active)

    def sheet(name, headers, rows):
        ws = book.create_sheet(name)
        if len(rows) + 1 > EXCEL_MAX_ROWS:
            raise ValueError("A reconciliation sheet exceeds Excel's row limit. Select a shorter range.")
        ws.append(headers)
        for number, row in enumerate(rows, 2):
            if cancellation and number % 500 == 0:
                cancellation.check()
            ws.append([float(value) if isinstance(value, Decimal) else value for value in row])
            # Preserve source text literally, including text starting with '='.
            for cell in ws[number]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                elif isinstance(cell.value, (date, datetime)):
                    cell.number_format = "yyyy-mm-dd"
        apply_default_font(ws)
        for cell in ws[1]:
            cell.font = Font(name="Aptos", size=12, bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[1].height = 44
        # This helper sets widths and number formats; no Excel Table is added.
        currency = {header for header in headers if any(word in header for word in ("Debit", "Credit", "Difference", "Activity"))
                    and header not in {"Included in Activity"}}
        currency |= {"Amount", "Balance", "'Amount'", "'Balance'"}
        currency |= {header for header in headers if header.endswith(" Amt")}
        format_table_columns(ws, headers=headers, currency_headers=currency,
                             count_headers={header for header in headers if "Count" in header})
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        return ws

    try:
        # Filter the Combined view only. Source/financial controls retain all
        # selected activity, including manual rows and invalid posting references.
        combined_rows = [row for row in review.combined if re.fullmatch(r"[0-9]{6,7}", str(row[3]))]
        combined = sheet("1305 Combined", COMBINED_OUTPUT_HEADERS,
                         [combined_output_row(row) for row in combined_rows])
        grey = PatternFill(fill_type="solid", fgColor="FFEFEFEF")
        previous_group = None
        shaded = True
        for number, row in enumerate(combined_rows, 2):
            # Keep the same journal in separate periods visually separate too.
            # Manual activity retains its individual source-row grouping.
            reference = combined.cell(number, 4).value or row[3]
            if str(row[18]).startswith("Other Workday activity"):
                reference = (row[4], row[5])
            group = (row[0], row[1], reference)
            if group != previous_group:
                shaded = not shaded
                previous_group = group
            if shaded:
                for cell in combined[number]:
                    cell.fill = grey
        # Present signed amounts while retaining separate debit/credit matching
        # controls internally: net zero must never authorize substitution.
        documents = []
        for row in review.documents:
            workday_amount, banner_amount = row[4] - row[7], row[5] - row[8]
            documents.append([row[0], row[1], row[3], workday_amount, banner_amount,
                              banner_amount - workday_amount, row[13]])
        sheet("1305 Doc Recon", DOC_OUTPUT_HEADERS, documents)
        students = []
        for row in review.students:
            banner_amount, workday_amount = row[4] - row[5], row[6] - row[7]
            tags = worktags.by_cwid.get(str(row[3]), {}) if worktags else {}
            # Reuse the reconciler's signed Banner + non-SIS Workday net activity.
            students.append([row[0], row[1], row[3], banner_amount, workday_amount,
                row[10], row[14],
                *[tags.get(kind, "") for kind in ("Project", "Grant", "Operating Account")],
                row[15], row[16]])
        # Keep each student's periods together without combining their amounts.
        students.sort(key=lambda row: (row[2], row[0], row[1]))
        student_sheet = sheet("1305 CWID Recon", CWID_OUTPUT_HEADERS, students)
        previous_cwid = None
        shaded = True
        for number, row in enumerate(students, 2):
            if row[2] != previous_cwid:
                shaded = not shaded
                previous_cwid = row[2]
            if shaded:
                for cell in student_sheet[number]:
                    cell.fill = grey
        for column in ("H", "I", "J"):
            for cell in book["1305 CWID Recon"][column][1:]:
                cell.number_format = "@"
        periods = []
        for row in review.periods:
            sis, other, banner_amount = row[3] - row[6], row[9] - row[10], row[4] - row[7]
            periods.append([*row[:3], sis, other, banner_amount, banner_amount - (sis + other), ""])
        sheet("1305 Period Totals", PERIOD_OUTPUT_HEADERS, periods)
        book["1305 Combined"].freeze_panes = "E2"
        for name in ("1305 Doc Recon", "1305 CWID Recon"):
            book[name].freeze_panes = "D2"
        book["1305 Period Totals"].freeze_panes = "D2"
        if worktags:
            worktag_sheet = sheet("Worktags", WORKTAG_HEADERS, worktags.rows)
            for column in ("A", "B", "E"):
                for cell in worktag_sheet[column][1:]:
                    cell.number_format = "@"
        for source in review.sources:
            if any(len(values) > len(source.headers) for _, values in source.rows):
                raise ValueError(f"{source.source} has source rows wider than the headers.")
            if source.source == "Banner":
                columns = [index for index, header in enumerate(source.headers)
                           if header_key(header) not in {"sourcerow", "aidyear", "selectionstatus"}]
                rows = []
                for _, values in source.rows:
                    row = []
                    for index in columns:
                        value = values[index] if index < len(values) else None
                        if index == source.columns["feed_date"]:
                            try:
                                value = parse_date(value)
                            except (ValueError, TypeError, OverflowError):
                                # Keep blank/invalid source values visible for review.
                                pass
                        row.append(value)
                    rows.append(row)
                sheet("Banner Data", [source.headers[index] for index in columns], rows)
                continue
            sheet(f"{source.source} Data", ["Source Row", *source.headers, "Selection Status"],
                [[number, *values, *([None] * (len(source.headers) - len(values))),
                  source.dispositions.get(number, "")] for number, values in source.rows])
        wrap_headers = {"Notes", "Notes / Explanation", "Description / Original Memo",
                        "Status", "Memo / Feed Document", "Issue", "Project", "Grant",
                        "Operating Account", "Ledger Acct", "Spend Cat"}
        for ws in book:
            for cell in ws[1]:
                if cell.value in {"Notes", "Notes / Explanation", "Status", "Review Status"}:
                    ws.column_dimensions[cell.column_letter].width = 45
            for row in ws.iter_rows(min_row=2):
                lines = 1
                for cell in row:
                    header = ws.cell(1, cell.column).value
                    if header in wrap_headers:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
                        width = max(1, ws.column_dimensions[cell.column_letter].width - 2)
                        lines = max(lines, ceil(len(str(cell.value or "")) / width))
                if lines > 1:
                    ws.row_dimensions[row[0].row].height = min(409, lines * 18)

        save_new_workbook(book, output, cancellation=cancellation)
    finally:
        book.close()
