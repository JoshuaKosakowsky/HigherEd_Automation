"""Read source exports without modifying them or assuming fixed column positions."""

from __future__ import annotations

import csv
import math
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from data_processing.shared.tgiaccd import clean_tgiaccd_header
from shared.insights.banner_activity import BannerActivityParameters


ZERO = Decimal("0.00")
CENT = Decimal("0.01")
ALIASES = {
    "date": ("Accounting Date",), "kind": ("Journal Source", "Source"),
    "debit": ("Debit Amount", "Workday Debit Amount", "Ledger Debit Amount", "Workday Debit"),
    "credit": ("Credit Amount", "Workday Credit Amount", "Ledger Credit Amount", "Workday Credit"),
    "memo": ("Memo", "Line Memo", "Journal Line Memo"), "ledger": ("Ledger Account",),
    "reference": ("External Reference", "External Reference ID"),
    "journal": ("Journal Entry", "Journal Number"), "line": ("Journal Line ID",),
    "cwid": ("CWID", "Student ID"),
    "id": ("ID",), "description": ("Description",), "code": ("Detail Code",),
    "amount": ("Amount",), "feed": ("Feed Document",), "feed_date": ("Feed Date",),
    "transaction": ("Transaction Number",), "term": ("Term",),
    "user": ("Transaction User", "Cashier User ID"),
}
REQUIRED = {
    "Workday": ("date", "kind", "debit", "credit", "memo", "ledger"),
    "Banner": ("id", "description", "code", "amount", "feed", "feed_date"),
}


def header_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", clean_tgiaccd_header(value).casefold())


def text(value: object) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    if isinstance(value, (float, Decimal)) and math.isfinite(value) and value == int(value):
        return str(int(value))
    return str(value).strip()


def reference(value: object) -> str:
    value = text(value)
    # Numeric export cells and text such as 12345.0 represent the same document.
    if re.fullmatch(r"[0-9]+\.0+", value):
        value = value.split(".")[0]
    return " ".join(value.split()).upper()


def fiscal_scope(day: date) -> tuple[int, str, str]:
    """Match shared/fiscal_period.ps1: July P01, June P12, ending-year FY."""
    year = day.year + (day.month >= 7)
    period = (day.month + 5) % 12 + 1
    return year, f"P{period:02d}", f"Q{(period - 1) // 3 + 1}"


def parse_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if type(value) is date:
        return value
    if isinstance(value, (float, int)) and not isinstance(value, bool):
        if not 1 <= value <= 2958465:
            raise ValueError("Invalid date")
        return from_excel(value).date()
    value = text(value)
    try:
        # Preserve the source calendar date rather than shifting time zones.
        return datetime.fromisoformat(value.replace("Z", "+00:00")).date()
    except ValueError:
        for pattern in ("%m/%d/%Y", "%m/%d/%y", "%m/%d/%Y %H:%M:%S", "%m/%d/%Y %I:%M:%S %p"):
            try:
                return datetime.strptime(value, pattern).date()
            except ValueError:
                continue
    raise ValueError("Invalid date")


def money(value: object, *, blank_zero: bool = False) -> Decimal:
    value = text(value)
    if not value and blank_zero:
        return ZERO
    if value.startswith("(") and value.endswith(")"):
        value = "-" + value[1:-1]
    try:
        amount = Decimal(value.replace(",", "").replace("$", ""))
        if not amount.is_finite() or amount != amount.quantize(CENT):
            raise ValueError("Amount must have cents precision")
        return amount.quantize(CENT)
    except (InvalidOperation, ValueError):
        raise ValueError("Invalid amount or cents precision") from None


def student_id(value: object) -> str:
    value = text(value)
    return value if re.fullmatch(r"10[0-9]{6}", value) else ""


def leading_student(value: object) -> str:
    match = re.match(r"^(10[0-9]{6})(?=\s|$)", text(value))
    return match[1] if match else ""


@dataclass
class SourceData:
    source: str
    label: str
    headers: list[str]
    rows: list[tuple[int, list[object]]]
    columns: dict[str, int]
    dispositions: dict[int, str] = field(default_factory=dict)

    def value(self, row: list[object], name: str) -> object:
        col = self.columns.get(name)
        return row[col] if col is not None and col < len(row) else None


def _columns(headers: list[object], source: str) -> dict[str, int] | None:
    keys = [header_key(value) for value in headers]
    columns = {}
    for name, choices in ALIASES.items():
        matches = [keys.index(header_key(choice)) for choice in choices if header_key(choice) in keys]
        if matches:
            # Two financial aliases in one source could mean pre-manipulated data.
            if name in REQUIRED[source] and len(set(matches)) > 1:
                raise ValueError(f"{source} has ambiguous columns for {choices[0]}.")
            columns[name] = matches[0]
    if not all(name in columns for name in REQUIRED[source]):
        return None
    populated = [key for key in keys if key]
    if len(populated) != len(set(populated)):
        raise ValueError(f"{source} has duplicate column headers.")
    return columns


def read_source(path: Path, source: str) -> SourceData:
    """Locate one matching source sheet, including exports with title rows."""
    if source not in REQUIRED:
        raise ValueError("Unknown reconciliation source.")
    candidates = []

    def inspect(name, values):
        iterator = iter(values)
        for row_number, row in enumerate(iterator, 1):
            if row_number > 30:
                break
            columns = _columns(list(row), source)
            if columns is not None:
                headers = [text(value) for value in row]
                candidates.append(SourceData(source, f"{path.name} / {name}", headers,
                    [(number, list(values)) for number, values in enumerate(iterator, row_number + 1)], columns))
                break

    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            inspect("CSV", csv.reader(handle))
    elif path.suffix.lower() in {".xlsx", ".xlsm"}:
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            for sheet in book:
                # Some Workday exports declare A1 as the used range despite
                # containing a full report. Read actual cells, not that metadata.
                sheet.reset_dimensions()
                inspect(sheet.title, sheet.iter_rows(values_only=True))
        finally:
            book.close()
    else:
        raise ValueError(f"{source} must be an XLSX, XLSM, or CSV export.")
    if not candidates:
        needed = ", ".join(ALIASES[name][0] for name in REQUIRED[source])
        raise ValueError(f"No {source} source sheet with the required headers: {needed}. Upload the original export.")
    if len(candidates) != 1:
        raise ValueError(f"Multiple {source} source sheets found. Upload a workbook with only one matching source sheet.")
    return candidates[0]


def banner_from_frame(frame: pd.DataFrame) -> SourceData:
    columns = _columns(list(frame.columns), "Banner")
    if columns is None:
        raise ValueError("Insights Banner extract is missing required columns.")
    rows = [[None if pd.isna(value) else value for value in row]
            for row in frame.itertuples(index=False, name=None)]
    return SourceData("Banner", "PROD Insights / Banner Activity by Date & Detail Code",
                      list(frame.columns), list(enumerate(rows, 2)), columns)


@dataclass(frozen=True)
class Transaction:
    source: str
    row: int
    day: date
    kind: str
    memo: str
    feed: str
    cwid: str
    account: str
    code: str
    debit: Decimal
    credit: Decimal
    journal: str = ""
    term: str = ""
    user: str = ""
    concerns: tuple[str, ...] = ()

    @property
    def group_key(self) -> tuple[int, str, str]:
        year, period, _ = fiscal_scope(self.day)
        return year, period, self.feed


@dataclass(frozen=True)
class Issue:
    source: str
    row: int
    issue: str


def prepare_source(
    source: SourceData, parameters: BannerActivityParameters,
) -> tuple[list[Transaction], list[Issue]]:
    transactions, issues = [], []
    identities = set()
    for number, values in source.rows:
        get = lambda name: source.value(values, name)
        if not any(text(value) for value in values):
            source.dispositions[number] = "Blank row"
            continue
        if any(text(value).casefold() in {"grand total", "report total", "subtotal"} for value in values[:5]):
            source.dispositions[number] = "Excluded report total"
            continue
        if source.source == "Workday":
            if not text(get("ledger")):
                source.dispositions[number] = "Excluded missing ledger account"
                issues.append(Issue(source.source, number, "Invalid ledger account; source coverage requires review"))
                continue
            if not re.match(r"^1305(?:\s|:|$)", text(get("ledger"))):
                source.dispositions[number] = "Excluded ledger account"
                continue
        else:
            code = text(get("code")).upper()
            if re.fullmatch(r"[A-Z0-9]{4}", code) is None:
                source.dispositions[number] = "Excluded invalid detail code"
                issues.append(Issue(source.source, number, "Invalid detail code; source coverage requires review"))
                continue
            if code not in parameters.detail_codes:
                source.dispositions[number] = "Excluded detail code"
                continue
        try:
            day = parse_date(get("date" if source.source == "Workday" else "feed_date"))
        except (ValueError, TypeError, OverflowError):
            source.dispositions[number] = "Excluded invalid date"
            issues.append(Issue(source.source, number, "Invalid source date; selection and amount coverage require review"))
            continue
        if not parameters.start_date <= day <= parameters.end_date:
            source.dispositions[number] = "Excluded date range"
            continue
        try:
            if source.source == "Workday":
                debit = money(get("debit"), blank_zero=True)
                credit = money(get("credit"), blank_zero=True)
            else:
                amount = money(get("amount"))
                debit, credit = max(amount, ZERO), max(-amount, ZERO)
        except ValueError:
            source.dispositions[number] = "Excluded invalid amount"
            issues.append(Issue(source.source, number, "Invalid amount; selected totals are incomplete"))
            continue
        concerns = []
        kind = text(get("kind")) if source.source == "Workday" else "Banner"
        memo = text(get("memo" if source.source == "Workday" else "description"))
        feed = reference(get("feed")) if source.source == "Banner" else (
            reference(get("reference")) or reference(get("memo")) if kind.upper() == "SIS" else ""
        )
        if source.source == "Workday":
            if not kind:
                concerns.append("Missing journal source; activity classification requires review")
            if debit < 0 or credit < 0 or (debit and credit):
                concerns.append("Unusual debit/credit signs or both populated")
            cwid = student_id(get("cwid")) or leading_student(memo)
            account = ""
            identity = (text(get("journal")), text(get("line"))) if text(get("line")) else None
        else:
            account = text(get("id"))
            cwid = student_id(account) or leading_student(memo)
            if student_id(account) and leading_student(memo) and student_id(account) != leading_student(memo):
                concerns.append("Conflicting account ID and description student ID")
                cwid = ""
            identity = (account, text(get("transaction"))) if text(get("transaction")) else None
        if identity is not None:
            if identity in identities:
                concerns.append("Duplicate source transaction identity; rows retained")
            identities.add(identity)
        if not feed and (source.source == "Banner" or kind.upper() == "SIS"):
            concerns.append("Missing feed document; cannot match this row")
        if not cwid and (source.source == "Banner" or kind.upper() != "SIS"):
            concerns.append("Missing valid CWID; student attribution incomplete")
        issues.extend(Issue(source.source, number, concern) for concern in concerns)
        source.dispositions[number] = "Selected"
        transactions.append(Transaction(source.source, number, day, kind, memo, feed, cwid, account,
            text(get("code")).upper() if source.source == "Banner" else "", debit, credit,
            text(get("journal")), text(get("term")), text(get("user")), tuple(concerns)))
    return transactions, issues
