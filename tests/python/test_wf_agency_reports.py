"""Synthetic named agency layouts, conservative totals and batch routing."""
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

from data_processing.wf_mines_balance_recon.agency_reports import (
    COLLECTION_HEADERS, COLLECTION_OUTPUT_HEADERS, RETURN_HEADERS, RETURN_OUTPUT_HEADERS,
    identify_report, read_agency_report,
)


def agency_file(root, creditor='12003', *, returned=False, fees='WITH FEES', rows=None, footer=True):
    kind = 'ReturnedAccounts' if returned else 'AccountStatusReport_NoBreakout'
    path = root / f'Report - {kind} - {creditor}-2099-10-01.xlsx'
    book = Workbook()
    sheet = book.active
    sheet.title = 'Report'
    sheet['A4'] = f'Creditor ignored for ID; accounts are {fees}'
    headers = RETURN_HEADERS if returned else COLLECTION_HEADERS
    for column, value in enumerate(headers, 1):
        sheet.cell(6, column, value)
    rows = rows if rows is not None else (
        [['0000123', 'Synthetic Student', 'CONSUMER-1', 100, 75, None, date(2099, 1, 1), date(2099, 9, 1), 'Test reason']]
        if returned else [['0000123', 'Synthetic Student', date(2099, 1, 1), 100, 75, 25, None]]
    )
    for number, values in enumerate(rows, 7):
        for column, value in enumerate(values, 1):
            sheet.cell(number, column, value)
    if footer:
        row = 7 + len(rows)
        sheet.cell(row, 1, len(rows))
        for column, header in enumerate(headers, 1):
            if header in {'Orig. Balance', 'Curr. Balance', 'Collected'}:
                sheet.cell(row, column, sum(values[column - 1] or 0 for values in rows))
    book.save(path)
    book.close()
    return path


class WFAgencyReportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_names_use_any_five_digit_creditor_and_validate_dates(self):
        path = Path('Report - ReturnedAccounts- 00999-2099-10-01.xlsx')
        report = identify_report(path)
        self.assertEqual(report.creditor, '00999')
        self.assertEqual(report.report_type, 'returned')
        self.assertEqual(report.report_date, date(2099, 10, 1))
        self.assertEqual(identify_report(Path('Report - AccountStatusReport_NoBreakout - 99999-2099-10-01.xlsx')).creditor, '99999')
        self.assertIsNone(identify_report(Path('Demographics.xlsx')))
        for name in ('Report - ReturnedAccounts- 1203-2099-10-01.xlsx',
                     'Report - ReturnedAccounts- 12003-2099-02-30.xlsx'):
            with self.subTest(name=name), self.assertRaises(ValueError):
                identify_report(Path(name))

    def test_collections_fee_types_metadata_and_unlabeled_immediate_totals(self):
        for creditor, fee_text, expected in (('12003', 'WITH FEES', 'Fees'), ('12561', 'no fees', 'No Fees')):
            path = agency_file(self.root, creditor, fees=fee_text)
            result = read_agency_report(path, identify_report(path))
            self.assertEqual(tuple(result.frame.columns), COLLECTION_OUTPUT_HEADERS)
            self.assertEqual(len(result.frame), 1)
            row = result.frame.iloc[0]
            self.assertEqual(row['Your Accnt #'], '0000123')
            self.assertEqual(row['Creditor'], creditor)
            self.assertEqual(row['Type'], expected)
            self.assertEqual(row['Source File'], path.name)
            self.assertIn('1 verified total row(s) excluded', result.sheets[0].notes)

    def test_returned_headers_dates_creditor_and_source_filename(self):
        path = agency_file(self.root, '00999', returned=True)
        result = read_agency_report(path, identify_report(path))
        self.assertEqual(tuple(result.frame.columns), RETURN_OUTPUT_HEADERS)
        self.assertEqual(len(result.frame), 1)
        self.assertEqual(result.frame.iloc[0]['Consumer #'], 'CONSUMER-1')
        self.assertEqual(result.frame.iloc[0]['Creditor'], '00999')
        self.assertIsNone(result.frame.iloc[0]['Last Pay Date'])
        self.assertNotIn('Type', result.frame)
        self.assertEqual(result.frame.iloc[0]['Source File'], path.name)

    def test_returned_total_accounts_labels_remove_entire_row(self):
        for label in ('Total Accounts:1', 'Total Accounts:5', 'Total Accounts: 2', ' total accounts: 15 '):
            with self.subTest(label=label):
                path = agency_file(self.root, returned=True)
                book = load_workbook(path)
                book.active['A8'] = label
                book.active['B8'] = 'Summary text'
                book.active['D8'] = 999  # Explicit labels do not require matching totals.
                book.save(path)
                book.close()
                result = read_agency_report(path, identify_report(path))
                self.assertEqual(len(result.frame), 1)
                self.assertEqual(result.frame.iloc[0]['Your Accnt #'], '0000123')
                self.assertIn('1 verified total row(s) excluded', result.sheets[0].notes)

    def test_account_after_labeled_returned_footer_still_stops(self):
        path = agency_file(self.root, returned=True)
        book = load_workbook(path)
        book.active['A8'] = 'Total Accounts:1'
        for column in range(1, 10):
            book.active.cell(9, column, book.active.cell(7, column).value)
        book.save(path)
        book.close()
        with self.assertRaisesRegex(ValueError, 'account data follows a summary row'):
            read_agency_report(path, identify_report(path))

    def test_missing_account_fields_and_unverified_totals_are_retained(self):
        rows = [['0000123', 'Synthetic Student', None, 100, 75, 25, None],
                [None, None, None, 100, 75, 25, None]]
        path = agency_file(self.root, rows=rows)
        book = load_workbook(path)
        book.active['E9'] = 999  # Totals do not reconcile: do not discard them.
        book.save(path)
        book.close()
        result = read_agency_report(path, identify_report(path))
        self.assertEqual(len(result.frame), 3)
        self.assertIn('8, 9', result.sheets[0].notes)

    def test_missing_or_ambiguous_fee_text_imports_as_unspecified(self):
        for fee_text in ('WITH FEES and NO FEES', 'not specified', ''):
            path = agency_file(self.root, fees=fee_text)
            if not fee_text:
                book = load_workbook(path)
                book.active['A4'] = None
                book.save(path)
                book.close()
            result = read_agency_report(path, identify_report(path))
            self.assertEqual(len(result.frame), 1)
            self.assertEqual(result.frame.iloc[0]['Type'], 'unspecified')
            self.assertIn('Fee classification unspecified', result.sheets[0].notes)

    def test_singular_fee_labels_preserve_classification(self):
        for text, expected in (('WITH FEE', 'Fees'), ('NO FEE', 'No Fees')):
            path = agency_file(self.root, fees=text)
            result = read_agency_report(path, identify_report(path))
            self.assertEqual(result.frame.iloc[0]['Type'], expected)

    def test_schema_changes_still_fail_without_guessing(self):
        path = agency_file(self.root)
        book = load_workbook(path)
        book.active['B6'] = 'Changed name heading'
        book.save(path)
        book.close()
        with self.assertRaisesRegex(ValueError, 'row 6'):
            read_agency_report(path, identify_report(path))

    def test_numeric_identifiers_preserve_zero_formats_and_signed_amounts(self):
        path = agency_file(self.root, rows=[[123, 'Synthetic', None, -20, -15, -5, None]])
        book = load_workbook(path)
        book.active['A7'].number_format = '0000000'
        book.save(path)
        book.close()
        result = read_agency_report(path, identify_report(path))
        self.assertEqual(result.frame.iloc[0]['Your Accnt #'], '0000123')
        self.assertEqual(result.frame.iloc[0]['Orig. Balance'], -20)
        self.assertEqual(len(result.frame), 1)

    def test_aggregate_formulas_can_identify_footer_without_a_blank_separator(self):
        path = agency_file(self.root)
        book = load_workbook(path)
        book.active['A8'] = None
        for column in ('D', 'E', 'F'):
            book.active[f'{column}8'] = f'=SUM({column}7:{column}7)'
        book.save(path)
        book.close()
        result = read_agency_report(path, identify_report(path))
        self.assertEqual(len(result.frame), 1)


if __name__ == '__main__':
    unittest.main()
