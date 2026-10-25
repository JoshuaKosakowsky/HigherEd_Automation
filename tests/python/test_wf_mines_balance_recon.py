"""Synthetic end-to-end W&F source collection and safe publication checks."""
import csv
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import MagicMock, patch

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from app.gui.models import WorkflowContext, WorkflowMode
from app.gui.services.wf_mines_balance_recon import run_wf_mines_balance_recon
from data_processing.wf_mines_balance_recon.workbook import (
    ACTIVITY_HEADERS, HOLD_HEADERS, create_review_workbook, read_insights_file,
)
from shared.cancellation import CancellationToken, WorkflowCancelled


def holds_frame():
    return pd.DataFrame([['TEST-1', 'Synthetic', 'Student', -5, '2099-10-01',
                          'CO', '2099-01-01', '9999-12-31']], columns=HOLD_HEADERS)


def activity_frame():
    return pd.DataFrame([['TEST-1', 'Synthetic, Student', 'WOFF', 'Synthetic reversal',
                          -25, 0, '209910', None, 'TEST-FEED', None,
                          'Synthetic write-off', 'TEST_USER']], columns=ACTIVITY_HEADERS)


def write_frame(path, frame):
    book = Workbook()
    sheet = book.active
    sheet.append(list(frame.columns))
    for row in frame.itertuples(index=False, name=None):
        sheet.append(list(row))
    book.save(path)
    book.close()


class WFMinesBalanceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.output = self.root / 'review.xlsx'
        self.holds = self.root / 'holds.xlsx'
        self.activity = self.root / 'activity.xlsx'
        write_frame(self.holds, holds_frame())
        write_frame(self.activity, activity_frame())
        self.agency = self.root / 'agency.xlsx'
        book = Workbook()
        sheet = book.active
        sheet.title = 'Account Data'
        sheet.append(['Account', 'Amount'])
        sheet.append(['0000123', -12.34])
        sheet['A2'].number_format = '@'
        sheet['B2'].number_format = '$#,##0.00'
        sheet['A1'].font = Font(bold=True)
        sheet.column_dimensions['A'].width = 24
        sheet.freeze_panes = 'A2'
        sheet['A4'] = '=literal source text'
        sheet['A4'].data_type = 's'
        sheet.merge_cells('A5:B5')
        sheet['A5'] = 'Original note'
        extra = book.create_sheet('Hidden Data')
        extra['A1'] = 'Hidden source row'
        extra.sheet_state = 'hidden'
        book.create_sheet('Empty')
        book.save(self.agency)
        book.close()
        self.params = dict(banner_source='manual', holds_file=self.holds,
                           activity_file=self.activity, wf_files=(self.agency,), output_file=self.output)

    def context(self, mode=None, cancellation=None):
        return WorkflowContext('wf_mines_balance_recon', self.params, mode, cancellation=cancellation)

    def test_manual_combines_every_worksheet_preserves_values_formats_and_sources(self):
        originals = {path: path.read_bytes() for path in (self.holds, self.activity, self.agency)}
        with patch('app.gui.services.wf_mines_balance_recon.build_authenticated_client') as auth:
            result = run_wf_mines_balance_recon(self.context())
        auth.assert_not_called()
        self.assertTrue(result.success)
        book = load_workbook(self.output)
        try:
            self.assertEqual(len(book.sheetnames), 6)
            self.assertEqual(book.sheetnames[:2], ['Active CO Holds', 'Banner Detail Codes'])
            self.assertEqual(book['Active CO Holds']['H2'].value, datetime(9999, 12, 31))
            self.assertEqual(book['Banner Detail Codes']['E2'].value, -25)
            self.assertIsNone(book['Banner Detail Codes']['J2'].value)
            agency = book[book.sheetnames[2]]
            self.assertEqual(agency['A2'].value, '0000123')
            self.assertEqual(agency['B2'].number_format, '$#,##0.00')
            self.assertEqual(agency['A4'].value, '=literal source text')
            self.assertEqual(agency['A4'].data_type, 's')
            self.assertEqual(agency.column_dimensions['A'].width, 24)
            self.assertTrue(agency['A1'].font.bold)
            self.assertIn('A5:B5', str(agency.merged_cells))
            self.assertEqual(agency.freeze_panes, 'A2')
            self.assertTrue(all(sheet.sheet_state == 'visible' for sheet in book))
            sources = list(book['Sources'].values)
            self.assertEqual(sources[-1][3], 'Empty')
            self.assertIn('coverage unconfirmed', sources[1][2])
            self.assertFalse(any(sheet.tables for sheet in book))
        finally:
            book.close()
        self.assertTrue(all(path.read_bytes() == content for path, content in originals.items()))

    def test_batch_same_names_long_names_and_sheet_name_collisions_preserve_all_files(self):
        files = []
        for directory in ('one', 'two'):
            folder = self.root / directory
            folder.mkdir()
            path = folder / ('[agency]_' + 'long' * 10 + '.xlsx')
            write_frame(path, pd.DataFrame({'ID': [directory]}))
            files.append(path)
        self.params['wf_files'] = tuple(files)
        run_wf_mines_balance_recon(self.context())
        book = load_workbook(self.output)
        try:
            self.assertEqual(len(book.sheetnames), 5)
            self.assertEqual(len({name.casefold() for name in book.sheetnames}), 5)
            self.assertTrue(all(len(name) <= 31 and '[' not in name for name in book.sheetnames))
            self.assertEqual([book[name]['A2'].value for name in book.sheetnames[2:4]], ['one', 'two'])
        finally:
            book.close()

    def test_mixed_twelve_file_batch_consolidates_known_reports_and_retains_other_sheets(self):
        from test_wf_agency_reports import agency_file
        files = [agency_file(self.root, str(12003 + index),
                             fees='WITH FEES' if index % 2 else 'NO FEES') for index in range(8)]
        files.extend(agency_file(self.root, str(12559 + index), returned=True) for index in range(3))
        files.append(self.agency)
        self.params['wf_files'] = tuple(files)
        result = run_wf_mines_balance_recon(self.context())
        self.assertTrue(result.success)
        book = load_workbook(self.output)
        try:
            self.assertEqual(book.sheetnames[:4], [
                'Active CO Holds', 'Banner Detail Codes', 'W&F Collections', 'W&F Returned',
            ])
            self.assertEqual(book['W&F Collections'].max_row, 9)
            self.assertEqual(book['W&F Returned'].max_row, 4)
            collections = list(book['W&F Collections'].values)
            self.assertEqual(collections[0][-3:], ('Creditor', 'Type', 'Source File'))
            self.assertEqual({row[-2] for row in collections[1:]}, {'Fees', 'No Fees'})
            self.assertEqual({row[-3] for row in collections[1:]}, {str(12003 + index) for index in range(8)})
            self.assertTrue(all(row[0] == '0000123' for row in collections[1:]))
            self.assertTrue(all(row[-1] in {path.name for path in files} for row in collections[1:]))
            self.assertEqual(len(book.sheetnames), 8)
            self.assertTrue(any('Unclassified report' in str(row[-1]) for row in list(book['Sources'].values)[1:]))
        finally:
            book.close()

    def test_matching_sheets_from_one_report_are_both_imported_and_mapped(self):
        from test_wf_agency_reports import agency_file
        path = agency_file(self.root)
        source = load_workbook(path)
        source.copy_worksheet(source.active).title = 'Second report'
        source.save(path)
        source.close()
        self.params['wf_files'] = (path,)
        run_wf_mines_balance_recon(self.context())
        book = load_workbook(self.output)
        try:
            self.assertEqual(book['W&F Collections'].max_row, 3)
            manifest = [row for row in list(book['Sources'].values)[1:] if row[0] == 'W&F Collections']
            self.assertEqual({row[3] for row in manifest}, {'Report', 'Second report'})
        finally:
            book.close()

    def test_invalid_missing_duplicate_and_unsupported_agency_inputs_cannot_publish(self):
        invalid = self.root / 'unsupported.pdf'
        invalid.write_bytes(b'synthetic unsupported input')
        for files in ((), (self.agency, self.agency), (self.root / 'missing.xlsx',), (invalid,)):
            with self.subTest(files=files):
                self.params['wf_files'] = files
                with self.assertRaises(ValueError):
                    run_wf_mines_balance_recon(self.context())
                self.assertFalse(self.output.exists())

    def test_source_or_existing_output_cannot_be_overwritten(self):
        self.output.write_bytes(b'existing output')
        with self.assertRaises(ValueError):
            run_wf_mines_balance_recon(self.context())
        self.assertEqual(self.output.read_bytes(), b'existing output')
        self.params['output_file'] = self.agency
        before = self.agency.read_bytes()
        with self.assertRaises(ValueError):
            run_wf_mines_balance_recon(self.context())
        self.assertEqual(self.agency.read_bytes(), before)

    def test_formula_without_cached_result_stops_instead_of_losing_financial_data(self):
        book = load_workbook(self.agency)
        book.active['B2'] = '=10-5'
        book.save(self.agency)
        book.close()
        with self.assertRaisesRegex(ValueError, 'formula without a saved value'):
            run_wf_mines_balance_recon(self.context())
        self.assertFalse(self.output.exists())

    def test_empty_insights_csv_headers_and_year_9999_are_supported(self):
        path = self.root / 'holds.csv'
        with path.open('w', newline='', encoding='utf-8-sig') as handle:
            writer = csv.writer(handle)
            writer.writerow(['Report title'])
            writer.writerow(HOLD_HEADERS)
            writer.writerows(holds_frame().itertuples(index=False, name=None))
        frame = read_insights_file(path, 'holds')
        self.assertEqual(frame.iloc[0]['To Date'].year, 9999)
        write_frame(self.holds, pd.DataFrame(columns=HOLD_HEADERS))
        write_frame(self.activity, pd.DataFrame(columns=ACTIVITY_HEADERS))
        run_wf_mines_balance_recon(self.context())
        book = load_workbook(self.output)
        try:
            self.assertEqual(book['Active CO Holds'].max_row, 1)
            self.assertEqual(book['Banner Detail Codes'].max_row, 1)
        finally:
            book.close()

    def test_wrong_or_ambiguous_insights_uploads_fail_without_authentication(self):
        self.params['holds_file'] = self.activity
        with self.assertRaisesRegex(ValueError, 'exactly one matching holds'):
            run_wf_mines_balance_recon(self.context())
        self.params['holds_file'] = self.holds
        book = load_workbook(self.holds)
        book.copy_worksheet(book.active)
        book.save(self.holds)
        book.close()
        with self.assertRaisesRegex(ValueError, 'found 2'):
            run_wf_mines_balance_recon(self.context())
        self.assertFalse(self.output.exists())

    def test_cancel_and_failed_publication_leave_no_partial_output(self):
        token = CancellationToken()
        token.request()
        with self.assertRaises(WorkflowCancelled):
            run_wf_mines_balance_recon(self.context(cancellation=token))
        with patch('data_processing.shared.workbook_output.shutil') as transfer:
            transfer.copyfileobj.side_effect = OSError('synthetic')
            with self.assertRaises(OSError):
                run_wf_mines_balance_recon(self.context())
        self.assertFalse(self.output.exists())

    def test_cancellation_during_save_and_racing_destination_preserve_safety(self):
        original_save = Workbook.save
        token = CancellationToken()
        def cancel_after_save(book, filename):
            original_save(book, filename)
            token.request()
        with patch.object(Workbook, 'save', cancel_after_save):
            with self.assertRaises(WorkflowCancelled):
                run_wf_mines_balance_recon(self.context(cancellation=token))
        self.assertFalse(self.output.exists())
        def race_after_save(book, filename):
            original_save(book, filename)
            self.output.write_bytes(b'other process output')
        with patch.object(Workbook, 'save', race_after_save):
            with self.assertRaises(FileExistsError):
                run_wf_mines_balance_recon(self.context())
        self.assertEqual(self.output.read_bytes(), b'other process output')

    def test_cached_formula_is_copied_as_its_saved_value(self):
        from zipfile import ZipFile
        book = load_workbook(self.agency)
        book.active['B2'] = '=10-5'
        book.save(self.agency)
        book.close()
        # Excel normally supplies this cached value; synthetic XML avoids needing Excel.
        with ZipFile(self.agency) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        entries['xl/worksheets/sheet1.xml'] = entries['xl/worksheets/sheet1.xml'].replace(
            b'<f>10-5</f><v></v>', b'<f>10-5</f><v>5</v>',
        ).replace(
            b'<f>10-5</f><v />', b'<f>10-5</f><v>5</v>',
        )
        with ZipFile(self.agency, 'w') as archive:
            for name, content in entries.items():
                archive.writestr(name, content)
        run_wf_mines_balance_recon(self.context())
        book = load_workbook(self.output)
        try:
            cell = book[book.sheetnames[2]]['B2']
            self.assertEqual(cell.value, 5)
            self.assertEqual(cell.data_type, 'n')
        finally:
            book.close()

    def test_auth_errors_are_sanitized_and_manual_mode_cannot_request_prod(self):
        with self.assertRaisesRegex(ValueError, 'does not use TEST or PROD'):
            run_wf_mines_balance_recon(self.context(WorkflowMode.PRODUCTION))
        self.params['banner_source'] = 'sql'
        with patch('app.gui.services.wf_mines_balance_recon.load_access_configuration') as policy, patch(
            'app.gui.services.wf_mines_balance_recon.load_department_profiles',
        ), patch('app.gui.services.wf_mines_balance_recon.build_authenticated_client',
                 side_effect=RuntimeError('https://example.invalid?token=SECRET')):
            policy.return_value.is_administrator.return_value = True
            with self.assertRaisesRegex(ValueError, 'Insights extraction failed') as error:
                run_wf_mines_balance_recon(self.context(WorkflowMode.PRODUCTION))
        self.assertNotIn('SECRET', str(error.exception))
        self.assertTrue(error.exception.__suppress_context__)
        self.assertFalse(self.output.exists())

    def test_admin_prod_queries_both_complete_reports_without_date_filter(self):
        self.params.update(banner_source='sql', detail_codes='WOFF, BDRC')
        client = MagicMock()
        client.__enter__.return_value = client
        holds, activity = holds_frame(), activity_frame()
        holds['Report Row Count'] = 1
        activity['Report Row Count'] = 1
        client.run_sql_file.return_value = holds
        client.run_sql.return_value = activity
        with patch('app.gui.services.wf_mines_balance_recon.load_access_configuration') as policy, patch(
            'app.gui.services.wf_mines_balance_recon.load_department_profiles',
        ), patch('app.gui.services.wf_mines_balance_recon.build_authenticated_client', return_value=(client, 'cached')) as auth:
            policy.return_value.is_administrator.return_value = True
            result = run_wf_mines_balance_recon(self.context(WorkflowMode.PRODUCTION))
        self.assertTrue(result.success)
        self.assertTrue(auth.call_args.kwargs['use_saved_mines_login'])
        self.assertEqual(client.run_sql_file.call_args.args[0].name, 'collections_tracking.sql')
        sql = client.run_sql.call_args.args[0]
        self.assertIn("IN ('WOFF', 'BDRC')", sql)
        self.assertNotIn('CAST(', sql)
        book = load_workbook(self.output)
        try:
            self.assertEqual(book['Active CO Holds'].max_column, 8)
            self.assertEqual(book['Banner Detail Codes'].max_column, 12)
        finally:
            book.close()

    def test_forged_admin_wrong_mode_invalid_codes_and_capped_results_fail_closed(self):
        self.params['banner_source'] = 'sql'
        for admin, mode, codes, count in (
            (False, WorkflowMode.PRODUCTION, 'WOFF', 1),
            (True, WorkflowMode.TEST, 'WOFF', 1),
            (True, WorkflowMode.PRODUCTION, '', 1),
            (True, WorkflowMode.PRODUCTION, 'WOFF', 2),
        ):
            with self.subTest(admin=admin, mode=mode, codes=codes, count=count):
                self.params['detail_codes'] = codes
                frame = holds_frame()
                frame['Report Row Count'] = count
                client = MagicMock()
                client.__enter__.return_value = client
                client.run_sql_file.return_value = frame
                with patch('app.gui.services.wf_mines_balance_recon.load_access_configuration') as policy, patch(
                    'app.gui.services.wf_mines_balance_recon.load_department_profiles',
                ), patch('app.gui.services.wf_mines_balance_recon.build_authenticated_client', return_value=(client, 'cached')) as auth:
                    policy.return_value.is_administrator.return_value = admin
                    with self.assertRaises(ValueError):
                        run_wf_mines_balance_recon(self.context(mode))
                if count == 1:
                    auth.assert_not_called()
                self.assertFalse(self.output.exists())


if __name__ == '__main__':
    unittest.main()
