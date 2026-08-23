from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

from data_processing.population_testing.assignment import ADDITIONAL_ACCOUNTS
from data_processing.population_testing.export import (
    SOURCE_OUTPUT_COLUMNS,
    export_population_workbook,
)
from data_processing.shared.xlsx_output_format import (
    COUNT_FORMAT,
    CURRENCY_FORMAT,
)
from data_processing.trends.config import REQUIRED_COLUMNS, TrendsConfig
from data_processing.trends.pipeline import run_trends_pipeline


def worksheet_headers(worksheet) -> list[object]:
    return [cell.value for cell in worksheet[1]]


def assert_standard_header(
    test_case: unittest.TestCase,
    worksheet,
) -> None:
    first_header = worksheet["A1"]
    test_case.assertTrue(first_header.font.bold)
    test_case.assertEqual(first_header.font.color.type, "rgb")
    test_case.assertTrue(first_header.font.color.rgb.endswith("FFFFFF"))
    test_case.assertEqual(first_header.fill.fill_type, "solid")
    test_case.assertTrue(first_header.fill.fgColor.rgb.endswith("1F4E78"))


class PopulationWorkbookContractTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary_directory.cleanup)
        self.root = Path(self.temporary_directory.name)

    def test_population_workbook_preserves_required_output_contract(self) -> None:
        rows = [
            self.population_row(
                cwid="900000001",
                bucket_id="B001",
                bucket_sheet="B001 RL UG 6-19",
                bucket_name="Resident, Undergraduate, 6-19 credits",
                selected=True,
                assigned_to="Analyst One",
                environment="TEST",
                graduate_online=False,
            ),
            self.population_row(
                cwid="900000002",
                bucket_id="B001",
                bucket_sheet="B001 RL UG 6-19",
                bucket_name="Resident, Undergraduate, 6-19 credits",
                selected=False,
                assigned_to="",
                environment="",
                graduate_online=False,
            ),
            self.population_row(
                cwid="900000003",
                bucket_id="B000",
                bucket_sheet="B000 Graduate Online",
                bucket_name="Graduate Online",
                selected=True,
                assigned_to=ADDITIONAL_ACCOUNTS,
                environment="TEST",
                graduate_online=True,
            ),
        ]
        population = pd.DataFrame(rows)
        validation = pd.DataFrame(
            [
                {
                    "Check": "Synthetic reconciliation",
                    "Status": "PASS",
                    "Detail": "All synthetic rows reconciled.",
                }
            ]
        )
        output = self.root / "reports" / "Population Testing.xlsx"

        returned_path = export_population_workbook(
            population,
            validation_df=validation,
            output_file=output,
            staff_names=("Analyst One", ADDITIONAL_ACCOUNTS),
        )

        self.assertEqual(returned_path, output)
        self.assertTrue(output.is_file())

        workbook = load_workbook(output, data_only=False)
        self.addCleanup(workbook.close)

        self.assertEqual(
            workbook.sheetnames,
            [
                "Summary",
                "Validation",
                "Testing Sample",
                "B000 Graduate Online",
                "B001 RL UG 6-19",
                "A01 Analyst One",
                "A02 Additional Accounts",
            ],
        )

        summary = workbook["Summary"]
        self.assertEqual(
            worksheet_headers(summary),
            [
                "Bucket ID",
                "Worksheet",
                "Bucket Description",
                "Population Count",
                "Selected for Testing",
                "Selection Rate",
                "Assigned Count",
                "TEST Count",
                "PROD Count",
                "Test All",
            ],
        )
        self.assertEqual(summary["A2"].value, "B000")
        self.assertEqual(summary["D2"].value, 1)
        self.assertEqual(summary["E2"].value, 1)
        self.assertEqual(summary["F2"].value, 1)
        self.assertEqual(summary["F2"].number_format, "0.0%")
        self.assertEqual(summary["A3"].value, "B001")
        self.assertEqual(summary["D3"].value, 2)
        self.assertEqual(summary["E3"].value, 1)
        self.assertEqual(summary["F3"].value, 0.5)
        self.assertEqual(summary.freeze_panes, "A2")
        self.assertEqual(summary.auto_filter.ref, summary.dimensions)
        assert_standard_header(self, summary)

        self.assertEqual(
            worksheet_headers(workbook["Validation"]),
            ["Check", "Status", "Detail"],
        )
        self.assertEqual(
            worksheet_headers(workbook["Testing Sample"]),
            SOURCE_OUTPUT_COLUMNS + ["Assigned Staff", "Environment"],
        )
        self.assertEqual(
            worksheet_headers(workbook["B001 RL UG 6-19"]),
            SOURCE_OUTPUT_COLUMNS,
        )
        self.assertEqual(
            worksheet_headers(workbook["A01 Analyst One"]),
            SOURCE_OUTPUT_COLUMNS + ["Environment", "Comments"],
        )
        self.assertEqual(
            worksheet_headers(workbook["A02 Additional Accounts"]),
            SOURCE_OUTPUT_COLUMNS + ["Environment"],
        )
        self.assertEqual(workbook["Testing Sample"]["D2"].value, "900000001")

    @staticmethod
    def population_row(
        *,
        cwid: str,
        bucket_id: str,
        bucket_sheet: str,
        bucket_name: str,
        selected: bool,
        assigned_to: str,
        environment: str,
        graduate_online: bool,
    ) -> dict[str, object]:
        return {
            "Term": "202680",
            "Student Last Name": "Synthetic",
            "Student First Name": cwid[-1],
            "CWID": cwid,
            "Primary Student Level Desc": (
                "Graduate" if graduate_online else "Undergraduate"
            ),
            "Primary Program": "X-ONLINE" if graduate_online else "BS-CS",
            "Primary 1st Major Desc": "Synthetic Major",
            "Registered Credits": 9,
            "Student Residency Desc": "Resident",
            "BUCKET_ID": bucket_id,
            "BUCKET_SHEET": bucket_sheet,
            "BUCKET_NAME": bucket_name,
            "SELECTED_FOR_TESTING": selected,
            "ASSIGNED_TO": assigned_to,
            "TESTING_ENVIRONMENT": environment,
            "IS_GRAD_ONLINE": graduate_online,
        }


class TrendsWorkbookContractTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary_directory.cleanup)
        self.root = Path(self.temporary_directory.name)

    def test_trends_pipeline_preserves_required_workbook_contract(self) -> None:
        input_directory = self.root / "input"
        input_directory.mkdir()
        source = input_directory / "TGIACCD_fy26.xlsx"
        self.write_tgiaccd_source(source)

        detail_codes_path = self.root / "detail_codes.json"
        detail_codes_path.write_text(
            json.dumps(self.detail_codes()),
            encoding="utf-8",
        )
        output = self.root / "reports" / "Trends.xlsx"

        returned_path = run_trends_pipeline(
            TrendsConfig(
                input_dir=input_directory,
                detail_codes_file=detail_codes_path,
                output_file=output,
            )
        )

        self.assertEqual(returned_path, output)
        self.assertTrue(output.is_file())

        workbook = load_workbook(output, data_only=False)
        self.addCleanup(workbook.close)

        self.assertEqual(
            workbook.sheetnames,
            [
                "Summary",
                "TUI Averages",
                "Reporting Group Averages",
                "Reporting Group Totals",
                "Reporting Group Totals by Term",
                "Collections by Term",
                "Category Averages",
                "Detail Code Totals",
                "Checks",
                "Definitions",
            ],
        )

        summary = workbook["Summary"]
        self.assertEqual(
            worksheet_headers(summary),
            [
                "Fiscal Year",
                "Source File",
                "Rows",
                "Students in File",
                "Students with Positive AR",
                "Gross AR",
                "Credit Balances",
                "Net AR",
                "Average Positive Balance",
                "Average Net Balance - All Students",
                "Charge Activity",
                "Payment/Credit Activity",
                "Net AR Activity",
                "Gross Tuition Charges",
                "Tuition Offsets",
                "Net Tuition Revenue",
                "WOFF Transactions",
                "WOFF Students",
                "WOFF Amount",
            ],
        )
        self.assertEqual(summary["A2"].value, 2026)
        self.assertEqual(summary["B2"].value, "TGIACCD_fy26.xlsx")
        self.assertEqual(summary["C2"].value, 5)
        self.assertEqual(summary["N2"].value, 1000)
        self.assertEqual(summary["O2"].value, 200)
        self.assertEqual(summary["P2"].value, 800)
        self.assertEqual(summary["Q2"].value, 1)
        self.assertEqual(summary["S2"].value, 50)
        self.assertEqual(summary["C2"].number_format, COUNT_FORMAT)
        self.assertEqual(summary["F2"].number_format, CURRENCY_FORMAT)
        self.assertEqual(summary.freeze_panes, "A2")
        self.assertFalse(summary.sheet_view.showGridLines)
        self.assertIn("FiscalYearSummaryTable", summary.tables)
        assert_standard_header(self, summary)

        checks = workbook["Checks"]
        self.assertEqual(
            worksheet_headers(checks),
            [
                "Fiscal Year",
                "Source File",
                "Status",
                "Rows",
                "Mapped Rows",
                "Unmapped Rows",
                "Unmapped Detail Codes",
            ],
        )
        self.assertEqual(checks["C2"].value, "PASS")
        self.assertEqual(checks["D2"].value, 5)
        self.assertEqual(checks["E2"].value, 5)
        self.assertEqual(checks["F2"].value, 0)

        totals = workbook["Reporting Group Totals"]
        rows_by_label = {
            totals.cell(row=row_number, column=2).value: row_number
            for row_number in range(2, totals.max_row + 1)
        }
        writeoff_row = rows_by_label["Write-Off Bad Debt"]
        recovery_row = rows_by_label["BDRC - Bad Debt Recovery"]
        self.assertEqual(totals.cell(writeoff_row, 11).value, -50)
        self.assertEqual(totals.cell(recovery_row, 11).value, 30)
        self.assertEqual(totals.cell(writeoff_row, 11).number_format, CURRENCY_FORMAT)

        collections = workbook["Collections by Term"]
        collection_rows = {
            collections.cell(row=row_number, column=3).value: row_number
            for row_number in range(2, collections.max_row + 1)
        }
        coll_row = collection_rows["COLL"]
        self.assertEqual(collections.cell(coll_row, 7).value, 1)
        self.assertEqual(collections.cell(coll_row, 9).value, 70)
        self.assertEqual(collections.cell(coll_row, 10).value, 70)

        self.assertTrue(workbook.calculation.fullCalcOnLoad)
        self.assertTrue(workbook.calculation.forceFullCalc)

    @staticmethod
    def write_tgiaccd_source(path: Path) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(list(REQUIRED_COLUMNS))
        for row in [
            ["900000001", "TUIT", "Tuition charge", 1000, 800, "202610"],
            ["900000002", "TUIR", "Tuition credit", 200, -200, "202610"],
            ["900000003", "WOFF", "Write-off", 50, 0, "202610"],
            ["900000004", "BDRC", "Bad debt recovery", 30, 0, "202620"],
            ["900000005", "COLL", "Collections", 70, 70, "202620"],
        ]:
            worksheet.append(row)
        workbook.save(path)
        workbook.close()

    @staticmethod
    def detail_codes() -> dict[str, dict[str, str]]:
        return {
            "TUIT": {
                "type": "C",
                "category": "TUI",
                "detail_code_description": "Tuition charge",
            },
            "TUIR": {
                "type": "P",
                "category": "TUI",
                "detail_code_description": "Tuition credit",
            },
            "WOFF": {
                "type": "P",
                "category": "WOF",
                "detail_code_description": "Write-off",
            },
            "BDRC": {
                "type": "C",
                "category": "BDR",
                "detail_code_description": "Bad debt recovery",
            },
            "COLL": {
                "type": "C",
                "category": "COL",
                "detail_code_description": "Collections",
            },
            "PCCS": {
                "type": "C",
                "category": "COL",
                "detail_code_description": "Collection service",
            },
        }


if __name__ == "__main__":
    unittest.main()
