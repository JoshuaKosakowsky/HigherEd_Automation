from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook, load_workbook

from app.gui.models import WorkflowContext, WorkflowDefinition, WorkflowMode, WorkflowResult
from app.gui.services.access import AccessConfiguration, filter_workflows_for_view
from app.gui.services.textbook_recon import _normalized_banner, _render_sql, run_recon
from app.gui.workflow_registry import get_workflow, get_workflows
from data_processing.textbook_brokers.recon import (
    BLUE_HEADER, BLUE_TAB_DARK, BLUE_TAB_LIGHT, GREEN_HEADER,
    GREEN_TAB_DARK, GREEN_TAB_LIGHT, _banner_rows, _broker_rows,
    _write_recon, build_recon,
)
from shared.progress import ProgressReporter


class TextbookReconTests(unittest.TestCase):
    def setUp(self) -> None:
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)

    def _brokers(self) -> Path:
        path = self.root / "brokers.xlsx"
        book = Workbook()
        ia = book.active
        ia.title = "IA Charge Report"
        ia.append(["Customer ID", "First Name", "Last Name", "Sub Total", "Tax", "Total"])
        ia.append(["90000001", "Ada", "Able", 12, 0, 12])
        ia.append(["90000001", "ADA", "ABLE", -2, 0, -2])
        ia.append(["90000002", "Bo", "Baker", 5, 0, 5])
        ia.append([None, None, None, 15, 0, 15])  # Vendor grand total.
        fa = book.create_sheet("FA Charge Report")
        fa.append(["Student ID", "First", "Last", "Subtotal", "Tax", "Spent"])
        fa.append(["90000003", "Cy", "Cedar", 7, 0, 7])
        book.save(path)
        return path

    def _banner(self, code: str, rows: list[tuple[str, str, int]]) -> Path:
        path = self.root / f"{code}.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.append(["'ID'", "'Name'", "'Detail Code'", "'Amount'", "'Feed Date'"])
        for cwid, name, amount in rows:
            sheet.append([cwid, name, code, amount, "2026-07-15"])
        book.save(path)
        return path

    def test_net_recon_aligns_missing_students(self) -> None:
        _, broker = _broker_rows(self._brokers(), "IA")
        _, _, banner = _banner_rows(self._banner("FRST", [
            ("90000001", "Able, Ada", 15),
            ("90000001", "Able, Ada", -5),
            ("90000004", "Dove, Dee", 3),
        ]), "FRST", "2026-07")
        sheet = Workbook().active
        _write_recon(sheet, broker, banner)
        self.assertEqual((sheet["A4"].value, sheet["D4"].value,
                          sheet["G4"].value, sheet["I4"].value),
                         ("90000001", 10, "90000001", 10))
        self.assertIsNone(sheet["G5"].value)
        self.assertIsNone(sheet["A6"].value)
        self.assertEqual(sheet["G6"].value, "90000004")
        self.assertIn("ROUND", sheet["M4"].value)

    def test_python_creates_complete_workbook_with_native_pivots(self) -> None:
        output = self.root / "recon.xlsx"
        updates = []
        counts = build_recon(
            brokers=self._brokers(),
            frst=self._banner("FRST", [("90000001", "Able, Ada", 15),
                                       ("90000001", "Able, Ada", -5),
                                       ("90000004", "Dove, Dee", 3)]),
            book=self._banner("BOOK", [("90000003", "Cedar, Cy", 7)]),
            output=output, month="2026-07",
            progress_reporter=ProgressReporter(updates.append),
        )
        stages = [update.stage for update in updates]
        self.assertIn("Reading IA Textbook Brokers charges", stages)
        self.assertIn("Reading BOOK Banner transactions", stages)
        self.assertEqual(stages[-1], "Saving reconciliation workbook")
        self.assertEqual(counts, {"IA": 3, "FA": 1})
        result = load_workbook(output)
        self.assertEqual(result.sheetnames, [
            "IA Recon", "IA Pivot", "IA Charge Report", "FRST Pivot", "FRST",
            "BOOK Recon", "FA Pivot", "FA Charge Report", "BOOK Pivot", "BOOK",
        ])
        self.assertEqual([result["IA Recon"].cell(4, 1).value,
                          result["IA Recon"].cell(4, 4).value], ["90000001", 10])
        self.assertEqual(result["IA Recon"].cell(4, 9).value, 10)
        self.assertEqual(result["IA Recon"].cell(5, 7).value, None)
        self.assertEqual(result["IA Recon"].cell(6, 1).value, None)
        self.assertEqual(result["IA Recon"].cell(6, 7).value, "90000004")
        self.assertEqual(result["BOOK Recon"].cell(4, 4).value, 7)
        self.assertEqual(result["BOOK Recon"].cell(4, 9).value, 7)
        for name, expected in {
            "IA Recon": BLUE_TAB_DARK, "IA Pivot": BLUE_TAB_LIGHT,
            "IA Charge Report": BLUE_TAB_LIGHT, "FRST Pivot": BLUE_TAB_LIGHT,
            "FRST": BLUE_TAB_LIGHT, "BOOK Recon": GREEN_TAB_DARK,
            "FA Pivot": GREEN_TAB_LIGHT, "FA Charge Report": GREEN_TAB_LIGHT,
            "BOOK Pivot": GREEN_TAB_LIGHT, "BOOK": GREEN_TAB_LIGHT,
        }.items():
            self.assertEqual(result[name].sheet_properties.tabColor.rgb[-6:], expected)
        for recon, title_color in (("IA Recon", BLUE_HEADER), ("BOOK Recon", GREEN_HEADER)):
            sheet = result[recon]
            self.assertFalse(sheet.merged_cells.ranges)
            for row_range in ("A1:D1", "G1:I1", "K1:N1"):
                for row in sheet[row_range]:
                    for cell in row:
                        self.assertEqual(cell.alignment.horizontal, "centerContinuous")
            self.assertEqual(sheet["A1"].fill.fgColor.rgb[-6:], title_color)
        for name in ("IA Pivot", "FA Pivot", "FRST Pivot", "BOOK Pivot"):
            pivot = result[name]._pivots[0]
            self.assertTrue(pivot.cache.refreshOnLoad)
            self.assertEqual(pivot.cache.recordCount, 0)
            self.assertIsNone(result[name]["A4"].value)
        result.close()

    def test_wrong_feed_month_is_rejected(self) -> None:
        with self.assertRaisesRegex(ValueError, "Feed Date outside"):
            _banner_rows(self._banner("FRST", [("90000001", "Able, Ada", 10)]),
                         "FRST", "2026-08")

    def test_recon_does_not_request_template(self) -> None:
        definition = get_workflow("textbook_recon_manual")
        self.assertNotIn("template_file", [item.key for item in definition.parameters])
        self.assertEqual(definition.name, "Textbook Recon")
        self.assertEqual([item.workflow_id for item in get_workflows()
                          if item.name.startswith("Textbook Recon")], ["textbook_recon_manual"])

    def test_sql_month_is_bounded(self) -> None:
        sql = _render_sql("2026-12")
        self.assertIn("DATE '2026-12-01'", sql)
        self.assertIn("DATE '2027-01-01'", sql)

    def test_sql_result_keeps_recon_fields_and_raw_transaction_columns(self) -> None:
        frame = pd.DataFrame([{
            "'ID'": "90000001", "'Name'": "Able, Ada",
            "tbraccd_detail_code": "FRST", "tbraccd_amount": 10,
            "tbraccd_feed_date": "2026-07-15", "tbraccd_tran_number": 123,
            "extract_row_count": 1,
        }])
        result = _normalized_banner(frame, "FRST")
        self.assertEqual(result.loc[0, "'Amount'"], 10)
        self.assertEqual(result.loc[0, "'Detail Code'"], "FRST")
        self.assertEqual(result.loc[0, "tbraccd_tran_number"], 123)

    def test_nonadmin_view_cannot_show_sql_route(self) -> None:
        run = lambda context: WorkflowResult(True, "done")
        workflows = (
            WorkflowDefinition("manual", "Manual", "", "", run),
            WorkflowDefinition("sql", "SQL", "", "", run, administrator_only=True),
        )
        policy = AccessConfiguration({}, {
            "administrator": frozenset({"*"}),
            "analyst": frozenset({"manual", "sql"}),
        })
        self.assertEqual(
            [item.workflow_id for item in filter_workflows_for_view(workflows, policy, "analyst")],
            ["manual"],
        )
        self.assertEqual(len(filter_workflows_for_view(workflows, policy, "administrator")), 2)

    def test_legacy_sql_card_grant_does_not_break_home_page(self) -> None:
        policy = AccessConfiguration({}, {
            "administrator": frozenset({"*"}),
            "analyst": frozenset({"textbook_recon_manual", "textbook_recon_sql"}),
        })
        self.assertEqual(
            [item.name for item in filter_workflows_for_view(get_workflows(), policy, "analyst")],
            ["Textbook Recon"],
        )

    def test_sql_choice_accepts_only_production_mode(self) -> None:
        policy = SimpleNamespace(is_administrator=lambda login: True)
        with patch("app.gui.services.textbook_recon.get_shared_gui_access_path", return_value=self.root / "access.json"), \
             patch("app.gui.services.textbook_recon.load_access_configuration", return_value=policy), \
             patch("app.gui.services.textbook_recon.get_current_login", return_value="admin"):
            with self.assertRaisesRegex(ValueError, "PROD Insights"):
                run_recon(WorkflowContext("textbook_recon_manual", {"banner_source": "sql"}, WorkflowMode.TEST))

    def test_nonadmin_sql_choice_is_rechecked_by_runner(self) -> None:
        policy = SimpleNamespace(is_administrator=lambda login: False)
        with patch("app.gui.services.textbook_recon.get_shared_gui_access_path", return_value=self.root / "access.json"), \
             patch("app.gui.services.textbook_recon.load_access_configuration", return_value=policy), \
             patch("app.gui.services.textbook_recon.get_current_login", return_value="staff"):
            with self.assertRaisesRegex(ValueError, "Only an administrator"):
                run_recon(WorkflowContext("textbook_recon_manual", {"banner_source": "sql"}, WorkflowMode.PRODUCTION))


if __name__ == "__main__":
    unittest.main()
