"""Synthetic-only trends SQL, streaming and GUI authorization contracts."""
from dataclasses import asdict
import csv
import json
from pathlib import Path
import re
import tempfile
import unittest
from unittest.mock import MagicMock, patch

import pandas as pd
from openpyxl import Workbook, load_workbook

from app.gui.models import ParameterDefinition, ParameterKind, WorkflowContext, WorkflowMode
from app.gui.services.parameters import parse_parameters
from app.gui.services.trends import run_trends
from app.gui.services.access import AccessConfiguration, filter_workflows_for_view
from app.gui.workflow_registry import get_workflow
from data_processing.trends.analysis import analyze_all_files
from data_processing.trends.config import REQUIRED_COLUMNS, TrendsConfig
from data_processing.trends.extract import TrendsExtractSettings, extract_trends, run_sql_trends_pipeline
from shared.cancellation import CancellationToken, WorkflowCancelled
from shared.insights.client import InsightsAPIError

COLUMNS = list(REQUIRED_COLUMNS) + ["extract_pidm", "extract_transaction", "extract_remaining_count"]
METADATA = {"TUIT": {"type": "C", "category": "TUI", "detail_code_description": "Tuition"}}


class FakeInsights:
    def __init__(self, rows=(), cap=1000):
        self.rows = rows
        self.cap = cap
        self.sql = []

    def run_sql(self, sql):
        self.sql.append(sql)
        if "FROM taismgr.tbbdetc" in sql:
            return pd.DataFrame([{"detail_code": "TUIT", **METADATA["TUIT"], "extract_remaining_count": 1}])
        start, end = re.findall(r"DATE '([0-9-]+)'", sql)
        cursor = re.search(r"> \((\d+), (\d+)\)", sql)
        key = tuple(map(int, cursor.groups())) if cursor else (-1, -1)
        eligible = sorted([row for row in self.rows if start <= row[0] < end and tuple(row[1:3]) > key], key=lambda row: row[1:3])
        limit = int(re.search(r"LIMIT (\d+)", sql).group(1))
        return pd.DataFrame([list(row[3:]) + list(row[1:3]) + [len(eligible)] for row in eligible[:min(limit, self.cap)]], columns=COLUMNS)


def transaction(day, pidm, number, student="SYNTH001", amount=100):
    return (day, pidm, number, student, "TUIT", "Tuition", amount, 50, "202610")


class TrendsSQLTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        access_path = patch("app.gui.services.trends.get_shared_gui_access_path", return_value=self.root / "access.json")
        access_path.start(); self.addCleanup(access_path.stop)
        self.settings = TrendsExtractSettings(2026, 2026, self.root / "sources", page_size=2)
        self.metadata = self.root / "codes.json"
        self.metadata.write_text(json.dumps(METADATA))
        reference_path = patch("app.gui.services.trends.get_shared_detail_codes_path", return_value=self.metadata)
        reference_path.start(); self.addCleanup(reference_path.stop)
        self.config = TrendsConfig(detail_codes_file=self.metadata, output_file=self.root / "review.xlsx")

    def read(self, path):
        with path.open(newline="") as stream:
            return list(csv.DictReader(stream))

    def test_month_year_boundaries_pages_and_legitimate_duplicates(self):
        executor = FakeInsights([
            transaction("2025-06-30", 1, 1), transaction("2025-07-01", 1, 2),
            transaction("2025-07-01", 1, 3), transaction("2025-07-31", 2, 1),
            transaction("2025-08-01", 3, 1), transaction("2026-06-30", 4, 1),
            transaction("2026-07-01", 5, 1),
        ])
        files = extract_trends(executor, self.settings)
        self.assertEqual(len(self.read(files[0])), 5)
        self.assertGreater(len(executor.sql), 12)
        self.assertTrue(all("__" not in sql and "COUNT(*) OVER ()" in sql for sql in executor.sql))
        self.assertFalse(list(self.settings.extract_directory.glob(".*.tmp.csv")))

    def test_smaller_api_cap_retries_cursor_without_losing_rows(self):
        executor = FakeInsights([transaction("2025-07-01", 1, n) for n in range(1, 6)], cap=1)
        files = extract_trends(executor, self.settings)
        self.assertEqual(len(self.read(files[0])), 5)
        self.assertIn("LIMIT 2", executor.sql[0])
        self.assertIn("LIMIT 1", executor.sql[1])
        self.assertIn("AND TRUE", executor.sql[1])

    def test_csv_and_workbook_analysis_match(self):
        rows = [transaction("2025-07-01", 1, 1, "000123", 150), transaction("2025-07-01", 1, 2, "000123", -20)]
        files = extract_trends(FakeInsights(rows), self.settings)
        workbook = Workbook(); sheet = workbook.active
        sheet.append(list(REQUIRED_COLUMNS))
        for row in rows: sheet.append(list(row[3:]))
        path = self.root / "TGIACCD_fy2026.xlsx"
        workbook.save(path); workbook.close()
        csv_analysis = asdict(analyze_all_files(files, detail_codes=METADATA, config=self.config))
        excel_analysis = asdict(analyze_all_files([path], detail_codes=METADATA, config=self.config))
        csv_analysis["fiscal_years"][0]["source_file"] = path.name
        self.assertEqual(csv_analysis, excel_analysis)
        self.assertEqual(self.read(files[0])[0]["ID"], "000123")

    def test_rejects_missing_guards_invalid_counts_and_duplicate_keys(self):
        for kind in ("missing", "fractional", "inconsistent", "duplicates"):
            with self.subTest(kind=kind):
                executor = FakeInsights([transaction("2025-07-01", 1, 1), transaction("2025-07-01", 1, 2)])
                original = executor.run_sql
                def bad(sql):
                    frame = original(sql)
                    if kind == "missing": return frame.drop(columns="extract_remaining_count")
                    if kind == "fractional": frame["extract_remaining_count"] = 2.5
                    if kind == "inconsistent": frame["extract_remaining_count"] = [2, 3]
                    if kind == "duplicates": frame["extract_transaction"] = 1
                    return frame
                executor.run_sql = bad
                settings = TrendsExtractSettings(2026, 2026, self.root / kind, 2)
                with self.assertRaises(InsightsAPIError): extract_trends(executor, settings)
                self.assertFalse(list(settings.extract_directory.iterdir()))

    def test_missing_or_changed_later_pages_fail_closed(self):
        for kind in ("empty", "changed"):
            with self.subTest(kind=kind):
                executor = FakeInsights([transaction("2025-07-01", 1, n) for n in range(1, 5)])
                original = executor.run_sql
                def bad(sql):
                    frame = original(sql)
                    if "> (" in sql:
                        if kind == "empty": return pd.DataFrame(columns=COLUMNS)
                        frame["extract_remaining_count"] = 3
                    return frame
                executor.run_sql = bad
                settings = TrendsExtractSettings(2026, 2026, self.root / kind, 2)
                with self.assertRaises(InsightsAPIError): extract_trends(executor, settings)
                self.assertFalse(list(settings.extract_directory.iterdir()))

    def test_admin_view_sees_card_and_staff_view_cannot(self):
        definition = get_workflow("trends")
        policy = AccessConfiguration({}, {"administrator": frozenset({"*"}), "analyst": frozenset({"*"})})
        self.assertEqual(filter_workflows_for_view([definition], policy, "administrator"), (definition,))
        self.assertEqual(filter_workflows_for_view([definition], policy, "analyst"), ())

    def test_empty_year_has_headers_and_zero_analysis(self):
        files = extract_trends(FakeInsights(), self.settings)
        self.assertEqual(self.read(files[0]), [])
        analysis = analyze_all_files(files, detail_codes=METADATA, config=self.config)
        self.assertEqual(analysis.fiscal_years[0].row_count, 0)

    def test_multiple_years_are_independent(self):
        settings = TrendsExtractSettings(2025, 2026, self.root / "two-years")
        files = extract_trends(FakeInsights([transaction("2025-06-30", 1, 1), transaction("2025-07-01", 1, 2)]), settings)
        self.assertEqual([p.name for p in files], ["TGIACCD_fy2025.csv", "TGIACCD_fy2026.csv"])
        self.assertEqual([len(self.read(p)) for p in files], [1, 1])

    def test_output_and_sources_are_not_overwritten(self):
        self.config.output_file.write_text("existing")
        executor = FakeInsights()
        with self.assertRaises(ValueError): run_sql_trends_pipeline(executor, self.settings, self.config)
        self.assertEqual(executor.sql, [])
        self.assertEqual(self.config.output_file.read_text(), "existing")
        self.settings.extract_directory.mkdir()
        with self.assertRaises(FileExistsError): extract_trends(executor, self.settings)

    def test_cancelled_page_does_not_publish_annual_source_or_report(self):
        token = CancellationToken()
        executor = FakeInsights([transaction("2025-07-01", 1, 1)])
        original = executor.run_sql
        def cancel(sql):
            frame = original(sql)
            if "FROM taismgr.tbraccd" in sql:
                token.request()
            return frame
        executor.run_sql = cancel
        with self.assertRaises(WorkflowCancelled): run_sql_trends_pipeline(executor, self.settings, self.config, cancellation=token)
        self.assertFalse(self.config.output_file.exists())
        self.assertFalse(list(self.settings.extract_directory.iterdir()))

    def test_end_to_end_creates_existing_report_sheets(self):
        output = run_sql_trends_pipeline(FakeInsights([transaction("2025-07-01", 1, 1)]), self.settings, self.config)
        workbook = load_workbook(output)
        self.addCleanup(workbook.close)
        self.assertIn("Checks", workbook.sheetnames)
        self.assertIn("Detail Code Totals", workbook.sheetnames)
        self.assertTrue((self.settings.extract_directory / "TGIACCD_fy2026.csv").exists())

    def test_invalid_settings(self):
        for first, last, size in [(2026, 2025, 1000), (26, 2026, 1000), (2026, 2026, 0), (2026, 2026, True)]:
            with self.assertRaises(ValueError): TrendsExtractSettings(first, last, self.root / "invalid", size)

    def test_registry_and_service_enforce_admin_prod(self):
        definition = get_workflow("trends")
        self.assertTrue(definition.administrator_only)
        self.assertTrue(definition.cancellable)
        self.assertEqual(definition.supported_modes, ())
        with patch("app.gui.services.trends.load_access_configuration") as policy, patch("app.gui.services.trends.build_authenticated_client") as auth:
            policy.return_value.is_administrator.return_value = False
            with self.assertRaisesRegex(ValueError, "administrator"):
                run_trends(WorkflowContext("trends", {}, WorkflowMode.PRODUCTION))
            policy.return_value.is_administrator.return_value = True
            with self.assertRaisesRegex(ValueError, "PROD"):
                run_trends(WorkflowContext("trends", {"banner_source": "sql"}, WorkflowMode.TEST))
            auth.assert_not_called()

    def test_gui_service_authenticates_and_runs_checked_pipeline(self):
        inputs = {"banner_source": "sql", "first_fiscal_year": "2026", "last_fiscal_year": "2026",
                  "output_file": self.config.output_file}
        executor = FakeInsights([transaction("2025-07-01", 1, 1)])
        client = MagicMock(); client.run_sql.side_effect = executor.run_sql
        with patch("app.gui.services.trends.load_access_configuration") as policy, patch("app.gui.services.trends.load_department_profiles") as profiles, patch("app.gui.services.trends.build_authenticated_client", return_value=(client, None)):
            policy.return_value.is_administrator.return_value = True
            profiles.return_value = {"PROD": MagicMock()}
            result = run_trends(WorkflowContext("trends", inputs, WorkflowMode.PRODUCTION))
        self.assertTrue(result.success)
        self.assertTrue(result.output_path.exists())

    def test_sql_initializes_reference_without_existing_excel_json(self):
        self.metadata.unlink()
        inputs = {"banner_source": "sql", "first_fiscal_year": "2026", "last_fiscal_year": "2026",
                  "output_file": self.config.output_file}
        executor = FakeInsights([transaction("2025-07-01", 1, 1)])
        client = MagicMock(); client.run_sql.side_effect = executor.run_sql
        with patch("app.gui.services.trends.load_access_configuration") as policy, patch("app.gui.services.trends.load_department_profiles", return_value={"PROD": MagicMock()}), patch("app.gui.services.trends.build_authenticated_client", return_value=(client, None)):
            policy.return_value.is_administrator.return_value = True
            result = run_trends(WorkflowContext("trends", inputs, WorkflowMode.PRODUCTION))
        self.assertTrue(result.success)
        self.assertEqual(json.loads(self.metadata.read_text()), METADATA)
        self.assertIn("FROM taismgr.tbbdetc", executor.sql[0])

    def test_manual_annual_uploads_ignore_sql_year_inputs_and_never_authenticate(self):
        annual = self.root / "TGIACCD_fy2026.xlsx"
        workbook = Workbook(); sheet = workbook.active
        sheet.append(list(REQUIRED_COLUMNS)); sheet.append(list(transaction("2025-07-01", 1, 1)[3:]))
        workbook.save(annual); workbook.close()
        original = annual.read_bytes()
        inputs = {"banner_source": "manual", "banner_file": (annual,),
                  "first_fiscal_year": "invalid", "last_fiscal_year": "invalid",
                  "output_file": self.config.output_file}
        with patch("app.gui.services.trends.load_access_configuration") as policy, patch("app.gui.services.trends.build_authenticated_client") as auth:
            policy.return_value.is_administrator.return_value = True
            result = run_trends(WorkflowContext("trends", inputs))
            auth.assert_not_called()
        self.assertTrue(result.success)
        self.assertEqual(annual.read_bytes(), original)
        self.assertFalse((self.root / "review_sources").exists())

    def test_manual_rejects_prod_existing_output_and_duplicate_fiscal_years(self):
        inputs = {"banner_source": "manual", "banner_file": (), "detail_codes_file": self.metadata,
                  "output_file": self.config.output_file}
        with patch("app.gui.services.trends.load_access_configuration") as policy, patch("app.gui.services.trends.build_authenticated_client") as auth:
            policy.return_value.is_administrator.return_value = True
            with self.assertRaisesRegex(ValueError, "TEST or PROD"):
                run_trends(WorkflowContext("trends", inputs, WorkflowMode.PRODUCTION))
            self.config.output_file.write_text("keep")
            with self.assertRaisesRegex(ValueError, "new .xlsx"):
                run_trends(WorkflowContext("trends", inputs))
            self.config.output_file.unlink()
            annual = self.root / "TGIACCD_fy2026.xlsx"; annual.touch()
            inputs["banner_file"] = (annual, annual)
            with self.assertRaisesRegex(ValueError, "More than one"):
                run_trends(WorkflowContext("trends", inputs))
            auth.assert_not_called()

    def test_year_parameter_parses_integers_and_rejects_invalid_or_out_of_range(self):
        definition = ParameterDefinition("year", "Fiscal year", ParameterKind.INTEGER,
                                         minimum=1900, maximum=9998)
        self.assertEqual(parse_parameters((definition,), {"year": "2026"}), {"year": 2026})
        for value in ("2026.5", "abc", "1899", "9999", ""):
            with self.assertRaises(ValueError): parse_parameters((definition,), {"year": value})

    def test_missing_shared_reference_explains_setup_before_authentication(self):
        self.metadata.unlink()
        inputs = {"banner_source": "manual", "banner_file": (),
                  "output_file": self.config.output_file}
        with patch("app.gui.services.trends.load_access_configuration") as policy, patch("app.gui.services.trends.build_authenticated_client") as auth:
            policy.return_value.is_administrator.return_value = True
            with self.assertRaisesRegex(ValueError, "beside gui_access.json"):
                run_trends(WorkflowContext("trends", inputs))
            auth.assert_not_called()
        self.assertFalse(self.config.output_file.exists())

    def test_auth_exception_hides_sensitive_chain(self):
        inputs = {"banner_source": "sql", "first_fiscal_year": "2026", "last_fiscal_year": "2026", "output_file": self.config.output_file}
        with patch("app.gui.services.trends.load_access_configuration") as policy, patch("app.gui.services.trends.load_department_profiles", return_value={"PROD": MagicMock()}), patch("app.gui.services.trends.build_authenticated_client", side_effect=RuntimeError("secret-token")):
            policy.return_value.is_administrator.return_value = True
            with self.assertRaisesRegex(ValueError, "Could not connect") as caught:
                run_trends(WorkflowContext("trends", inputs, WorkflowMode.PRODUCTION))
        self.assertTrue(caught.exception.__suppress_context__)
        self.assertNotIn("secret", str(caught.exception))


if __name__ == "__main__": unittest.main()
