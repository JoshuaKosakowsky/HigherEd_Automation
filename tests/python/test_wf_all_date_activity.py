"""Run the actual all-history SQL under a synthetic native API row cap."""
import unittest
from unittest.mock import MagicMock, patch

import pandas as pd

from app.gui.models import WorkflowMode
from openpyxl import load_workbook
from shared.insights.all_date_activity import extract_all_date_activity
from shared.insights.banner_activity import BannerDetailCodeParameters
from shared.insights.client import InsightsAPIError
from shared.insights.query_catalog import get_query
from shared.insights.report_results import InsightsReportResultError
import test_tgiaccd_query as sql_fixtures
import test_wf_mines_balance_recon as workflow_fixtures


class AllDateActivityTests(unittest.TestCase):
    def setUp(self):
        self.fixture = sql_fixtures.TgiaccdQueryTests()
        self.fixture.setUp()
        self.addCleanup(self.fixture.doCleanups)
        self.parameters = BannerDetailCodeParameters(('WOFF', 'BDRC'))
        self.template = get_query('wf_mines_balance_transactions').sql_path.read_text()
        self.calls = []
        self.cap = 2

    def run_sql(self, sql):
        self.calls.append(sql)
        return pd.read_sql_query(sql, self.fixture.db).iloc[:self.cap].copy()

    def extract(self):
        executor = MagicMock()
        executor.run_sql.side_effect = self.run_sql
        return extract_all_date_activity(executor, self.parameters, sql_template=self.template)

    def populate(self):
        for number, pidm, day, code, amount in (
            (1, 1, '1900-01-01', 'WOFF', 10),
            (2, 1, '1900-01-01', 'WOFF', 10),  # Legitimate duplicate output rows.
            (3, 1, None, 'BDRC', -5),
            (4, 2, '2199-12-31', 'WOFF', 0),
            (5, None, '2099-01-01', 'BDRC', 20),
            (None, 3, None, 'WOFF', -2),
        ):
            self.fixture.add_transaction(number, day, code, amount, pidm=pidm)

    def test_caps_partition_all_history_including_null_keys_and_duplicate_transactions(self):
        self.populate()
        expected = pd.read_sql_query(self.parameters.render_sql(self.template), self.fixture.db)
        actual = self.extract()
        self.assertEqual(list(actual.columns), list(expected.columns[:-1]))
        self.assertEqual(len(actual), 6)
        self.assertEqual(actual["'Amount'"].sum(), 33)
        self.assertEqual(actual["'Amount'"].tolist().count(10), 2)
        self.assertEqual(set(actual["'Feed Date'"].dropna()), {'1900-01-01', '2199-12-31', '2099-01-01'})
        self.assertEqual(actual["'Feed Date'"].isna().sum(), 2)
        actual_rows = actual.astype(object).where(pd.notna(actual), None)
        expected = expected.drop(columns=['Report Row Count'])
        expected_rows = expected.astype(object).where(pd.notna(expected), None)
        self.assertCountEqual(list(actual_rows.itertuples(index=False, name=None)),
                              list(expected_rows.itertuples(index=False, name=None)))
        self.assertTrue(any('t.tbraccd_pidm IS NULL' in sql for sql in self.calls))
        self.assertTrue(any('t.tbraccd_tran_number' in sql and '__min_key' in sql for sql in self.calls))
        self.assertFalse(any('CURRENT_DATE' in sql or 'CAST(' in sql for sql in self.calls))

    def test_complete_and_empty_results_need_no_partition_discovery(self):
        self.assertTrue(self.extract().empty)
        self.assertEqual(len(self.calls), 1)
        self.calls.clear()
        self.fixture.add_transaction(1, None, 'WOFF')
        self.assertEqual(len(self.extract()), 1)
        self.assertEqual(len(self.calls), 1)

    def test_null_transaction_keys_within_one_capped_account_are_kept(self):
        for number in (None, 1, 2, 3):
            self.fixture.add_transaction(number, None, 'WOFF', 10, pidm=1)
        result = self.extract()
        self.assertEqual(len(result), 4)
        self.assertEqual(result["'Amount'"].sum(), 40)
        self.assertTrue(any('t.tbraccd_tran_number IS NULL' in sql for sql in self.calls))

    def test_missing_and_invalid_counts_fail_without_retrying(self):
        for frame in (pd.DataFrame({'Value': [1]}), pd.DataFrame({'Report Row Count': [1.0]})):
            executor = MagicMock()
            executor.run_sql.return_value = frame
            with self.assertRaises(InsightsReportResultError):
                extract_all_date_activity(executor, self.parameters, sql_template=self.template)
            executor.run_sql.assert_called_once()

    def test_one_unsplittable_key_group_fails_without_deduplication(self):
        for _ in range(3):
            self.fixture.add_transaction(1, None, 'WOFF')
        with self.assertRaises(InsightsAPIError) as caught:
            self.extract()
        self.assertEqual(caught.exception.code, 'key_group_cap')

    def test_source_count_changes_fail_instead_of_exporting_a_partial_result(self):
        self.populate()
        base = self.run_sql
        def changing_source(sql):
            result = base(sql)
            if len(self.calls) == 1:
                self.fixture.db.execute('DELETE FROM taismgr.tbraccd WHERE tbraccd_tran_number = 4')
            return result
        executor = MagicMock()
        executor.run_sql.side_effect = changing_source
        with self.assertRaises(InsightsAPIError) as caught:
            extract_all_date_activity(executor, self.parameters, sql_template=self.template)
        self.assertEqual(caught.exception.code, 'source_changed')

    def test_legacy_capped_all_date_query_still_recovers_in_connections(self):
        from app.gui.models import WorkflowContext
        from app.gui.services.insights import run_insights_connection
        self.populate()
        fixture = workflow_fixtures.WFMinesBalanceTests()
        fixture.setUp()
        self.addCleanup(fixture.doCleanups)
        client = MagicMock()
        client.__enter__.return_value = client
        client.run_sql.side_effect = self.run_sql
        with patch('app.gui.services.insights.load_access_configuration') as policy, patch(
            'app.gui.services.insights.get_shared_gui_access_path', return_value=fixture.root / 'access.json',
        ), patch('app.gui.services.insights.load_department_profiles'), patch(
            'app.gui.services.insights.build_authenticated_client', return_value=(client, 'cached'),
        ):
            policy.return_value.is_administrator.return_value = True
            result = run_insights_connection(WorkflowContext('insights_query_export', {
                'action': 'query_export', 'query_id': 'wf_mines_balance_transactions',
                'output_path': fixture.output,
            }, WorkflowMode.PRODUCTION))
        self.assertTrue(result.success, result.message)
        client.run_sql_file.assert_not_called()
        book = load_workbook(fixture.output)
        try:
            self.assertEqual(book.active.max_row, 7)
            self.assertEqual(book.active.max_column, 12)
        finally:
            book.close()
