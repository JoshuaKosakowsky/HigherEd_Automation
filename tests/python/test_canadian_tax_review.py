"""Synthetic financial controls and administrator adapter checks; no student data."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import MagicMock, patch
from zipfile import ZipFile
import json

import pandas as pd
from openpyxl import load_workbook

from app.gui.models import WorkflowContext, WorkflowMode
from app.gui.services.canadian_tax import run_canadian_tax
from app.gui.services.access import AccessConfiguration, UserAccessProfile, filter_workflows_for_view
from app.gui.workflow_registry import get_workflow
from data_processing.canadian_tax.exchange_rates import CRA_GUIDANCE, RATE_API, RATE_PAGE
from data_processing.canadian_tax.pipeline import prepare_review, read_source_package, write_package
from data_processing.canadian_tax.preparation import load_rules
from shared.cancellation import CancellationToken, WorkflowCancelled


def sources() -> tuple[dict[str, pd.DataFrame], dict]:
    transactions = []
    for number, code, kind, amount, balance in [
        (1, "TU01", "C", "1000.00", "200.00"),
        (2, "FEIT", "C", "30.00", "0.00"), (3, "CFEE", "C", "10.00", "0.00"),
        (4, "SCH1", "P", "600.00", "0.00"), (5, "CASH", "P", "240.00", "0.00"),
    ]:
        transactions.append(dict(tran_number=number, term_code="202555", term_description="Summer 2025",
            detail_code=code, type_ind=kind, amount_usd=amount, stored_balance_usd=balance,
            year_term_candidate="Y", effective_date="2025-05-01"))
    applications = [dict(application_id=index, payment_tran_number=pay, charge_tran_number=charge,
        application_amount_usd=amount, reapplication_ind=None, direct_payment_ind=None,
        direct_payment_type=None) for index, (pay, charge, amount) in enumerate([
            (4, 1, "600.00"), (5, 1, "200.00"), (5, 2, "30.00"), (5, 3, "10.00")], 1)]
    enrollment = [dict(term_code="202555", term_description="Summer 2025", crn=101,
        counts_in_enrollment="Y", credit_hours=6, section_start_date="2025-05-01", section_end_date="2025-05-21",
        part_of_term_start_date="2025-05-01", part_of_term_end_date="2025-05-21")]
    raw = dict(identity=[dict(cwid="SYNTHETIC001", first_name="Synthetic", last_name="Example")],
               transactions=transactions, payment_applications=applications, enrollment=enrollment,
               programs=[dict(term_code="202555", program_code="BS-TEST", degree_code="BS", level_code="UG")])
    frames = {}
    for name, records in raw.items():
        frames[name] = pd.DataFrame([{**row, "pidm": 1, "extract_row_count": len(records), "extract_tax_year": 2025} for row in records])
    frames["schema"] = pd.DataFrame({"validation_status": ["FOUND"]})
    frames["payment_application_columns"] = pd.DataFrame({"column_name": ["synthetic"]})
    manifest = dict(app_version="synthetic", environment="TEST", tax_year=2025,
        exchange_rate=dict(status="published", cad_per_usd="1.3978", rate_year=2025,
            direction="CAD per 1 USD", series="FXAUSDCAD", method="Bank of Canada published annual average",
            source_url=RATE_PAGE, api_url=RATE_API, guidance_url=CRA_GUIDANCE),
        bank_of_canada_response={"observations": [{"d": "2025-01-01", "FXAUSDCAD": {"v": "1.3978"}}]})
    return frames, manifest


def complete(frames: dict[str, pd.DataFrame]) -> None:
    for frame in frames.values():
        if "extract_row_count" in frame:
            frame["extract_row_count"] = len(frame)


class TuitionReviewTests(unittest.TestCase):
    def review(self, frames=None, manifest=None):
        default_frames, default_manifest = sources()
        return prepare_review(frames if frames is not None else default_frames,
                              manifest if manifest is not None else default_manifest, load_rules())

    def test_partial_paid_scholarships_exclusions_summer_and_rounding(self):
        report = self.review()
        self.assertEqual(report.paid_usd, Decimal("800.00"))  # not the 1,000 charge or 840 payment pool
        self.assertEqual(report.paid_cad, Decimal("1118.24"))
        self.assertEqual(report.sessions.retained_paid_usd.tolist(), [Decimal("800.00")])
        self.assertEqual(report.charges.retained_paid_usd.tolist(), [Decimal("800"), Decimal(0), Decimal(0)])
        self.assertTrue(report.reconciliation.difference_usd.eq(0).all())
        self.assertEqual(report.checks.iloc[-1].result, "PENDING")

    def test_paired_reapplications_do_not_double_count(self):
        frames, manifest = sources()
        a = frames["payment_applications"].iloc[0].to_dict()
        frames["payment_applications"] = pd.concat([frames["payment_applications"], pd.DataFrame([
            {**a, "application_id": 5, "reapplication_ind": "Y"},
            {**a, "application_id": 6, "payment_tran_number": 1, "charge_tran_number": 4, "reapplication_ind": "Y"},
        ])], ignore_index=True)
        complete(frames)
        report = self.review(frames, manifest)
        self.assertEqual(report.paid_usd, Decimal("800"))
        self.assertEqual(report.applications.calculation_treatment.tolist()[-2:], ["cancelled_reapplication_pair"] * 2)

    def test_unmatched_pair_and_sequential_balance_mismatch_withhold_amounts(self):
        for scenario in ("missing_pair", "balance"):
            frames, manifest = sources()
            if scenario == "missing_pair":
                frames["payment_applications"].loc[0, "reapplication_ind"] = "Y"
            else:
                frames["transactions"].loc[0, "stored_balance_usd"] = "0.00"
            report = self.review(frames, manifest)
            self.assertIsNone(report.paid_usd)
            self.assertIsNone(report.paid_cad)

    def test_charge_credit_and_payment_reversal_are_not_paid_tuition(self):
        frames, manifest = sources()
        t = frames["transactions"].iloc[0].to_dict()
        frames["transactions"] = pd.concat([frames["transactions"], pd.DataFrame([
            {**t, "tran_number": 6, "amount_usd": "-100", "stored_balance_usd": "0"},
            {**t, "tran_number": 7, "type_ind": "P", "amount_usd": "-50", "stored_balance_usd": "0"},
        ])], ignore_index=True)
        frames["transactions"].loc[0, "stored_balance_usd"] = "100"
        frames["transactions"].loc[4, "amount_usd"] = "290"
        a = frames["payment_applications"].iloc[0].to_dict()
        frames["payment_applications"] = pd.concat([frames["payment_applications"], pd.DataFrame([
            {**a, "application_id": 5, "payment_tran_number": 6, "application_amount_usd": "100"},
            {**a, "application_id": 6, "payment_tran_number": 5, "charge_tran_number": 7, "application_amount_usd": "50"},
        ])], ignore_index=True)
        complete(frames)
        report = self.review(frames, manifest)
        self.assertEqual(report.paid_usd, Decimal("800"))
        self.assertEqual(report.applications.calculation_treatment.tolist()[-2:], ["charge_credit", "payment_reversal"])

    def test_advance_late_missing_dates_preserve_usd_withhold_cad(self):
        for date in ("2024-12-15", "2026-01-05", None, "bad-date", "NaT"):
            frames, manifest = sources()
            frames["transactions"].loc[3, "effective_date"] = date
            report = self.review(frames, manifest)
            self.assertEqual(report.paid_usd, Decimal("800"))
            self.assertIsNone(report.paid_cad)

    def test_unavailable_mismatched_rate_keeps_usd_without_fallback(self):
        for scenario in ("unavailable", "wrong_year", "wrong_value"):
            frames, manifest = sources()
            if scenario == "unavailable":
                manifest["bank_of_canada_response"] = None
                manifest["exchange_rate"].update(status="unavailable", cad_per_usd=None)
            elif scenario == "wrong_year":
                manifest["exchange_rate"]["rate_year"] = 2024
            else:
                manifest["exchange_rate"]["cad_per_usd"] = "1.5"
            report = self.review(frames, manifest)
            self.assertEqual(report.paid_usd, Decimal("800"))
            self.assertIsNone(report.paid_cad)

    def test_unsupported_sign_flags_missing_link_and_short_course_withhold(self):
        for scenario in ("negative", "direct", "link", "short", "dropped", "program_tie", "cross_session"):
            frames, manifest = sources()
            if scenario == "negative":
                frames["payment_applications"].loc[0, "application_amount_usd"] = "-600"
            elif scenario == "direct":
                frames["payment_applications"].loc[0, "direct_payment_ind"] = "Y"
            elif scenario == "link":
                frames["payment_applications"].loc[0, "payment_tran_number"] = 999
            elif scenario == "short":
                frames["enrollment"].loc[0, ["section_end_date", "part_of_term_end_date"]] = "2025-05-20"
            elif scenario == "dropped":
                frames["enrollment"].loc[0, "counts_in_enrollment"] = "N"
            elif scenario == "program_tie":
                frames["programs"] = pd.concat([frames["programs"]] * 2, ignore_index=True)
            else:
                frames["enrollment"].loc[0, ["section_start_date", "part_of_term_start_date"]] = "2024-12-01"
            complete(frames)
            with self.subTest(scenario=scenario):
                self.assertIsNone(self.review(frames, manifest).paid_usd)

    def test_package_workbook_audit_sources_exclusivity_and_zip_validation(self):
        frames, manifest = sources()
        report = self.review(frames, manifest)
        with TemporaryDirectory() as directory:
            output = Path(directory) / "review"
            write_package(output, frames, manifest, report)
            workbook = load_workbook(output / "tl11a_review.xlsx")
            try:
                self.assertEqual(workbook["Summary"]["B6"].value, 800)
                self.assertEqual(workbook["Summary"]["B7"].value, 1118.24)
                self.assertEqual(workbook["Summary"]["B4"].value, "SYNTHETIC001")
                self.assertEqual(workbook["Sessions"].max_row, 2)
                self.assertIn(RATE_PAGE, [r[1].value for r in workbook["Rate and Policy"].iter_rows(min_row=2)])
            finally:
                workbook.close()
            with self.assertRaises(FileExistsError):
                write_package(output, frames, manifest, report)
            archive_path = Path(directory) / "source.zip"
            with ZipFile(archive_path, "w") as archive:
                for path in output.iterdir():
                    if path.suffix in {".csv", ".json"}:
                        archive.write(path, f"nested/{path.name}")
            imported_frames, imported_manifest = read_source_package(archive_path)
            self.assertEqual(self.review(imported_frames, imported_manifest).paid_cad, Decimal("1118.24"))
            with ZipFile(archive_path, "a") as archive:
                archive.writestr("another/identity.csv", "not source data")
            with self.assertRaisesRegex(ValueError, "duplicate"):
                read_source_package(archive_path)


class CanadianTaxAdapterTests(unittest.TestCase):
    def setUp(self):
        patcher = patch("app.gui.services.canadian_tax.get_shared_gui_access_path", return_value=Path("synthetic-policy.json"))
        patcher.start()
        self.addCleanup(patcher.stop)

    def test_registry_and_explicit_grants_cannot_expose_to_staff(self):
        definition = get_workflow("canadian_tax")
        self.assertTrue(definition.administrator_only)
        self.assertEqual(definition.default_mode, WorkflowMode.TEST)
        policy = AccessConfiguration(
            users={"staff": UserAccessProfile("staff", "Example", "Analyst", "analyst")},
            workflows_by_view={"analyst": frozenset({"canadian_tax"}), "administrator": frozenset({"*"})})
        self.assertEqual(filter_workflows_for_view((definition,), policy, "analyst"), ())
        self.assertEqual(filter_workflows_for_view((definition,), policy, "administrator"), (definition,))

    def test_adapter_unauthorized_and_cancelled_stop_before_auth(self):
        context = WorkflowContext("canadian_tax", {"cwid": "SYNTHETIC001", "tax_year": "2025",
            "output_directory": "unused"}, WorkflowMode.TEST)
        with patch("app.gui.services.canadian_tax.load_access_configuration") as policy, patch("app.gui.services.canadian_tax.build_authenticated_client") as auth:
            policy.return_value.is_administrator.return_value = False
            with self.assertRaisesRegex(ValueError, "administrator"):
                run_canadian_tax(context)
            policy.return_value.is_administrator.return_value = True
            token = CancellationToken()
            token.request()
            with self.assertRaises(WorkflowCancelled):
                run_canadian_tax(WorkflowContext(context.workflow_id, context.parameters, context.mode, token))
            auth.assert_not_called()

    def test_gui_runner_mocked_auth_publishes_same_review_and_sanitizes_errors(self):
        frames, manifest = sources()
        client = MagicMock()
        client.run_sql_file.side_effect = [frames["schema"], frames["payment_application_columns"]]
        client.run_sql.side_effect = [frames[name] for name in ("identity", "transactions", "enrollment", "programs", "payment_applications")]
        with TemporaryDirectory() as directory:
            context = WorkflowContext("canadian_tax", {"cwid": "SYNTHETIC001", "tax_year": "2025",
                "output_directory": str(Path(directory) / "new")}, WorkflowMode.TEST, CancellationToken())
            with patch("app.gui.services.canadian_tax.load_access_configuration") as policy, patch("app.gui.services.canadian_tax.load_department_profiles", return_value={"TEST": MagicMock()}), patch("app.gui.services.canadian_tax.build_authenticated_client", return_value=(client, "synthetic")), patch("app.gui.services.canadian_tax.fetch_annual_rate", return_value=(manifest["exchange_rate"], manifest["bank_of_canada_response"])):
                policy.return_value.is_administrator.return_value = True
                result = run_canadian_tax(context)
                self.assertTrue(result.success)
                self.assertTrue(result.output_path.is_file())
                saved = json.loads((result.output_path.parent / "manifest.json").read_text())
                self.assertEqual(saved["eligible_paid_cad"], "1118.24")
                # Use a different destination to reach extraction again.
                params = {**context.parameters, "output_directory": str(Path(directory) / "failed")}
                client.run_sql_file.side_effect = RuntimeError("sensitive synthetic SQL literal")
                with self.assertRaises(RuntimeError) as caught:
                    run_canadian_tax(WorkflowContext("canadian_tax", params, WorkflowMode.TEST))
                self.assertNotIn("sensitive", str(caught.exception))
                self.assertTrue(caught.exception.__suppress_context__)
                self.assertFalse(Path(params["output_directory"]).exists())


if __name__ == "__main__":
    unittest.main()
