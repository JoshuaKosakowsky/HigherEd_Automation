"""Run the report SQL with synthetic history and the existing app export service.

SQLite adapters cover date/age and ordered string aggregation syntax only.
These checks do not certify Mines' live PostgreSQL schema or local mappings.
"""
import re
import sqlite3
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from unittest.mock import MagicMock, patch

import pandas as pd
from openpyxl import load_workbook

from app.gui.models import WorkflowContext, WorkflowMode
from app.gui.services.insights import run_insights_connection
from shared.insights.query_catalog import get_query


class OrderedStringAggregate:
    def __init__(self):
        self.values = []
        self.separator = "; "

    def step(self, value, separator, key):
        if value is not None:
            self.values.append((key, value))
            self.separator = separator

    def finalize(self):
        return self.separator.join(value for _, value in sorted(self.values))


def sqlite_sql(sql):
    sql = sql.replace("CURRENT_DATE", "'2099-10-02'")
    sql = sql.replace(
        "CAST(EXTRACT(YEAR FROM AGE(h.from_date, CAST(b.spbpers_birth_date AS date))) AS integer)",
        "age_years(h.from_date, b.spbpers_birth_date)",
    )
    sql = re.sub(r"CAST\(([\w.]+) AS date\)", r"date(\1)", sql)
    return re.sub(
        r"STRING_AGG\(([\w.]+), '; ' ORDER BY ([\w.]+)\)",
        r"ordered_string_agg(\1, '; ', \2)", sql,
    )


def age_years(end, birth):
    if end is None or birth is None:
        return None
    end_date, birth_date = date.fromisoformat(end[:10]), date.fromisoformat(birth[:10])
    return end_date.year - birth_date.year - (
        (end_date.month, end_date.day) < (birth_date.month, birth_date.day)
    )


class CollectionsProfileSQLFixture(unittest.TestCase):
    def setUp(self):
        self.db = sqlite3.connect(":memory:")
        self.db.row_factory = sqlite3.Row
        self.addCleanup(self.db.close)
        self.db.create_function("age_years", 2, age_years)
        self.db.create_aggregate("ordered_string_agg", 3, OrderedStringAggregate)
        self.db.executescript("""
            ATTACH DATABASE ':memory:' AS saturn;
            ATTACH DATABASE ':memory:' AS general;
            CREATE TABLE saturn.sprhold (sprhold_pidm INTEGER, sprhold_hldd_code TEXT,
                sprhold_from_date TEXT, sprhold_to_date TEXT);
            CREATE TABLE saturn.spriden (spriden_pidm INTEGER, spriden_id TEXT,
                spriden_last_name TEXT, spriden_first_name TEXT, spriden_change_ind TEXT);
            CREATE TABLE saturn.spbpers (spbpers_pidm INTEGER, spbpers_birth_date TEXT,
                spbpers_sex TEXT, spbpers_gndr_code TEXT, spbpers_ethn_cde TEXT);
            CREATE TABLE saturn.stvterm (stvterm_code TEXT, stvterm_desc TEXT,
                stvterm_start_date TEXT, stvterm_end_date TEXT);
            CREATE TABLE saturn.sgbstdn (sgbstdn_pidm INTEGER, sgbstdn_term_code_eff TEXT,
                sgbstdn_levl_code TEXT, sgbstdn_styp_code TEXT, sgbstdn_stst_code TEXT,
                sgbstdn_degc_code_1 TEXT, sgbstdn_majr_code_1 TEXT, sgbstdn_majr_code_2 TEXT);
            CREATE TABLE saturn.stvlevl (stvlevl_code TEXT, stvlevl_desc TEXT);
            CREATE TABLE saturn.stvmajr (stvmajr_code TEXT, stvmajr_desc TEXT);
            CREATE TABLE general.gorprac (gorprac_pidm INTEGER, gorprac_race_cde TEXT);
            CREATE TABLE general.gorrace (gorrace_race_cde TEXT, gorrace_desc TEXT);
            CREATE TABLE saturn.sgrsatt (sgrsatt_pidm INTEGER, sgrsatt_term_code_eff TEXT,
                sgrsatt_atts_code TEXT);
            CREATE TABLE saturn.stvatts (stvatts_code TEXT, stvatts_desc TEXT);
            CREATE TABLE saturn.sfrstcr (sfrstcr_pidm INTEGER, sfrstcr_term_code TEXT,
                sfrstcr_rsts_code TEXT, sfrstcr_credit_hr NUMERIC);
            CREATE TABLE saturn.stvrsts (stvrsts_code TEXT, stvrsts_incl_sect_enrl TEXT);
            CREATE TABLE saturn.shrtgpa (shrtgpa_pidm INTEGER, shrtgpa_term_code TEXT,
                shrtgpa_levl_code TEXT, shrtgpa_gpa_type_ind TEXT, shrtgpa_gpa NUMERIC,
                shrtgpa_quality_points NUMERIC, shrtgpa_gpa_hours NUMERIC, shrtgpa_hours_earned NUMERIC);
            CREATE TABLE saturn.shrttrm (shrttrm_pidm INTEGER, shrttrm_term_code TEXT,
                shrttrm_astd_code_end_of_term TEXT);
            CREATE TABLE saturn.stvastd (stvastd_code TEXT, stvastd_desc TEXT);
            INSERT INTO saturn.stvterm VALUES
                ('209810', 'Prior Spring', '2098-01-10', '2098-05-10'),
                ('209880', 'Prior Fall', '2098-08-20', '2098-12-15'),
                ('209910', 'Spring', '2099-01-10', '2099-05-10'),
                ('209980', 'Fall', '2099-08-20', '2099-12-15'),
                ('210010', 'Future Spring', '2100-01-10', '2100-05-10');
            INSERT INTO saturn.stvlevl VALUES ('UG', 'Undergraduate'), ('GR', 'Graduate');
            INSERT INTO saturn.stvmajr VALUES ('MAJA', 'Synthetic Major A'), ('MAJB', 'Synthetic Major B');
            INSERT INTO saturn.stvrsts VALUES ('RE', 'Y'), ('DD', 'N'), ('ZZ', NULL);
            INSERT INTO saturn.stvastd VALUES ('GOOD', 'Good standing'), ('PROB', 'Probation');
            INSERT INTO saturn.stvatts VALUES ('FGA', 'Synthetic First Generation Admissions'),
                ('FGF', 'Synthetic First Generation FAFSA'),
                ('BLSS', 'Synthetic 1st Gen Scholarship'), ('OTHER', 'Synthetic Other Attribute');
        """)
        self.query = get_query("collections_student_profile")

    def insert(self, table, *values):
        self.db.execute(f"INSERT INTO {table} VALUES ({', '.join('?' for _ in values)})", values)

    def student(self, pidm=1, start="2099-09-10", end="9999-12-31", level="UG"):
        self.insert("saturn.sprhold", pidm, " co ", start, end)
        self.insert("saturn.spriden", pidm, f"TEST-{pidm}", "Synthetic", "Student", None)
        self.insert("saturn.spbpers", pidm, "2079-09-11", "F", "SYNTHETIC", "2")
        self.insert("saturn.sgbstdn", pidm, "209810", level, "C", "AS", "BS", "MAJA", None)

    def gpa(self, pidm=1, term="209910", level="UG", kind="I", quality=45, hours=15, earned=15, gpa=3):
        self.insert("saturn.shrtgpa", pidm, term, level, kind, gpa, quality, hours, earned)

    def rows(self):
        return [dict(row) for row in self.db.execute(sqlite_sql(self.query.render_sql()))]


class CollectionsStudentProfileTests(CollectionsProfileSQLFixture):
    def test_all_history_active_boundaries_duplicates_and_missing_identity(self):
        for pidm, start, end in (
            (1, "2098-09-01", "2099-10-01"),
            (2, "2099-09-01", "2099-10-02 23:59:59"),
            (3, "2099-09-01", "9999-12-31"),
            (4, "2099-11-01", "9999-12-31"),
            (5, None, None),
            (6, "2099-10-01", "2099-09-01"),
        ):
            self.student(pidm, start, end)
        self.insert("saturn.sprhold", 3, "CO", "2099-09-01", "9999-12-31")
        self.insert("saturn.sprhold", 3, "CO", "2099-01-10", "2099-05-10")
        self.insert("saturn.sprhold", 7, "CO", "2099-09-01", "9999-12-31")
        self.insert("saturn.sprhold", 8, "TO", "2099-09-01", "9999-12-31")
        rows = self.rows()
        self.assertEqual(len(rows), 8)
        self.assertTrue(all(row["Report Row Count"] == 8 for row in rows))
        by_id = {row["Banner PIDM"]: row for row in rows}
        self.assertEqual([by_id[i]["Active CO Hold"] for i in range(1, 7)],
                         ["No", "No", "Yes", "No", "Unknown", "Unknown"])
        self.assertEqual(by_id[4]["CO Hold Status"], "Scheduled")
        self.assertEqual(by_id[3]["Source Hold Rows"], 2)
        self.assertEqual(by_id[3]["Student CO Hold Count"], 2)
        self.assertIsNone(by_id[7]["CWID"])
        self.assertEqual(by_id[7]["Identity Review"], "Review current identity")

    def test_containing_term_boundaries_gaps_and_overlaps(self):
        for pidm, start in ((1, "2099-08-20"), (2, "2099-12-15"), (3, "2099-07-01")):
            self.student(pidm, start)
        rows = self.rows()
        self.assertEqual([r["CO Term"] for r in rows], ["209980", "209980", None])
        self.assertEqual(rows[2]["Academic Context Review"], "No containing term")
        self.insert("saturn.stvterm", "209981", "Overlapping", "2099-08-01", "2099-12-31")
        rows = self.rows()
        self.assertTrue(all(r["CO Term"] is None for r in rows))
        self.assertEqual(rows[0]["Containing Term Count"], 2)
        self.assertIsNone(rows[0]["CO Term Primary Major"])
        self.assertIsNone(rows[0]["CO Term Registered Credits"])

    def test_program_history_excludes_future_changes_and_keeps_old_hold_context(self):
        self.student()
        self.insert("saturn.sgbstdn", 1, "209980", "UG", "C", "AS", "BS", "MAJB", "MAJA")
        self.insert("saturn.sgbstdn", 1, "210010", "GR", "C", "AS", "MS", "MAJA", None)
        row = self.rows()[0]
        self.assertEqual(row["Program Effective Term"], "209980")
        self.assertEqual(row["CO Term Primary Major"], "Synthetic Major B")
        self.assertEqual(row["CO Term Student Level"], "Undergraduate")
        self.insert("saturn.sprhold", 1, "CO", "2099-01-20", "2099-05-10")
        rows = self.rows()
        self.assertEqual(rows[0]["CO Term Primary Major"], "Synthetic Major A")
        self.assertEqual(rows[1]["CO Term Primary Major"], "Synthetic Major B")

    def test_age_birthday_and_separate_current_demographics_and_multirace(self):
        self.student()
        self.insert("general.gorrace", "A", "Synthetic race A")
        self.insert("general.gorrace", "B", "Synthetic race B")
        for code in ("B", "A", "A"):
            self.insert("general.gorprac", 1, code)
        row = self.rows()[0]
        self.assertEqual(row["Age at CO Start"], 19)
        self.assertEqual(row["Current Recorded Sex"], "Female")
        self.assertEqual(row["Current Gender Identity Code"], "SYNTHETIC")
        self.assertEqual(row["Current Race Codes"], "A; B")
        self.assertEqual(row["Current Races"], "Synthetic race A; Synthetic race B")
        self.assertEqual(row["Current Race Count"], 2)
        self.db.execute("UPDATE saturn.sprhold SET sprhold_from_date = '2099-09-11'")
        self.assertEqual(self.rows()[0]["Age at CO Start"], 20)
        self.db.execute("UPDATE saturn.spbpers SET spbpers_birth_date = '2100-01-01'")
        self.assertIsNone(self.rows()[0]["Age at CO Start"])

    def test_first_gen_sources_are_independent_and_missing_is_unknown(self):
        cases = (
            (1, (), 'Unknown', 'Unknown'),
            (2, ('FGA', 'FGA'), 'Yes', 'Unknown'),
            (3, ('FGF',), 'Unknown', 'Yes'),
            (4, ('FGA', 'FGF'), 'Yes', 'Yes'),
            (5, ('BLSS',), 'Unknown', 'Unknown'),
            (6, (' fga ', 'fgf'), 'Yes', 'Yes'),
        )
        for pidm, codes, _, _ in cases:
            self.student(pidm)
            for code in codes:
                self.insert('saturn.sgrsatt', pidm, '209980', code)
        rows = self.rows()
        self.assertEqual(len(rows), len(cases))
        for row, (pidm, _, admissions, fafsa) in zip(rows, cases):
            with self.subTest(pidm=pidm):
                self.assertEqual(row['First Generation Admissions'], admissions)
                self.assertEqual(row['First Generation FAFSA'], fafsa)
                self.assertEqual(row['Report Row Count'], len(cases))

    def test_first_gen_uses_latest_effective_attribute_set_for_each_hold(self):
        self.student()
        self.insert("saturn.sprhold", 1, "CO", "2099-02-01", "2099-05-10")
        self.insert("saturn.sgrsatt", 1, "209810", "FGA")
        self.insert("saturn.sgrsatt", 1, "209910", "FGF")
        self.insert("saturn.sgrsatt", 1, "209980", "OTHER")
        self.insert("saturn.sgrsatt", 1, "210010", "FGA")
        self.insert("saturn.sgrsatt", 1, "210010", "FGF")
        spring, fall = self.rows()
        self.assertEqual(spring['Attribute Effective Term'], '209910')
        self.assertEqual(spring['First Generation Admissions'], 'Unknown')
        self.assertEqual(spring['First Generation FAFSA'], 'Yes')
        self.assertEqual(fall['Attribute Effective Term'], '209980')
        self.assertEqual(fall['CO Term Attribute Codes'], 'OTHER')
        self.assertEqual(fall['First Generation Admissions'], 'Unknown')
        self.assertEqual(fall['First Generation FAFSA'], 'Unknown')

    def test_first_gen_future_only_or_ambiguous_term_is_unknown(self):
        self.student(1)
        self.insert('saturn.sgrsatt', 1, '210010', 'FGA')
        self.insert('saturn.sgrsatt', 1, '210010', 'FGF')
        self.student(2, start='2099-07-01')  # Between terms.
        self.student(3, start=None)
        for pidm in (2, 3):
            self.insert('saturn.sgrsatt', pidm, '209810', 'FGA')
            self.insert('saturn.sgrsatt', pidm, '209810', 'FGF')
        for row in self.rows():
            self.assertEqual(row['First Generation Admissions'], 'Unknown')
            self.assertEqual(row['First Generation FAFSA'], 'Unknown')
        self.insert('saturn.stvterm', '209981', 'Overlapping', '2099-08-01', '2099-12-31')
        self.student(4)
        self.insert('saturn.sgrsatt', 4, '209810', 'FGA')
        self.insert('saturn.sgrsatt', 4, '209810', 'FGF')
        for row in self.rows():
            self.assertEqual(row['First Generation Admissions'], 'Unknown')
            self.assertEqual(row['First Generation FAFSA'], 'Unknown')

    def test_weighted_prior_gpa_level_and_transfer_separation_and_later_outcomes(self):
        self.student()
        self.gpa(term="209880", quality=60, hours=20, earned=20)
        self.gpa(quality=40, hours=10, earned=10, gpa=4)
        self.gpa(kind="T", quality=0, hours=0, earned=30, gpa=None)
        self.gpa(level="GR", quality=400, hours=100, earned=100, gpa=4)
        self.gpa(term="209980", quality=6, hours=6, earned=0, gpa=1)
        self.insert("saturn.shrttrm", 1, "209910", "GOOD")
        self.insert("saturn.shrttrm", 1, "209980", "PROB")
        row = self.rows()[0]
        self.assertAlmostEqual(row["Reconstructed Institution GPA Before CO"], 3.333)
        self.assertEqual(row["Reconstructed Earned Credits Before CO"], 60)
        self.assertEqual(row["Reconstructed Class at CO Start"], "Junior")
        self.assertEqual(row["Prior Academic History Through Term"], "209910")
        self.assertEqual(row["Prior Academic Standing"], "Good standing")
        self.assertEqual(row["CO Term Institution GPA - Later Outcome"], 1)
        self.assertEqual(row["CO Term End Standing - Later Outcome"], "Probation")

    def test_class_thresholds_graduate_and_no_history_are_not_assumed_freshmen(self):
        for pidm, credits, expected in ((1, 0, "Freshman"), (2, 29.99, "Freshman"),
                                      (3, 30, "Sophomore"), (4, 60, "Junior"), (5, 90, "Senior")):
            self.student(pidm)
            self.gpa(pidm, earned=credits)
            self.assertEqual(self.rows()[pidm - 1]["Reconstructed Class at CO Start"], expected)
        self.student(6, level="GR")
        self.student(7)
        rows = self.rows()
        self.assertEqual(rows[5]["Reconstructed Class at CO Start"], "Graduate")
        self.assertEqual(rows[6]["Reconstructed Class at CO Start"], "Unknown")
        self.assertIsNone(rows[6]["Reconstructed Institution GPA Before CO"])

    def test_missing_history_values_zero_gpa_hours_and_unknown_term_require_review(self):
        self.student()
        self.gpa(quality=None, earned=None)
        row = self.rows()[0]
        self.assertIsNone(row["Reconstructed Institution GPA Before CO"])
        self.assertIsNone(row["Reconstructed Earned Credits Before CO"])
        self.assertEqual(row["Reconstructed Class at CO Start"], "Unknown")
        self.db.execute("DELETE FROM saturn.shrtgpa")
        self.gpa(quality=0, hours=0, earned=10)
        self.assertIsNone(self.rows()[0]["Reconstructed Institution GPA Before CO"])
        self.gpa(term="000000", kind="T", quality=0, hours=0, earned=30)
        row = self.rows()[0]
        self.assertIsNone(row["Reconstructed Earned Credits Before CO"])
        self.assertEqual(row["Prior History Review"], "Missing or invalid history - review")

    def test_registration_excludes_drops_and_preserves_zero_and_unknown(self):
        self.student()
        self.insert("saturn.sfrstcr", 1, "209980", "RE", 12)
        self.insert("saturn.sfrstcr", 1, "209980", "RE", 0)
        self.insert("saturn.sfrstcr", 1, "209980", "DD", 9)
        row = self.rows()[0]
        self.assertEqual(row["CO Term Registered Credits"], 12)
        self.assertEqual(row["CO Term Counting Courses"], 2)
        self.insert("saturn.sfrstcr", 1, "209980", "UNKNOWN", 3)
        self.assertIsNone(self.rows()[0]["CO Term Registered Credits"])
        self.db.execute("DELETE FROM saturn.sfrstcr")
        self.insert("saturn.sfrstcr", 1, "209980", "RE", None)
        self.assertIsNone(self.rows()[0]["CO Term Registered Credits"])

    def test_missing_program_does_not_use_future_program_or_drop_hold(self):
        self.student()
        self.db.execute("UPDATE saturn.sgbstdn SET sgbstdn_term_code_eff = '210010'")
        self.insert('saturn.sgrsatt', 1, '209980', 'FGA')
        row = self.rows()[0]
        self.assertEqual(row["Academic Context Review"], "No program effective by CO term")
        self.assertIsNone(row["CO Term Student Level"])
        self.assertEqual(row['First Generation Admissions'], 'Yes')
        self.assertEqual(row['First Generation FAFSA'], 'Unknown')
        self.assertEqual(row["Report Row Count"], 1)

    def test_empty_result_retains_headers_and_one_authoritative_catalog_source(self):
        cursor = self.db.execute(sqlite_sql(self.query.render_sql()))
        self.assertEqual(cursor.fetchall(), [])
        headings = [column[0] for column in cursor.description]
        self.assertEqual(len(set(headings)), len(headings))
        self.assertIn("CO Term", headings)
        self.assertIn("First Generation Admissions", headings)
        self.assertIn("First Generation FAFSA", headings)
        self.assertNotIn("First Generation", headings)
        self.assertNotIn("First Generation Source", headings)
        self.assertEqual(headings[-1], "Report Row Count")
        self.assertTrue(self.query.administrator_only)
        self.assertIsNone(self.query.term_variable)
        self.assertNotIn("{{", self.query.render_sql())

    def test_schema_inventory_covers_every_source_and_discovery_is_explicit(self):
        inventory = self.query.sql_path.with_name("validate_collections_profile_schema.sql").read_text(encoding="utf-8")
        expected = set(re.findall(r"\('([A-Z]+)', '([A-Z]+)', '([A-Z_0-9]+)'\)", inventory))
        for database in ("saturn", "general"):
            for table_row in self.db.execute(f"SELECT name FROM {database}.sqlite_master WHERE type = 'table'"):
                table = table_row[0]
                for column in self.db.execute(f"PRAGMA {database}.table_info({table})"):
                    self.assertIn((database.upper(), table.upper(), column[1][len(table) + 1:].upper()), expected)
        self.student()
        self.insert("saturn.sgrsatt", 1, "209980", "FGA")
        discovery = self.query.sql_path.with_name("inspect_collections_profile_codes.sql").read_text(encoding="utf-8")
        rows = [dict(row) for row in self.db.execute(discovery)]
        candidate = next(row for row in rows if row["Attribute Code"] == "FGA")
        self.assertEqual(candidate["Historical CO Students Ever With Attribute"], 1)
        self.assertIn("owner must verify", candidate["Mapping Review"])

    def test_app_export_access_complete_results_dates_and_existing_file(self):
        self.student()
        self.insert('saturn.sgrsatt', 1, '209980', 'FGA')
        cursor = self.db.execute(sqlite_sql(self.query.render_sql()))
        frame = pd.DataFrame.from_records(cursor.fetchall(), columns=[c[0] for c in cursor.description])
        client = MagicMock()
        client.__enter__.return_value = client
        policy = MagicMock()
        with tempfile.TemporaryDirectory() as directory, patch(
            "app.gui.services.insights.load_department_profiles", return_value={"TEST": MagicMock()},
        ), patch("app.gui.services.insights.load_access_configuration", return_value=policy), patch(
            "app.gui.services.insights.get_shared_gui_access_path", return_value=Path(directory) / "access.json",
        ), patch("app.gui.services.insights.build_authenticated_client", return_value=(client, "synthetic")) as build:
            output = Path(directory) / "profile.xlsx"
            context = WorkflowContext("insights_query_export", {
                "action": "query_export", "query_id": self.query.query_id,
                "output_path": str(output)}, WorkflowMode.TEST)
            for administrator, result_frame, success in (
                (False, frame, False),
                (True, frame.assign(**{"Report Row Count": 2}), False),
                (True, frame.drop(columns=["Report Row Count"]), False),
                (True, frame, True),
                (True, frame.iloc[:0], True),
            ):
                with self.subTest(administrator=administrator, rows=len(result_frame), success=success):
                    build.reset_mock()
                    policy.is_administrator.return_value = administrator
                    client.run_sql_file.return_value = result_frame
                    result = run_insights_connection(context)
                    self.assertEqual(result.success, success, result.message)
                    self.assertEqual(output.exists(), success)
                    if not administrator:
                        build.assert_not_called()
                    if success:
                        client.run_sql_file.assert_called_with(self.query.sql_path)
                        book = load_workbook(output)
                        try:
                            headings = [cell.value for cell in book.active[1]]
                            self.assertNotIn("Report Row Count", headings)
                            self.assertIn('First Generation Admissions', headings)
                            self.assertIn('First Generation FAFSA', headings)
                            if len(result_frame):
                                self.assertEqual(book.active.cell(2, headings.index("CWID") + 1).value, "TEST-1")
                                self.assertEqual(book.active.cell(2, headings.index('First Generation Admissions') + 1).value,
                                                 'Yes')
                                self.assertEqual(book.active.cell(2, headings.index('First Generation FAFSA') + 1).value,
                                                 'Unknown')
                                self.assertEqual(book.active.cell(2, headings.index("CO To Date") + 1).value,
                                                 datetime(9999, 12, 31))
                        finally:
                            book.close()
                        original = output.read_bytes()
                        self.assertFalse(run_insights_connection(context).success)
                        self.assertEqual(output.read_bytes(), original)
                        output.unlink()


if __name__ == "__main__":
    unittest.main()
