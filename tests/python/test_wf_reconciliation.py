"""Synthetic acceptance cases for issue #15's financial and matching rules."""
import unittest
import tempfile
from pathlib import Path
from unittest.mock import patch
from decimal import Decimal

import pandas as pd
from openpyxl import load_workbook

from data_processing.wf_mines_balance_recon.reconciliation import (
    add_principal_amount, reconcile_accounts, submission_fraction, load_submission_percentage,
)
import test_collections_tracking as collections_tests
import test_wf_mines_balance_recon as wf_tests
from test_wf_mines_balance_recon import holds_frame, write_frame
from test_wf_agency_reports import agency_file
from app.gui.services.wf_mines_balance_recon import run_wf_mines_balance_recon


class ReconciliationRulesTests(unittest.TestCase):
    def test_fee_only_division_signed_amounts_rounding_and_custom_rate(self):
        frame = pd.DataFrame({
            'Orig. Balance': [125, 125, 125, -125, 'unknown', None, 1.00625],
            'Type': ['Fees', 'No Fees', 'unspecified', 'Fees', 'Fees', 'Fees', 'Fees'],
        })
        result = add_principal_amount(frame, 0.25)
        self.assertEqual(list(result['Principal Amount']), [
            Decimal('100.00'), Decimal('125.00'), Decimal('125.00'), Decimal('-100.00'),
            None, None, Decimal('0.81'),
        ])
        self.assertEqual(list(frame.columns), ['Orig. Balance', 'Type'])
        self.assertEqual(add_principal_amount(frame, 0)['Principal Amount'][0], 125)
        self.assertEqual(add_principal_amount(frame, 1)['Principal Amount'][0], Decimal('62.50'))
        original = pd.DataFrame({'Orig. Balance': [Decimal('1.00625')], 'Type': ['No Fees']})
        self.assertEqual(add_principal_amount(original, 0.25)['Principal Amount'][0], Decimal('1.00625'))
        for invalid in (-0.01, 25, 'NaN', 'Infinity', True, None, 'bad'):
            with self.subTest(value=invalid), self.assertRaises(ValueError):
                submission_fraction(invalid)

    def test_outer_match_sums_accounts_preserves_ids_and_flags_returned(self):
        students = pd.DataFrame([
            ['0000123', 'Synthetic', 'Student', 175],
            ['B-only', 'Banner', 'Only', 20],
        ], columns=['CWID', 'First Name', 'Last Name', 'Total Balance'])
        collections = pd.DataFrame([
            ['0000123', 100], ['0000123', 50], ['W-only', -10], ['123', 5],
        ], columns=['Your Accnt #', 'Principal Amount'])
        returned = pd.DataFrame({'Your Accnt #': ['0000123', '0000123', 'W-only', 'returned-only']})
        result = reconcile_accounts(students, collections, returned).set_index('CWID')
        self.assertEqual(list(result.index), ['0000123', 'B-only', 'W-only', '123'])
        self.assertEqual(result.loc['0000123', 'W&F Principal Amount'], 150)
        self.assertEqual(result.loc['0000123', 'Difference'], 25)
        self.assertEqual(result['W&F Returned'].tolist(), [1, 0, 1, 0])
        self.assertIsNone(result.loc['B-only', 'W&F Principal Amount'])
        self.assertIsNone(result.loc['W-only', 'Banner Total Balance'])
        self.assertIsNone(result.loc['W-only', 'First Name'])
        self.assertTrue(pd.isna(result.loc['W-only', 'Difference']))

    def test_unknown_amounts_poison_aggregate_and_duplicates_or_missing_ids_fail(self):
        students = pd.DataFrame({'CWID': ['S'], 'Total Balance': [100]})
        collections = pd.DataFrame({'Your Accnt #': ['S', 'S'], 'Principal Amount': [50, None]})
        row = reconcile_accounts(students, collections, pd.DataFrame()).iloc[0]
        self.assertIsNone(row['W&F Principal Amount'])
        self.assertIsNone(row['Difference'])
        with self.assertRaisesRegex(ValueError, 'duplicate CWIDs'):
            reconcile_accounts(pd.concat([students, students]), collections, pd.DataFrame())
        for source in ('students', 'collections', 'returned'):
            with self.subTest(source=source), self.assertRaisesRegex(ValueError, 'missing or invalid'):
                reconcile_accounts(
                    pd.DataFrame({'CWID': [None]}) if source == 'students' else students,
                    pd.DataFrame({'Your Accnt #': [''], 'Principal Amount': [5]}) if source == 'collections' else collections,
                    pd.DataFrame({'Your Accnt #': [None]}) if source == 'returned' else pd.DataFrame(),
                )


class CommissionConfigurationTests(unittest.TestCase):
    def test_institution_default_and_admin_configured_fraction(self):
        self.assertEqual(load_submission_percentage(), Decimal('0.25'))
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'commission.json'
            path.write_text('{"submission_percentage": 0.30}')
            self.assertEqual(load_submission_percentage(path), Decimal('0.30'))

    def test_missing_malformed_or_invalid_config_fails(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'commission.json'
            with self.assertRaises(ValueError):
                load_submission_percentage(path)
            for text in ('invalid JSON', '[]', '{}', '{"submission_percentage": true}',
                         '{"submission_percentage": 25}', '{"submission_percentage": -0.1}',
                         '{"submission_percentage": "NaN"}'):
                path.write_text(text)
                with self.subTest(config=text), self.assertRaises(ValueError):
                    load_submission_percentage(path)


class CombinedBannerSQLTests(unittest.TestCase):
    setUp = collections_tests.CollectionsTrackingTests.setUp
    account = collections_tests.CollectionsTrackingTests.account
    transaction = collections_tests.CollectionsTrackingTests.transaction

    def test_to_holds_never_expand_population_or_multiply_balances(self):
        self.account(1)
        self.account(2, code='TO')
        self.account(3)
        self.transaction(1, 100)
        self.db.executescript("""
            INSERT INTO saturn.sprhold VALUES
                (1, 'TO', '2099-01-01', '9999-12-31'),
                (1, ' to ', '2099-02-01', '9999-12-31'),
                (3, 'TO', '2099-01-01', '2099-10-02');
        """)
        frame = pd.read_sql_query(self.query, self.db).set_index('CWID')
        self.assertEqual(list(frame.index), ['TEST-1', 'TEST-3'])
        self.assertEqual(frame.loc['TEST-1', 'TO Hold'], 'TO')
        self.assertIsNone(frame.loc['TEST-3', 'TO Hold'])
        self.assertEqual(frame.loc['TEST-1', 'Current Account Balance'], 100)

    def test_signed_all_history_code_amount_totals_and_unknown_values(self):
        self.db.executescript("INSERT INTO taismgr.tbbdetc VALUES ('WOFF', 'P'), ('BDRC', 'P');")
        self.account(1)
        self.account(2)
        self.transaction(1, 500)
        self.transaction(1, 125, 'WOFF', None, '199910')
        self.transaction(1, -25, 'WOFF', '1999-01-01')
        self.transaction(1, 50, 'BDRC', '2100-01-01')
        self.transaction(2, 100)
        self.transaction(2, None, 'WOFF')
        frame = pd.read_sql_query(self.query, self.db).set_index('CWID')
        row = frame.loc['TEST-1']
        self.assertEqual((row['Current Account Balance'], row['WOFF'], row['BDRC'], row['Total Balance']),
                         (350, 100, 50, 500))
        self.assertTrue(pd.isna(frame.loc['TEST-2', 'WOFF']))
        self.assertTrue(pd.isna(frame.loc['TEST-2', 'Total Balance']))


class ReconciliationWorkbookTests(unittest.TestCase):
    setUp = wf_tests.WFMinesBalanceTests.setUp
    context = wf_tests.WFMinesBalanceTests.context

    def test_real_workbook_matches_fee_principals_returned_and_excludes_entire_unisa_files(self):
        students = holds_frame()
        students['CWID'] = ['0000123']
        students['Total Balance'] = [150]
        write_frame(self.holds, students)
        collections = agency_file(self.root, rows=[
            ['0000123', 'Synthetic', None, 125, 100, 25, None],
            ['0000123', 'Synthetic', None, 62.5, 50, 12.5, None],
            ['W-only', 'Agency only', None, 25, 25, 0, None],
        ])
        returned = agency_file(self.root, returned=True)
        skipped = []
        for code in ('12007', '12009', '12559', '12561'):
            path = self.root / f'malformed-{code}.xlsx'
            path.write_bytes(b'Unisa file must not be parsed or copied')
            skipped.append(path)
        self.params['wf_files'] = (collections, returned, *skipped)
        self.params['submission_percentage'] = 0.99  # Run inputs cannot override institution policy.
        run_wf_mines_balance_recon(self.context())
        book = load_workbook(self.output)
        try:
            self.assertEqual(book.sheetnames, ['Reconciliation', 'Students in Collections', 'W&F Collections', 'W&F Returned'])
            rows = list(book['Reconciliation'].values)
            self.assertEqual(rows[1], ('0000123', 'Synthetic', 'Student', 150, 150, 0, 1))
            self.assertEqual(rows[2], ('W-only', None, None, None, 20, None, 0))
            self.assertEqual(book['W&F Collections']['E2'].value, 100)
            self.assertIn('0.25', book.properties.description)
            self.assertIn('$#,##0.00', book['Reconciliation']['F2'].number_format)
        finally:
            book.close()

    def test_collections_totals_never_reach_workbook_or_reconciliation(self):
        files = []
        for creditor in ('12003', '12004', '12005', '12006'):
            path = agency_file(self.root, creditor)
            book = load_workbook(path)
            book.active['A8'] = 'Total Accounts: 1'
            book.active['D8'] = 99999
            book.save(path)
            book.close()
            files.append(path)
        self.params['wf_files'] = tuple(files)
        run_wf_mines_balance_recon(self.context())
        book = load_workbook(self.output)
        try:
            rows = list(book['W&F Collections'].values)
            self.assertEqual(len(rows), 5)  # Header plus four actual account rows.
            self.assertEqual([row[0] for row in rows[1:]], ['0000123'] * 4)
            recon = {row[0]: row for row in list(book['Reconciliation'].values)[1:]}
            self.assertEqual(set(recon), {'TEST-1', '0000123'})
            self.assertEqual(recon['0000123'][4], 320)
        finally:
            book.close()

    def test_zero_difference_is_numeric_and_missing_difference_has_review_note(self):
        students = holds_frame()
        students['CWID'] = ['0000123']
        students['Total Balance'] = [80]
        write_frame(self.holds, students)
        self.params['wf_files'] = (agency_file(self.root, rows=[
            ['0000123', 'Synthetic', None, 100, 75, 25, None],
            ['W-only', 'Agency only', None, 0, 0, 0, None],
        ], footer=False),)
        run_wf_mines_balance_recon(self.context())
        book = load_workbook(self.output)
        try:
            sheet = book['Reconciliation']
            zero, missing = sheet['F2'], sheet['F3']
            self.assertEqual(zero.value, 0)
            self.assertEqual(zero.data_type, 'n')
            self.assertEqual(zero.number_format.split(';')[2], '$0.00')
            self.assertIsNone(zero.comment)
            self.assertIsNone(missing.value)
            self.assertEqual(missing.fill.fgColor.rgb, '00FFF2CC')
            self.assertIn('not a zero', missing.comment.text)
            self.assertEqual(sheet['E3'].value, 0)
            self.assertEqual(sheet['E3'].number_format.split(';')[2], '$0.00')
        finally:
            book.close()

    def test_workflow_uses_admin_configuration_instead_of_run_parameter(self):
        config = self.root / 'commission.json'
        config.write_text('{"submission_percentage": 0.5}')
        self.params['wf_files'] = (agency_file(self.root),)
        self.params['submission_percentage'] = 0
        with patch('app.gui.services.wf_mines_balance_recon.load_submission_percentage',
                   side_effect=lambda: load_submission_percentage(config)):
            run_wf_mines_balance_recon(self.context())
        book = load_workbook(self.output)
        try:
            self.assertEqual(book['W&F Collections']['E2'].value, 66.67)
            self.assertIn('0.5', book.properties.description)
        finally:
            book.close()

    def test_returned_without_collections_precedes_other_source_tabs(self):
        returned = agency_file(self.root, returned=True)
        self.params['wf_files'] = (self.agency, returned)
        run_wf_mines_balance_recon(self.context())
        book = load_workbook(self.output)
        try:
            self.assertEqual(book.sheetnames[:3], ['Reconciliation', 'Students in Collections', 'W&F Returned'])
            self.assertTrue(book.sheetnames[3].startswith('WF '))
        finally:
            book.close()

    def test_all_excluded_files_still_produce_banner_and_reconciliation(self):
        path = self.root / '12007.xlsx'
        path.write_bytes(b'excluded')
        self.params['wf_files'] = (path,)
        run_wf_mines_balance_recon(self.context())
        book = load_workbook(self.output)
        try:
            self.assertEqual(book.sheetnames, ['Reconciliation', 'Students in Collections'])
            self.assertIsNone(book['Reconciliation']['E2'].value)
        finally:
            book.close()
