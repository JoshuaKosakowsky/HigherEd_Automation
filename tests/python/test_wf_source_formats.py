"""Exercise actual CSV/BIFF XLS/OOXML inputs through the W&F runner."""
import base64
import csv
import json
import tempfile
import unittest
import zlib
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from app.gui.models import WorkflowContext
from app.gui.services.wf_mines_balance_recon import run_wf_mines_balance_recon
from app.gui.workflow_registry import get_workflow
from data_processing.wf_mines_balance_recon.source_workbooks import open_source_workbook
from test_wf_agency_reports import agency_file
from test_wf_mines_balance_recon import holds_frame, activity_frame, write_frame


FIXTURES = Path(__file__).parents[1] / 'fixtures' / 'wf_legacy_xls.json'


class WFSourceFormatTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.holds = self.root / 'holds.xlsx'
        self.activity = self.root / 'activity.xlsx'
        write_frame(self.holds, holds_frame())
        write_frame(self.activity, activity_frame())
        self.params = dict(banner_source='manual', banner_file=self.holds, output_file=self.root / 'review.xlsx')

    def xls(self, key, filename):
        data = json.loads(FIXTURES.read_text())['workbooks'][key]
        path = self.root / filename
        path.write_bytes(zlib.decompress(base64.b64decode(data)))
        self.assertEqual(path.read_bytes()[:8], bytes.fromhex('d0cf11e0a1b11ae1'))
        return path

    def csv(self, original, filename, encoding='utf-8-sig'):
        source = load_workbook(original)
        path = self.root / filename
        try:
            with path.open('w', encoding=encoding, newline='') as handle:
                csv.writer(handle).writerows(source.active.values)
        finally:
            source.close()
        return path

    def run_review(self, files):
        self.params['wf_files'] = tuple(files)
        with patch('app.gui.services.wf_mines_balance_recon.build_authenticated_client') as auth:
            result = run_wf_mines_balance_recon(WorkflowContext('wf_mines_balance_recon', self.params))
        self.assertTrue(result.success, result.message)
        auth.assert_not_called()
        book = load_workbook(self.params['output_file'])
        self.addCleanup(book.close)
        return book

    def test_four_formats_in_one_batch_use_same_named_parser_and_source_provenance(self):
        xlsx = agency_file(self.root, '12003')
        csv_path = self.csv(xlsx, xlsx.stem + '.csv')
        xls = self.xls('collections', xlsx.stem + '.xls')
        xlsm = self.root / (xlsx.stem + '.xlsm')
        xlsm.write_bytes(xlsx.read_bytes())
        originals = {path: path.read_bytes() for path in (csv_path, xls, xlsx, xlsm)}
        book = self.run_review((csv_path, xls, xlsx, xlsm))
        rows = list(book['W&F Collections'].values)
        self.assertEqual(len(rows), 5)
        self.assertEqual({row[-1] for row in rows[1:]}, {path.name for path in originals})
        self.assertTrue(all(row[0] == '0000123' for row in rows[1:]))
        self.assertTrue(all(row[3:7] == (100, 80, 75, 25) for row in rows[1:]))
        self.assertTrue(all(row[-2] == 'Fees' for row in rows[1:]))
        self.assertTrue(all(path.read_bytes() == data for path, data in originals.items()))
        notes_tab = next(name for name in book.sheetnames if name.startswith('WF '))
        self.assertEqual(book[notes_tab]['A2'].value, '=literal source text')
        self.assertEqual(book[notes_tab]['A2'].data_type, 's')

    def test_real_xls_insights_reports_and_returned_source_export_complete_workbook(self):
        legacy = self.xls('holds', 'holds.xls')
        from data_processing.wf_mines_balance_recon.workbook import read_insights_file
        with self.assertRaisesRegex(ValueError, 'exactly one matching students'):
            read_insights_file(legacy, 'students')
        self.params['banner_file'] = self.xls('students', 'combined.xls')
        returned = self.xls('returned', 'Report - ReturnedAccounts- 12003-2099-10-01.xls')
        book = self.run_review((returned,))
        self.assertEqual(book['Students in Collections']['I2'].value, datetime(9999, 12, 31))
        self.assertEqual(book['W&F Returned'].max_row, 2)
        self.assertEqual(book['W&F Returned']['A2'].value, '0000123')
        self.assertEqual(book['W&F Returned']['G2'].value, datetime(2099, 1, 1))

    def test_csv_insights_blanks_and_windows_encoded_agency_text_are_preserved(self):
        source = agency_file(self.root, fees='NO FEES')
        book = load_workbook(source)
        book.active['B7'] = 'Synthetic Caf\u00e9'
        book.save(source)
        book.close()
        agency = self.csv(source, source.stem + '.csv', encoding='cp1252')
        self.params['banner_file'] = self.csv(self.holds, 'holds.csv')
        output = self.run_review((agency,))
        self.assertEqual(output['W&F Collections']['B2'].value, 'Synthetic Caf\u00e9')
        self.assertEqual(output['W&F Collections']['J2'].value, 'No Fees')
        self.assertIsNone(output['Students in Collections']['G2'].value)

    def test_unknown_csv_keeps_leading_zero_ids_and_formula_like_text_literal(self):
        path = self.root / 'demographics.csv'
        path.write_text('Account,Note\n0000123,=literal text\n', encoding='utf-8')
        output = self.run_review((path,))
        sheet = next(sheet for sheet in output if sheet.title.startswith('WF demographics'))
        self.assertEqual(sheet['A2'].value, '0000123')
        self.assertEqual(sheet['B2'].value, '=literal text')
        self.assertEqual(sheet['B2'].data_type, 's')

    def test_utf16_csv_currency_fields_are_numeric_without_losing_identifier_text(self):
        source = agency_file(self.root)
        book = load_workbook(source)
        for row in (7, 8):
            book.active[f'D{row}'] = '$1,000.00'
            book.active[f'E{row}'] = '($25.00)'
            book.active[f'F{row}'] = '1,025.00'
        book.active['B7'] = 'Synthetic, Student'
        book.save(source)
        book.close()
        path = self.csv(source, source.stem + '.csv', encoding='utf-16')
        output = self.run_review((path,))
        sheet = output['W&F Collections']
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet['A2'].value, '0000123')
        self.assertEqual(sheet['B2'].value, 'Synthetic, Student')
        self.assertEqual([sheet.cell(2, column).value for column in (4, 5, 6, 7)], [1000, 800, -25, 1025])
        self.assertTrue(all(sheet.cell(2, column).data_type == 'n' for column in (4, 5, 6)))

    def test_legacy_date_epoch_number_formats_and_source_handles(self):
        path = self.xls('dates_1904', 'dates.xls')
        book = open_source_workbook(path)
        try:
            self.assertEqual(book.active['A2'].value, datetime(2099, 1, 1))
            self.assertEqual(book.active['A2'].number_format, 'yyyy-mm-dd')
        finally:
            book.close()
        path.unlink()  # No XLS file handle may remain open on Windows.

    def test_missing_dependency_invalid_xls_and_gui_filters_have_clear_contracts(self):
        path = self.root / 'invalid.xls'
        path.write_bytes(b'not an Excel workbook')
        with self.assertRaisesRegex(ValueError, 'unable to read legacy XLS'):
            open_source_workbook(path)
        with patch.dict('sys.modules', {'xlrd': None}):
            with self.assertRaisesRegex(ValueError, 'Run setup.ps1'):
                open_source_workbook(path)
        workflow = get_workflow('wf_mines_balance_recon')
        for parameter in workflow.parameters:
            if parameter.key in {'banner_file', 'wf_files'}:
                patterns = ' '.join(pattern for _, pattern in parameter.file_types)
                self.assertEqual(set(patterns.split()), {'*.csv', '*.xls', '*.xlsx', '*.xlsm'})
        help_text = next(item.help_text for item in workflow.parameters if item.key == 'wf_files')
        self.assertNotIn('Convert XLS', help_text)


if __name__ == '__main__':
    unittest.main()
