"""Synthetic financial, source parsing and publication checks for 1305."""

import csv
from io import BytesIO
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
import re
import tempfile
import unittest
from zipfile import ZipFile

import pandas as pd
from openpyxl import Workbook, load_workbook

from data_processing.graduate_contract_recon.export import export_reconciliation
from data_processing.graduate_contract_recon.reconciliation import banner_workday_classification, reconcile
from data_processing.graduate_contract_recon.sources import (
    SourceData, banner_from_frame, fiscal_scope, money, read_source, uploaded_scope,
)
from shared.cancellation import CancellationToken, WorkflowCancelled
from shared.insights.banner_activity import BannerActivityParameters


WD_HEADERS = ["Accounting Date", "Journal Source", "Ledger Account", "Workday Debit Amount",
              "Workday Credit Amount", "Memo", "CWID", "External Reference", "Journal Entry", "Journal Line ID"]
BN_HEADERS = ["'ID'", "'Name'", "'Detail Code'", "'Description'", "'Amount'", "'Balance'", "'Term'",
              "'Aid Year'", "'Feed Document'", "'Feed Date'", "detail_code_description", "Transaction User"]


def workday(rows):
    return SourceData("Workday", "synthetic.csv", WD_HEADERS, list(enumerate(rows, 2)),
        dict(date=0, kind=1, ledger=2, debit=3, credit=4, memo=5, cwid=6, reference=7, journal=8, line=9))


def wd_row(debit=60, credit=0, memo="9001", day="2026-09-30", kind="SIS", cwid="", ext="", line=""):
    return [day, kind, "1305:Graduate Contract Receivable", debit, credit, memo, cwid, ext, "TEST-JOURNAL", line]


def bn_row(amount=60, feed="9001", day="2026-09-30T23:59:59Z", code="TPDT", account="10000001", desc="Synthetic student"):
    return [account, "Synthetic Name", code, desc, amount, 0, "209980", None, feed, day, "Synthetic code", "TEST_USER"]


class GraduateContractReconTests(unittest.TestCase):
    def setUp(self):
        self.params = BannerActivityParameters.from_inputs("2026-09-01", "2026-09-30", "TPDT, Z0LE")

    def build(self, wd, bn):
        return reconcile(workday(wd), banner_from_frame(pd.DataFrame(bn, columns=BN_HEADERS)), self.params)

    def test_displays_summaries_above_details_without_cartesian_duplication(self):
        result = self.build([wd_row(30), wd_row(30), wd_row(5, memo="10000001 Manual", kind="Manual Journal EIB")],
                            [bn_row(40), bn_row(20)])
        self.assertEqual(result.documents[0][12], "Amounts agree")
        self.assertEqual(result.documents[0][10:12], [2, 2])
        self.assertEqual(len(result.combined), 5)
        self.assertEqual([row[4] for row in result.combined],
                         ["Workday", "Workday", "Banner", "Banner", "Workday"])
        self.assertEqual([row[15] for row in result.combined[:2]],
                         ["Summary reference", "Summary reference"])
        self.assertEqual(sum(row[16] for row in result.combined), Decimal("65"))
        self.assertEqual(result.students[0][8:11], [Decimal("65"), Decimal("0"), Decimal("65")])
        controls = dict(result.verification)
        self.assertEqual(controls["Workday debit minus combined activity debit"], 0)
        self.assertEqual(controls["Banner + non-SIS Workday debit minus CWID debit"], 0)
        self.assertEqual(len(result.sources[0].rows), 3)

    def test_zero_net_does_not_hide_debit_or_credit_variance(self):
        result = self.build([wd_row(100, 0), wd_row(0, 100)], [bn_row(0)])
        self.assertEqual(result.documents[0][12], "Amount difference")
        self.assertEqual(result.documents[0][6], 100)
        self.assertEqual(result.documents[0][9], 100)
        self.assertEqual(len(result.combined), 3)
        self.assertEqual([row[15] for row in result.combined if row[4] == "Banner"], ["Review only"])

    def test_feed_blocks_follow_workday_order_and_keep_banner_rows_once(self):
        result = self.build(
            [wd_row(30, memo="9002"), wd_row(10, memo="9001"), wd_row(30, memo="9002")],
            [bn_row(60, feed="9002"), bn_row(99, feed="9003"), bn_row(5, feed="9001")])
        self.assertEqual([(row[4], row[5]) for row in result.combined],
                         [("Workday", 2), ("Workday", 4), ("Banner", 2),
                          ("Workday", 3), ("Banner", 4), ("Banner", 3)])
        self.assertEqual(sum(row[16] for row in result.combined), 70)
        self.assertEqual(result.combined[-1][18], "Missing Workday")

    def test_combined_journal_column_uses_posting_codes_and_never_cwids(self):
        result = self.build(
            [wd_row(memo="123456"), wd_row(memo="10000001", kind="Manual"),
             wd_row(memo="1234567"), wd_row(memo="10000001")],
            [bn_row(feed="123456"), bn_row(feed="1234567"), bn_row(feed="10000001")])
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "review.xlsx"
            export_reconciliation(result, output, app_version="synthetic", banner_source="manual")
            book = load_workbook(output)
            rows = list(book["1305 Combined"].iter_rows(min_row=2, values_only=True))
            self.assertEqual([row[3] for row in rows],
                             ["123456", "123456", "1234567", "1234567"])
            self.assertEqual([row[9] for row in rows], [row[3] for row in rows])
            self.assertEqual(rows[1][0], "10000001")
            self.assertEqual(rows[3][0], "10000001")
            combined = book["1305 Combined"]
            headers = [cell.value for cell in combined[1]]
            self.assertEqual(len(headers), 14)
            for removed in ("Status", "Source Row", "Included in Activity", "Quarter",
                            "Activity Debit", "Activity Credit"):
                self.assertNotIn(removed, headers)
            # Filtering must not shift the shading away from its journal rows.
            self.assertEqual([combined.cell(number, 1).fill.patternType for number in range(2, 6)],
                             [None, None, "solid", "solid"])
            for number in (4, 5):
                self.assertTrue(all(cell.fill.fgColor.rgb == "FFEFEFEF" for cell in combined[number]))
            book.close()

    def test_combined_posting_references_require_six_or_seven_digits(self):
        feeds = ["12345", "260001", "2600001", "250001", "12345678", "26A001", "memo 260001", "=260001"]
        result = self.build([wd_row(memo=feed) for feed in feeds],
                            [bn_row(feed=feed) for feed in feeds])
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "review.xlsx"
            export_reconciliation(result, output, app_version="synthetic", banner_source="manual")
            book = load_workbook(output)
            try:
                rows = list(book["1305 Combined"].iter_rows(min_row=2, values_only=True))
                expected = [feed for feed in ("260001", "2600001", "250001") for _ in range(2)]
                self.assertEqual([row[3] for row in rows], expected)
                self.assertEqual([row[9] for row in rows], expected)
                self.assertEqual(book["1305 Combined"].max_row, 7)
                self.assertEqual(len(result.documents), len(feeds))
                self.assertEqual(len(result.combined), len(feeds) * 2)
                self.assertEqual([values[5] for _, values in result.sources[0].rows], feeds)
                self.assertEqual([values[8] for _, values in result.sources[1].rows], feeds)
            finally:
                book.close()

    def test_missing_zero_amount_sources_never_show_agreement(self):
        result = self.build([wd_row(0, memo="9001")], [bn_row(0, feed="9002")])
        self.assertEqual([row[12] for row in result.documents], ["Missing Banner", "Missing Workday"])
        self.assertTrue(dict(result.verification)["Conclusion"].startswith("OPEN"))

    def test_signed_review_sheets_keep_amounts_notes_and_requested_layout(self):
        result = self.build(
            [wd_row(60), wd_row(5, credit=2, memo="10000001 Manual", kind="Manual")],
            [bn_row(60)])
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "review.xlsx"
            export_reconciliation(result, output, app_version="synthetic", banner_source="manual")
            book = load_workbook(output)
            try:
                doc = book["1305 Doc Recon"]
                self.assertEqual([cell.value for cell in doc[1]],
                    ["Fiscal Year", "Period", "Feed Document", "Workday Total Amt",
                     "Banner Total Amt", "Difference", "Notes"])
                self.assertEqual([cell.value for cell in doc[2]][3:6], [60, 60, 0])
                cwid = book["1305 CWID Recon"]
                self.assertEqual([cell.value for cell in cwid[1]],
                    ["Fiscal Year", "Period", "CWID", "Banner Total Amt", "Workday Total Amt",
                     "Net Period Activity", "Notes / Explanation", "Project", "Grant",
                     "Operating Account", "Ledger Acct", "Spend Cat"])
                self.assertEqual([cell.value for cell in cwid[2]][3:6], [60, 3, 63])
                self.assertEqual([cell.value for cell in cwid[2]][7:10], [None, None, None])
                periods = book["1305 Period Totals"]
                self.assertEqual([cell.value for cell in periods[1]],
                    ["Fiscal Year", "Period", "Quarter", "Workday SIS Amt", "Workday Other Amt",
                     "Banner Amt", "Difference", "Notes"])
                self.assertTrue(all(cell.value in (None, "") for cell in periods["H"][1:]))
                self.assertEqual([cell.value for cell in periods[2]][3:7], [60, 3, 60, -3])
                for ws, columns in ((doc, "DEF"), (cwid, "DEF"), (periods, "DEFG")):
                    for column in columns:
                        self.assertNotEqual(ws[f"{column}2"].number_format, "General")
                self.assertEqual(doc.freeze_panes, "D2")
                self.assertEqual(cwid.freeze_panes, "D2")
            finally:
                book.close()

    def test_combined_filter_keeps_numeric_manual_memos_and_all_other_activity_views(self):
        result = self.build([wd_row(10, memo="250001", kind="Manual", cwid="10000001"),
                             wd_row(20, memo="10000001 adjustment", kind="Manual")],
                            [bn_row(30, feed="12345")])
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "review.xlsx"
            export_reconciliation(result, output, app_version="synthetic", banner_source="manual")
            book = load_workbook(output)
            try:
                combined = book["1305 Combined"]
                self.assertEqual(combined.max_row, 2)
                self.assertEqual(combined["D2"].value, "250001")
                self.assertEqual(combined["J2"].value, "250001")
                self.assertEqual(combined["G2"].value, 10)
                self.assertEqual(book["1305 CWID Recon"]["D2"].value, 30)
                self.assertEqual(book["1305 CWID Recon"]["E2"].value, 30)
                self.assertEqual(book["1305 Period Totals"]["E2"].value, 30)
                self.assertEqual(book["1305 Doc Recon"]["C2"].value, "12345")
                self.assertEqual(book["Workday Data"].max_row, 3)
                self.assertEqual(book["Banner Data"].max_row, 2)
            finally:
                book.close()

    def test_combined_filter_can_produce_a_header_only_sheet(self):
        result = self.build([wd_row()], [bn_row()])
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "review.xlsx"
            export_reconciliation(result, output, app_version="synthetic", banner_source="manual")
            book = load_workbook(output)
            try:
                self.assertEqual(book["1305 Combined"].max_row, 1)
                self.assertEqual(book["1305 Combined"].auto_filter.ref, "A1:N1")
                self.assertEqual(book["1305 Doc Recon"]["E2"].value, 60)
                self.assertEqual(book["1305 Period Totals"]["F2"].value, 60)
            finally:
                book.close()

    def test_reversals_duplicates_and_manual_journals_preserve_raw_signs(self):
        result = self.build([wd_row(0, 25)], [bn_row(-10), bn_row(-10), bn_row(-5)])
        self.assertEqual(result.documents[0][12], "Amounts agree")
        self.assertEqual(len(result.combined), 4)
        self.assertEqual(sum(row[17] for row in result.combined), 25)
        self.assertEqual(result.students[0][10], -25)

    def test_tps00080_sign_is_flipped_before_matching_without_changing_source(self):
        for amount in (Decimal("4000.00"), Decimal("-4000.00"), Decimal("0.00")):
            with self.subTest(amount=amount):
                debit, credit = max(-amount, Decimal("0")), max(amount, Decimal("0"))
                result = self.build([wd_row(debit, credit)], [bn_row(
                    amount, account="TPS00080", desc="10000001 Synthetic sponsor RA")])
                self.assertEqual(result.documents[0][12], "Amounts agree")
                self.assertEqual(result.documents[0][5], debit)
                self.assertEqual(result.documents[0][8], credit)
                self.assertEqual(result.students[0][4:6], [debit, credit])
                self.assertEqual(result.sources[1].rows[0][1][4], amount)
                self.assertEqual(result.combined[1][13:15], [debit, credit])

        regular = self.build([wd_row(60)], [bn_row(60, account="TPS00081",
            desc="10000001 Synthetic sponsor RA")])
        self.assertEqual(regular.documents[0][12], "Amounts agree")

    def test_banner_workday_coding_rules_and_exception_precedence(self):
        cases = [
            ("Synthetic RA", "1631e", "6152"),
            ("Synthetic TA", "1631e", "6153"),
            ("Synthetic GTF", "1631e", "6154"),
            ("Synthetic FELLOW T&F", "5893e", "7503"),
            ("Synthetic Grad Fellowship Living Stipend adjustment", "5893e", "7502"),
            ("Synthetic T&F", "Review", "7503"),
            ("Synthetic unknown", "Review", "Review"),
            ("", "Review", "Review"),
            ("  synthetic ra  ", "1631e", "6152"),
            ("FELLOW T&F Synthetic RA", "1631e", "6152"),
            ("Grad Fellowship Living Stipend T&F", "5893e", "7502"),
        ]
        for description, ledger, spend in cases:
            with self.subTest(description=description):
                self.assertEqual(banner_workday_classification("10000001", description), (ledger, spend))
                self.assertEqual(banner_workday_classification("TPS00080", description),
                                 ("No Ledger Applicable", "No Spend Applicable"))

    def test_cwid_coding_retains_all_distinct_banner_classifications_by_period(self):
        self.params = BannerActivityParameters.from_inputs("2026-09-01", "2026-10-31", "TPDT")
        result = self.build([wd_row(memo="260001"),
                             wd_row(5, kind="Manual", memo="10000002 Synthetic RA")], [
            bn_row(feed="260001", desc="Synthetic RA"),
            bn_row(feed="260001", desc="Synthetic TA"),
            bn_row(feed="260001", desc="Synthetic RA"),
            bn_row(feed="260001", desc="Synthetic FELLOW T&F"),
            bn_row(feed="260001", desc="10000001 Synthetic GTF", account="TPS00080"),
            bn_row(feed="260001", desc="Synthetic unknown"),
            bn_row(feed="260002", day="2026-10-01", desc="Synthetic GTF"),
        ])
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "review.xlsx"
            export_reconciliation(result, output, app_version="synthetic", banner_source="manual")
            book = load_workbook(output)
            try:
                combined = book["1305 Combined"]
                self.assertNotIn("Ledger Acct", [cell.value for cell in combined[1]])
                self.assertNotIn("Spend Cat", [cell.value for cell in combined[1]])
                cwid = book["1305 CWID Recon"]
                self.assertEqual([cell.value for cell in cwid[1]][-5:],
                                 ["Project", "Grant", "Operating Account", "Ledger Acct", "Spend Cat"])
                self.assertEqual(cwid["K2"].value, "1631e; 5893e; No Ledger Applicable; Review")
                self.assertEqual(cwid["L2"].value, "6152; 6153; 7503; No Spend Applicable; Review")
                # A Workday-only student's memo must not invent Banner coding.
                self.assertEqual(cwid["C3"].value, "10000002")
                self.assertEqual(cwid["K3"].value, "Review")
                self.assertEqual(cwid["L3"].value, "Review")
                self.assertEqual(cwid["D2"].value, 240)
                self.assertEqual(cwid["B4"].value, "P04")
                self.assertEqual(cwid["K4"].value, "1631e")
                self.assertEqual(cwid["L4"].value, "6154")
                self.assertEqual(book["Banner Data"].cell(2, 4).value, "Synthetic RA")
            finally:
                book.close()

    def test_signed_export_reversals_and_difference_direction(self):
        result = self.build([wd_row(0, 25), wd_row(0, 7, kind="Manual", memo="10000001 adjustment")],
                            [bn_row(-20)])
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "review.xlsx"
            export_reconciliation(result, output, app_version="synthetic", banner_source="manual")
            book = load_workbook(output)
            try:
                doc = list(book["1305 Doc Recon"].values)[1]
                self.assertEqual(doc[3:6], (-25, -20, 5))
                cwid = list(book["1305 CWID Recon"].values)[1]
                self.assertEqual(cwid[3:6], (-20, -7, -27))
                period = list(book["1305 Period Totals"].values)[1]
                self.assertEqual(period[3:7], (-25, -7, -20, 12))
                self.assertEqual(result.documents[0][12], "Amount difference")
            finally:
                book.close()

    def test_cwid_activity_adds_signed_banner_and_non_sis_workday_amounts(self):
        cases = [
            ("6334.62", "-6334.62", "0.00"),
            ("-6334.62", "6334.62", "0.00"),
            ("6334.62", "-6300.00", "34.62"),
            ("6300.00", "-6334.62", "-34.62"),
            ("25.00", "7.00", "32.00"),
            ("-25.00", "-7.00", "-32.00"),
            ("25.00", None, "25.00"),
            (None, "-7.00", "-7.00"),
        ]
        for banner, workday_amount, expected in cases:
            with self.subTest(banner=banner, workday=workday_amount):
                amount = Decimal(workday_amount or "0")
                wd = [wd_row(max(amount, 0), max(-amount, 0), kind="Manual",
                             memo="10000001 adjustment")] if workday_amount is not None else []
                bn = [bn_row(Decimal(banner))] if banner is not None else []
                result = self.build(wd, bn)
                self.assertEqual(result.students[0][10], Decimal(expected))
                with tempfile.TemporaryDirectory() as folder:
                    output = Path(folder) / "review.xlsx"
                    export_reconciliation(result, output, app_version="synthetic", banner_source="manual")
                    book = load_workbook(output)
                    try:
                        row = list(book["1305 CWID Recon"].values)[1]
                        self.assertEqual(row[3:6], (float(banner or 0), float(workday_amount or 0),
                                                   float(expected)))
                    finally:
                        book.close()

    def test_banner_output_omits_review_columns_and_formats_feed_dates_only(self):
        dates = ["2026-09-30T23:59:59Z", "2026-09-30T23:59:59-06:00",
                 datetime(2026, 9, 30, 15, 45, 12), date(2026, 9, 30), 46295.5,
                 "09/30/2026 11:59:59 PM", None, "invalid", "=TEST_FORMULA"]
        result = self.build([wd_row()], [bn_row(day=value) for value in dates])
        original_rows = [list(values) for _, values in result.sources[1].rows]
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "review.xlsx"
            export_reconciliation(result, output, app_version="synthetic", banner_source="sql")
            book = load_workbook(output)
            try:
                banner = book["Banner Data"]
                headers = [cell.value for cell in banner[1]]
                self.assertEqual(headers, [header for header in BN_HEADERS if header != "'Aid Year'"])
                self.assertEqual(banner.max_row, len(dates) + 1)
                date_column = headers.index("'Feed Date'") + 1
                for row in range(2, 8):
                    cell = banner.cell(row, date_column)
                    self.assertEqual(cell.value, datetime(2026, 9, 30))
                    self.assertEqual(cell.number_format, "yyyy-mm-dd")
                self.assertIsNone(banner.cell(8, date_column).value)
                self.assertEqual(banner.cell(9, date_column).value, "invalid")
                self.assertEqual(banner.cell(10, date_column).value, "=TEST_FORMULA")
                self.assertEqual(banner.cell(10, date_column).data_type, "s")
                self.assertEqual(banner["A2"].value, "10000001")
                self.assertEqual(banner["E2"].value, 60)
                self.assertEqual(banner.auto_filter.ref, "A1:K10")
                self.assertEqual([cell.value for cell in book["Workday Data"][1]],
                                 ["Source Row", *WD_HEADERS, "Selection Status"])
            finally:
                book.close()
        self.assertEqual(result.sources[1].headers, BN_HEADERS)
        self.assertEqual([values for _, values in result.sources[1].rows], original_rows)

    def test_uploaded_banner_output_handles_reordered_unquoted_headers_and_short_rows(self):
        headers = ["Aid Year", "Feed Date", "Amount", "Description", "ID", "Feed Document", "Detail Code"]
        values = ["2099", "2026-09-30T12:34:56Z", 60, "=TEST_FORMULA", "10000001", "260001", "TPDT"]
        with tempfile.TemporaryDirectory() as folder:
            source_path = Path(folder) / "banner.csv"
            with source_path.open("w", newline="", encoding="utf-8") as handle:
                writer = csv.writer(handle)
                writer.writerows([headers, values, ["2099", "invalid"]])
            original = source_path.read_bytes()
            result = reconcile(workday([wd_row(memo="260001")]), read_source(source_path, "Banner"), self.params)
            output = Path(folder) / "review.xlsx"
            export_reconciliation(result, output, app_version="synthetic", banner_source="manual")
            book = load_workbook(output)
            try:
                banner = book["Banner Data"]
                self.assertEqual([cell.value for cell in banner[1]], headers[1:])
                self.assertEqual(banner["A2"].value, datetime(2026, 9, 30))
                self.assertEqual(banner["A2"].number_format, "yyyy-mm-dd")
                self.assertEqual(banner["C2"].value, "=TEST_FORMULA")
                self.assertEqual(banner["C2"].data_type, "s")
                self.assertEqual(list(banner.values)[2], ("invalid", None, None, None, None, None))
            finally:
                book.close()
            self.assertEqual(source_path.read_bytes(), original)

    def test_blank_external_reference_falls_back_to_numeric_memo(self):
        result = self.build([wd_row(memo=9001.0, ext="")], [bn_row(feed="9001.0")])
        self.assertEqual(result.documents[0][12], "Amounts agree")
        result = self.build([wd_row(memo="Unrelated text", ext="9001")], [bn_row()])
        self.assertEqual(result.documents[0][12], "Amounts agree")

    def test_source_invalid_amount_cannot_leave_a_false_matched_group(self):
        result = self.build([wd_row(), wd_row(debit="bad")], [bn_row()])
        self.assertEqual(result.documents[0][12], "Source review required")
        self.assertEqual(result.sources[0].dispositions[3], "Excluded invalid amount")
        self.assertEqual(len(result.combined), 2)
        self.assertTrue(result.issues)

    def test_missing_document_invalid_date_and_empty_account_or_code_remain_open(self):
        missing_ledger = wd_row(); missing_ledger[2] = ""
        for wd, bn in (
            ([wd_row()], [bn_row(feed="")]),
            ([wd_row(), wd_row(day="invalid")], [bn_row()]),
            ([wd_row(), missing_ledger], [bn_row()]),
            ([wd_row()], [bn_row(), bn_row(code="")]),
        ):
            with self.subTest(rows=(len(wd), len(bn))):
                result = self.build(wd, bn)
                self.assertTrue(result.issues)
                self.assertTrue(dict(result.verification)["Conclusion"].startswith("OPEN"))
                self.assertTrue(all(row[12] != "Amounts agree" for row in result.documents))

    def test_account_cwid_fallback_conflicts_and_unallocated_student_controls(self):
        result = self.build([wd_row()], [bn_row(account="TPS-SYNTHETIC", desc="10000002 Synthetic sponsor student")])
        self.assertEqual(result.students[0][3], "10000002")
        result = self.build([wd_row()], [bn_row(desc="10000002 Conflicting synthetic identifier")])
        self.assertFalse(result.students)
        self.assertEqual(dict(result.verification)["Banner + non-SIS Workday debit minus CWID debit"], 60)
        self.assertTrue(any("Conflicting" in issue.issue for issue in result.issues))

    def test_date_code_and_exact_ledger_filters_are_independent(self):
        outside = wd_row(500); outside[2] = "11305 Other"
        result = self.build([wd_row(), outside, wd_row(500, day="2026-10-01")],
            [bn_row(), bn_row(500, day="2026-10-01T00:00:00Z"), bn_row(500, code="HLTH")])
        self.assertEqual(result.documents[0][4:7], [60, 60, 0])
        self.assertEqual(dict(result.verification)["Workday selected rows"], 1)
        self.assertEqual(dict(result.verification)["Banner selected rows"], 1)

    def test_multi_year_periods_quarters_and_reused_feed_are_not_collapsed(self):
        self.params = BannerActivityParameters.from_inputs("2025-06-30", "2026-07-01", "TPDT")
        result = self.build([wd_row(day="2025-06-30"), wd_row(day="2026-07-01")],
                            [bn_row(day="2025-06-30"), bn_row(day="2026-07-01")])
        self.assertEqual(len(result.periods), 14)
        self.assertEqual(result.documents[0][:3], ["FY2025", "P12", "Q4"])
        self.assertEqual(result.documents[1][:3], ["FY2027", "P01", "Q1"])
        self.assertTrue(all(row[12] == "Timing / reused reference review" for row in result.documents))
        self.assertEqual(result.periods[1][-1], "No supplied activity")

    def test_duplicate_workday_identity_blocks_substitution_but_retains_rows(self):
        result = self.build([wd_row(30, line="TEST-LINE"), wd_row(30, line="TEST-LINE")], [bn_row()])
        self.assertEqual(result.documents[0][12], "Source review required")
        self.assertEqual(len(result.sources[0].rows), 2)
        self.assertEqual(len(result.combined), 3)

    def test_fiscal_calendar_and_money_validation(self):
        self.assertEqual(fiscal_scope(date(2026, 6, 30)), (2026, "P12", "Q4"))
        self.assertEqual(fiscal_scope(date(2026, 7, 1)), (2027, "P01", "Q1"))
        self.assertEqual(money("($1,200.50)"), Decimal("-1200.50"))
        for value in ("NaN", "Infinity", "1.001", ""):
            with self.subTest(value=value), self.assertRaises(ValueError):
                money(value)

    def test_export_is_plain_literal_preserves_sources_and_refuses_overwrite_or_cancel(self):
        result = self.build([wd_row(memo="=TEST_FORMULA", kind="Manual", cwid="10000001")],
                            [bn_row(feed="260001")])
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "review.xlsx"
            export_reconciliation(result, output, app_version="synthetic", banner_source="manual")
            book = load_workbook(output)
            self.assertEqual(book.sheetnames[0], "1305 Combined")
            self.assertEqual(len(book.sheetnames), 6)
            self.assertNotIn("1305 Verification", book.sheetnames)
            self.assertNotIn("1305 Exceptions", book.sheetnames)
            self.assertTrue(all(not sheet.tables for sheet in book))
            cell = book["Workday Data"].cell(2, 7)
            self.assertEqual(cell.value, "=TEST_FORMULA")
            self.assertEqual(cell.data_type, "s")
            self.assertEqual([book["1305 Combined"].cell(1, column).value for column in range(1, 5)],
                             ["CWID", "Term", "Recon Period", "Journal Number"])
            self.assertIsInstance(book["1305 Combined"]["G2"].value, (int, float))
            book.close()
            existing = output.read_bytes()
            with self.assertRaises(ValueError):
                export_reconciliation(result, output, app_version="synthetic", banner_source="manual")
            self.assertEqual(output.read_bytes(), existing)
            token = CancellationToken(); token.request()
            cancelled = Path(folder) / "cancelled.xlsx"
            with self.assertRaises(WorkflowCancelled):
                export_reconciliation(result, cancelled, app_version="synthetic", banner_source="manual", cancellation=token)
            self.assertFalse(cancelled.exists())


class GraduateContractSourceTests(unittest.TestCase):
    def test_uploaded_scope_includes_invalid_amount_dates_and_all_valid_codes(self):
        wd = workday([wd_row(debit="bad", day="2026-07-01"), wd_row(day="2026-09-30")])
        outside = wd_row(day="2020-01-01"); outside[2] = "11305 Other"
        wd.rows.append((4, outside))
        bn = banner_from_frame(pd.DataFrame([bn_row(code="HLTH"), bn_row(code="TPDT", day="invalid")], columns=BN_HEADERS))
        scope = uploaded_scope(wd, bn)
        self.assertEqual(scope.start_date, date(2026, 7, 1))
        self.assertEqual(scope.end_date, date(2026, 9, 30))
        self.assertEqual(scope.detail_codes, ("HLTH", "TPDT"))
        result = reconcile(wd, bn, scope)
        self.assertTrue(any(issue.issue.startswith("Invalid amount") for issue in result.issues))
        self.assertTrue(any(issue.issue.startswith("Invalid source date") for issue in result.issues))

    def test_incorrect_worksheet_dimensions_do_not_truncate_headers_or_records(self):
        for source_name, headers, rows in (
            ("Workday", WD_HEADERS, [wd_row(debit=60), wd_row(debit=90)]),
            ("Banner", BN_HEADERS, [bn_row(amount=60), bn_row(amount=90)]),
        ):
            for dimension in ("A1", "A1:L2"):
                with self.subTest(source=source_name, dimension=dimension), tempfile.TemporaryDirectory() as folder:
                    path = Path(folder) / "export.xlsx"
                    book = Workbook()
                    book.active.append(headers)
                    for row in rows:
                        book.active.append(row)
                    book.save(path)
                    book.close()
                    # Reproduce export metadata that hides columns or later rows.
                    original = BytesIO(path.read_bytes())
                    with ZipFile(original) as archive, ZipFile(path, "w") as patched:
                        for member in archive.infolist():
                            content = archive.read(member.filename)
                            if member.filename == "xl/worksheets/sheet1.xml":
                                content = re.sub(rb'<dimension ref="[^"]+"',
                                    f'<dimension ref="{dimension}"'.encode(), content)
                            patched.writestr(member, content)
                    source = read_source(path, source_name)
                    self.assertEqual(source.headers, headers)
                    expected_rows = [[None if value == "" else value for value in row] for row in rows]
                    self.assertEqual(source.rows, list(enumerate(expected_rows, 2)))

    def test_header_based_workday_and_banner_reading_with_title_row_and_quotes(self):
        with tempfile.TemporaryDirectory() as folder:
            wd = Path(folder) / "wd.csv"
            with wd.open("w", newline="") as handle:
                csv.writer(handle).writerows([["Report title"], WD_HEADERS, wd_row()])
            source = read_source(wd, "Workday")
            self.assertEqual(source.rows[0][0], 3)
            bn = Path(folder) / "bn.xlsx"
            book = Workbook(); book.active.title = "Unrelated"
            sheet = book.create_sheet("Downloaded Insights")
            sheet.append(BN_HEADERS); sheet.append(bn_row()); book.save(bn); book.close()
            source = read_source(bn, "Banner")
            self.assertEqual(source.columns["feed"], 8)
            self.assertEqual(source.columns["description"], 3)
            # TGIACCD's additional columns do not change the required mapping.
            book = load_workbook(bn); sheet = book["Downloaded Insights"]
            sheet.cell(1, 13, "'Transaction Number'"); sheet.cell(2, 13, 1)
            book.save(bn); book.close()
            self.assertEqual(read_source(bn, "Banner").columns["transaction"], 12)

    def test_ambiguous_sources_duplicate_headers_or_workday_in_banner_are_rejected(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "bad.xlsx"
            book = Workbook(); book.active.append(BN_HEADERS); book.active.append(bn_row())
            other = book.create_sheet("Also Banner"); other.append(BN_HEADERS); other.append(bn_row())
            book.save(path); book.close()
            with self.assertRaisesRegex(ValueError, "Multiple Banner"):
                read_source(path, "Banner")
            book = Workbook(); book.active.append(WD_HEADERS); book.active.append(wd_row()); book.save(path); book.close()
            with self.assertRaisesRegex(ValueError, "No Banner source"):
                read_source(path, "Banner")
            book = Workbook(); book.active.append(BN_HEADERS + ["Amount"]); book.active.append(bn_row() + [60]); book.save(path); book.close()
            with self.assertRaisesRegex(ValueError, "duplicate"):
                read_source(path, "Banner")


if __name__ == "__main__":
    unittest.main()
