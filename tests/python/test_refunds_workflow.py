from __future__ import annotations

from dataclasses import dataclass, replace
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

import pandas as pd
from openpyxl import load_workbook

from data_processing.refunds import REPORT_COLUMNS, RefundParameters, allocate_refunds
from data_processing.refunds.export import REFUND_SHEETS, WORKBOOK_COLUMNS, export_refund_report
from data_processing.refunds.extract import (
    ExtractSettings,
    _validate_complete_result,
    extract_refund_data,
    read_refund_extracts,
    render_extract_sql,
    render_manual_extract_sql,
)
from data_processing.refunds.ingest import read_refund_download
from data_processing.refunds.terms import derive_target_term, fiscal_year_start, previous_term


TRANSACTION_COLUMNS = [
    "pidm", "term_code", "aidy_code", "tran_number", "detail_code", "amount",
    "stored_balance", "effective_date", "activity_date", "detail_desc", "type_ind",
    "priority", "category_code", "title_iv_ind",
]
CONTEXT_COLUMNS = [
    "pidm", "cwid", "last_name", "first_name", "deceased_ind", "deceased_date",
    "confidential_ind", "account_control_row_count", "refund_hold_count",
    "raw_delinquency_code", "raw_refund_account_ind", "account_control_activity_date",
    "active_ed_row_count", "ed_activity_date", "plus_auth_row_ind", "plus_auth_aidy_code",
    "plus_to_student", "plus_auth_activity_date",
]


@dataclass(frozen=True)
class Transaction:
    detail_code: str
    type_ind: str
    priority: str | None
    amount: int | str | None
    stored_balance: int | str | None = 0
    term: str = "209980"
    aid_year: str | None = "9900"
    effective_date: str | None = "2099-08-01"
    activity_date: str | None = "2099-08-01"
    category: str = "CSH"
    title_iv: str = "N"
    tran_number: int | None = None


class RefundAllocationTests(unittest.TestCase):
    def report(
        self,
        transactions: list[Transaction],
        *,
        authorizations: dict[str, str | list[str]] | None = None,
        refund_hold: bool = False,
        active_ed: bool = True,
        target_term: str = "209980",
        previous_term_override: str | None = None,
        run_date: date = date(2099, 8, 31),
        expected_rows: int = 1,
    ) -> dict[str, object]:
        transaction_rows = []
        occupied: set[int] = set()
        next_number = 1
        for transaction in transactions:
            number = transaction.tran_number
            if number is None:
                while next_number in occupied:
                    next_number += 1
                number = next_number
            occupied.add(number)
            next_number = max(next_number, number + 1)
            transaction_rows.append([
                1, transaction.term, transaction.aid_year, number,
                transaction.detail_code, transaction.amount, transaction.stored_balance,
                transaction.effective_date, transaction.activity_date, "Synthetic source",
                transaction.type_ind, transaction.priority, transaction.category,
                transaction.title_iv,
            ])

        authorization_values = {"9900": "N"} if authorizations is None else authorizations
        context_rows = []
        for aid_year, raw_values in authorization_values.items():
            values = raw_values if isinstance(raw_values, list) else [raw_values]
            for value in values:
                context_rows.append([
                    1, "TEST-1", "Synthetic", "Example", "N", None, "N", 1,
                    1 if refund_hold else 0, "RH" if refund_hold else None, "N",
                    "2099-08-31", 1 if active_ed else 0,
                    "2099-08-31" if active_ed else None, 1, aid_year, value, "2099-08-31",
                ])
        if not context_rows:
            context_rows.append([
                1, "TEST-1", "Synthetic", "Example", "N", None, "N", 1,
                1 if refund_hold else 0, "RH" if refund_hold else None, "N",
                "2099-08-31", 1 if active_ed else 0,
                "2099-08-31" if active_ed else None, 0, None, None, None,
            ])
        result = allocate_refunds(
            pd.DataFrame(transaction_rows, columns=TRANSACTION_COLUMNS),
            pd.DataFrame(context_rows, columns=CONTEXT_COLUMNS),
            RefundParameters(target_term, run_date, previous_term_override),
        )
        self.assertEqual(len(result), expected_rows)
        if expected_rows == 0:
            return {}
        return result.iloc[0].to_dict()

    @staticmethod
    def worked_example() -> list[Transaction]:
        return [
            Transaction("TUIN", "C", "899", 6000),
            Transaction("FEES", "C", "897", 1000),
            Transaction("OTHR", "C", "700", 500),
            Transaction("PAYA", "P", "899", 2000),
            Transaction("PAYB", "P", "890", 1500),
            Transaction("FDPL", "P", "800", 4000, -500, category="FA", title_iv="Y"),
            Transaction("PAYC", "P", "000", 2000, -1500),
        ]

    def assert_split(self, row: dict[str, object], parent: str, student: str) -> None:
        self.assertEqual(row["parent_refund_amount"], Decimal(parent))
        self.assertEqual(row["student_refund_amount"], Decimal(student))
        self.assertEqual(row["total_refund_amount"], Decimal(parent) + Decimal(student))

    def test_worked_example_and_output_contract(self) -> None:
        row = self.report(self.worked_example())
        self.assert_split(row, "500.00", "1500.00")
        self.assertEqual(list(row), REPORT_COLUMNS)
        self.assertEqual(row["review_status"], "READY_FOR_STAFF_REVIEW")
        self.assertIn("FDPL", str(row["balance_sources"]))
        self.assertNotIn("PAYA", str(row["balance_sources"]))

    def test_cross_fy_title_iv_uses_200_give_and_receive_cap(self) -> None:
        row = self.report([
            Transaction("OLD1", "C", "899", 150, term="209780"),
            Transaction("OLD2", "C", "899", 150, term="209880"),
            Transaction("TIVA", "P", "000", 500, -300, category="FA", title_iv="Y"),
        ])
        self.assert_split(row, "0.00", "300.00")
        self.assertEqual(row["total_unused_payment_amount"], Decimal("300.00"))
        self.assertEqual(row["unpaid_charge_amount"], Decimal("100.00"))
        self.assertEqual(row["title_iv_applied_to_older_fiscal_years"], Decimal("200.00"))
        self.assertEqual(row["review_status"], "MANUAL_REVIEW")

        destination = self.report([
            Transaction("OLD", "C", "899", 500, term="209780"),
            Transaction("TIVA", "P", "000", 300, -100, term="209880", aid_year="9899", category="FA", title_iv="Y"),
            Transaction("TIVB", "P", "000", 300, -300, category="FA", title_iv="Y"),
        ])
        self.assert_split(destination, "0.00", "400.00")
        self.assertEqual(destination["total_unused_payment_amount"], Decimal("400.00"))
        self.assertEqual(destination["unpaid_charge_amount"], Decimal("300.00"))

    def test_same_fy_title_iv_and_unrestricted_cross_fy_are_not_capped(self) -> None:
        same_fy = self.report([
            Transaction("OLD", "C", "899", 1000, term="209980"),
            Transaction("TIVA", "P", "000", 1200, -200, term="210010", category="FA", title_iv="Y"),
        ], target_term="210010", run_date=date(2100, 1, 31))
        self.assert_split(same_fy, "0.00", "200.00")
        self.assertEqual(same_fy["unpaid_charge_amount"], Decimal("0.00"))

        unrestricted = self.report([
            Transaction("OLD", "C", "899", 1000, term="209955"),
            Transaction("FREE", "P", "000", 1200, -200, category="FA", title_iv="N"),
        ])
        self.assert_split(unrestricted, "0.00", "200.00")
        self.assertEqual(unrestricted["unrestricted_applied_to_older_terms"], Decimal("1000.00"))

    def test_artificial_priority_bands_and_earliest_tie_breaker(self) -> None:
        first_800 = self.report([
            Transaction("CHG8", "C", "899", 100, tran_number=10),
            Transaction("FDPL", "P", "800", 100, -100, category="FA", title_iv="Y", tran_number=2),
            Transaction("TPDT", "P", "800", 100, tran_number=3),
        ])
        self.assert_split(first_800, "100.00", "0.00")

        early_fdpl = self.report([
            Transaction("CHG8", "C", "899", 100, tran_number=10),
            Transaction("FDPL", "P", "800", 100, category="FA", title_iv="Y", tran_number=1),
            Transaction("OT8", "P", "800", 100, -100, tran_number=2),
        ])
        self.assert_split(early_fdpl, "0.00", "100.00")

        regular_000 = self.report([
            Transaction("CHG", "C", "700", 100),
            Transaction("ACHK", "P", "000", 100, -100),
            Transaction("PAY0", "P", "000", 100),
        ])
        self.assertIn("PAY0", str(regular_000["balance_sources"]))
        self.assertNotIn("ACHK", str(regular_000["balance_sources"]))
        self.assertEqual(regular_000["proposed_student_delivery"], "ARFD (System)")

        first_000 = self.report([
            Transaction("CHG", "C", "700", 100),
            Transaction("PAY0", "P", "000", 100, -100),
            Transaction("COFP", "P", "000", 100),
        ])
        self.assertIn("PAY0", str(first_000["balance_sources"]))
        self.assertNotIn("COFP", str(first_000["balance_sources"]))

    def test_positional_wildcards_preserve_descending_charge_order(self) -> None:
        row = self.report([
            Transaction("C897", "C", "897", 100, tran_number=1),
            Transaction("C899", "C", "899", 100, tran_number=2),
            Transaction("P899", "P", "899", 100, tran_number=3),
            Transaction("P890", "P", "890", 100, tran_number=4),
            Transaction("FDPL", "P", "800", 100, -100, category="FA", title_iv="Y", tran_number=5),
        ])
        self.assert_split(row, "100.00", "0.00")
        self.assertEqual(row["unpaid_charge_amount"], Decimal("0.00"))

    def test_allocation_preserves_exact_cents(self) -> None:
        row = self.report([
            Transaction("CHG1", "C", "899", "0.10"),
            Transaction("CHG2", "C", "899", "0.20"),
            Transaction("FDPL", "P", "800", "0.35", "-0.05", category="FA", title_iv="Y"),
        ])
        self.assert_split(row, "0.05", "0.00")

    def test_reversals_and_cross_detail_charge_credits_net_correctly(self) -> None:
        reversal = self.report([
            Transaction("ACHK", "P", "000", 1000, -600, tran_number=1),
            Transaction("ACHK", "P", "000", -400, 0, tran_number=2),
        ])
        self.assert_split(reversal, "0.00", "600.00")
        self.assertEqual(reversal["negative_source_count"], 1)

        charges = self.report([
            Transaction("HLTH", "C", "879", 1589),
            Transaction("HIWR", "C", "879", -1589),
            Transaction("PAY0", "P", "000", 100, -100),
        ])
        self.assert_split(charges, "0.00", "100.00")
        self.assertEqual(charges["unpaid_charge_amount"], Decimal("0.00"))

    def test_tricky_cross_term_charge_credits_do_not_create_false_debt(self) -> None:
        row = self.report([
            Transaction("HLTH", "C", "879", "1111", term="209780"),
            Transaction("HIWR", "C", "879", "-1111", term="209780"),
            Transaction("HLTH", "C", "879", "2222", term="209880"),
            Transaction("HIWR", "C", "879", "-2222", term="209880"),
            Transaction("OLD", "C", "899", "4750.25", term="209955"),
            Transaction("HLTH", "C", "879", "3333", term="209955"),
            Transaction("HIWR", "C", "879", "-3333", term="209955"),
            Transaction("TUIN", "C", "899", "9000", tran_number=20),
            Transaction("FEES", "C", "897", "1500", tran_number=21),
            Transaction("HOUS", "C", "889", "2500", tran_number=22),
            Transaction("HLTH", "C", "879", "1600", tran_number=23),
            Transaction("HIWR", "C", "879", "-1600", tran_number=24),
            Transaction("FDPL", "P", "800", "12000", tran_number=25, category="PPL", title_iv="Y"),
            Transaction("FDSL", "P", "800", "1500", tran_number=26, category="FAL", title_iv="Y"),
            Transaction("FDUL", "P", "800", "500", tran_number=27, category="FAL", title_iv="Y"),
            Transaction("PELL", "P", "800", "1500", tran_number=28, category="FAG", title_iv="Y"),
            Transaction("SCHP", "P", "000", "6000", tran_number=29, category="FAS"),
        ])
        self.assert_split(row, "0.00", "3749.75")
        self.assertEqual(row["unpaid_charge_amount"], Decimal("0.00"))
        self.assertEqual(row["title_iv_applied_to_older_fiscal_years"], Decimal("200.00"))
        self.assertEqual(row["unrestricted_applied_to_older_terms"], Decimal("4550.25"))

    def test_historical_credit_reconstruction_uses_priority_eligibility(self) -> None:
        row = self.report([
            Transaction("OLD7", "C", "700", 100, term="209955"),
            Transaction("P899", "P", "899", 100, -100, term="209955"),
            Transaction("P000", "P", "000", 100, term="209955"),
        ])
        self.assert_split(row, "0.00", "100.00")
        self.assertIn("P899", str(row["balance_sources"]))
        self.assertNotIn("P000", str(row["balance_sources"]))

    def test_positive_account_balance_is_excluded_despite_restricted_funds(self) -> None:
        self.report([
            Transaction("OLD", "C", "899", 1500, 1300, term="209955"),
            Transaction("TIVA", "P", "000", 1200, -1000, category="FA", title_iv="Y"),
        ], expected_rows=0)

    def test_zero_balance_restricted_parent_plus_keeps_split_and_review(self) -> None:
        row = self.report([
            Transaction("OLD7", "C", "700", "4375.00", term="209880"),
            Transaction(
                "FDPL", "P", "800", "4375.00", "-4375.00",
                category="FA", title_iv="Y",
            ),
        ])
        self.assertEqual(row["full_account_balance"], Decimal("0.00"))
        self.assert_split(row, "4375.00", "0.00")
        self.assertEqual(row["unused_fdpl_amount"], Decimal("4375.00"))
        self.assertEqual(row["unpaid_charge_amount"], Decimal("4375.00"))
        self.assertEqual(row["proposed_parent_delivery"], "RFDP")
        self.assertEqual(row["refund_split_status"], "CALCULATED_SUBJECT_TO_REVIEW")
        self.assertEqual(row["review_status"], "MANUAL_REVIEW")
        self.assertIn(
            "POLICY_REFUND_DIFFERS_FROM_FULL_ACCOUNT_CREDIT",
            str(row["review_reasons"]),
        )

    def test_settled_old_term_is_not_reopened_by_current_priorities(self) -> None:
        row = self.report([
            Transaction("OLD7", "C", "700", 100, term="209880", tran_number=1),
            Transaction("P899", "P", "899", 100, term="209880", tran_number=2),
            Transaction("P000", "P", "000", 100, tran_number=3),
        ])
        self.assert_split(row, "0.00", "100.00")
        self.assertIn("P000", str(row["balance_sources"]))
        self.assertNotIn("P899", str(row["balance_sources"]))

    def test_exact_priority_refund_on_zero_account_in_current_or_open_old_term(self) -> None:
        for term in ("209980", "209880"):
            with self.subTest(term=term):
                row = self.report([
                    Transaction("TUIN", "C", "899", 10000, 500, term=term),
                    Transaction("FEE", "C", "897", 1000, 0, term=term),
                    Transaction("P899", "P", "899", 9500, 0, term=term),
                    Transaction("P897", "P", "897", 1500, -500, term=term),
                ])
                self.assert_split(row, "0.00", "500.00")
                self.assertEqual(row["full_account_balance"], Decimal("0.00"))
                self.assertEqual(row["unpaid_charge_amount"], Decimal("500.00"))
                self.assertEqual(row["proposed_student_delivery"], "ARFD (System)")
                self.assertIn("RESTRICTED_PAYMENT_REFUND_WITH_UNPAID_CHARGE", row["review_reasons"])

    def test_exact_match_remainders_and_cross_term_title_iv_limits(self) -> None:
        for priority in ("899", "869"):
            row = self.report([
                Transaction("EXCT", "C", priority, 800),
                Transaction("PAYX", "P", priority, 1000, -200),
            ])
            self.assert_split(row, "0.00", "200.00")
        for title_iv, term, target, expected, unpaid in (
            ("Y", "209980", "210010", "200.00", "0.00"),
            ("Y", "209955", "209980", "800.00", "600.00"),
            ("N", "209955", "209980", "200.00", "0.00"),
        ):
            with self.subTest(title_iv=title_iv, term=term):
                row = self.report([
                    Transaction("EXCT", "C", "869", 800, term=term),
                    Transaction("PAYX", "P", "869", 1000, -200, term=target,
                                title_iv=title_iv, category="FA"),
                ], target_term=target)
                self.assert_split(row, "0.00", expected)
                self.assertEqual(row["unpaid_charge_amount"], Decimal(unpaid))

    def test_card_and_ach_remainders_route_only_their_own_funds(self) -> None:
        for code in ("ACHK", "CRAM", "CRDS", "CRMC", "CRVC"):
            with self.subTest(code=code):
                row = self.report([
                    Transaction("CHG7", "C", "700", 400, tran_number=1),
                    Transaction("COFP", "P", "000", 100, tran_number=5),
                    Transaction(code, "P", "000", 600, -300, tran_number=2),
                    Transaction("SCHP", "P", "000", 1000, -1000, tran_number=3, category="FAS"),
                    Transaction("TIVA", "P", "800", 500, -500, tran_number=4, title_iv="Y", category="FA"),
                ])
                self.assert_split(row, "0.00", "1800.00")
                route = "AFRD" if code == "ACHK" else code
                self.assertEqual(row["proposed_student_delivery"], f"{route} (Transact) 300.00; ARFD (System) 1500.00")
                self.assertEqual(row["original_payment_total"], Decimal("300.00"))
                self.assertNotIn("000Z", row["balance_sources"])

    def test_cards_immediate_ach_day_16_and_card_reversals(self) -> None:
        for code in ("CRAM", "CRDS", "CRMC", "CRVC"):
            row = self.report([
                Transaction(code, "P", "000", 100, effective_date="2099-08-31"),
                Transaction(code, "P", "000", -40, effective_date="2099-08-31"),
            ])
            self.assert_split(row, "0.00", "60.00")
            self.assertEqual(row["proposed_student_delivery"], f"{code} (Transact)")
        ready = self.report([Transaction("ACHK", "P", "000", 50, -50, effective_date="2099-08-15")])
        self.assertEqual(ready["proposed_student_delivery"], "AFRD (Transact)")

    def test_c529_and_z0le_require_current_term_unused_payment(self) -> None:
        for code in ("C529", "Z0LE"):
            for term, charge, reversal, flag in (
                ("209980", 80, 0, True),
                ("209980", 100, 0, False),
                ("209980", 0, -100, False),
                ("209880", 0, 0, False),
            ):
                with self.subTest(code=code, term=term, charge=charge, reversal=reversal):
                    row = self.report([
                        Transaction("CHG8", "C", "899", charge, term=term),
                        Transaction(code, "P", "000", 100, term=term),
                        Transaction(code, "P", "000", reversal, term=term),
                        Transaction("FREE", "P", "000", 50, -50),
                    ])
                    self.assertEqual(row["third_party_review_required_ind"], "Y" if flag else "N")
                    self.assertEqual("Possible Third Party refund" in (row["review_reasons"] or ""), flag)
                    if flag:
                        self.assertEqual(row["third_party_match_source"], code)
                        self.assertEqual(row["proposed_student_delivery"], "THIRD_PARTY_REVIEW")
                        self.assertEqual(row["review_status"], "MANUAL_REVIEW")
                        self.assertIsNotNone(row["student_refund_amount"])

    def test_tppy_review_uses_surviving_payment_age_even_when_applied(self) -> None:
        run_date = date(2099, 8, 31)
        for term, age, reversal, flag in (
            ("209980", 90, 0, True),
            ("209955", 0, 0, True),
            ("209955", 32, 0, True),
            ("209955", 33, 0, False),
            ("209955", -1, 0, False),
            ("209955", None, 0, False),
            ("209980", None, 0, True),
            ("209980", 0, -100, False),
            ("209955", 32, -50, True),
        ):
            with self.subTest(term=term, age=age, reversal=reversal):
                effective = (run_date - timedelta(days=age)).isoformat() if age is not None else None
                row = self.report([
                    Transaction("CHG8", "C", "899", 100 + reversal, term=term),
                    Transaction("TPPY", "P", "800", 100, 0, term=term,
                                effective_date=effective),
                    Transaction("TPPY", "P", "800", reversal, 0, term=term),
                    Transaction("FREE", "P", "000", 50, -50),
                ])
                self.assert_split(row, "0.00", "50.00")
                self.assertEqual(row["third_party_review_required_ind"], "Y" if flag else "N")
                self.assertEqual("Possible Third Party refund" in (row["review_reasons"] or ""), flag)
                if flag:
                    self.assertEqual(row["third_party_match_source"], "TPPY")
                    self.assertEqual(row["proposed_student_delivery"], "THIRD_PARTY_REVIEW")
                    self.assertEqual(row["review_status"], "MANUAL_REVIEW")
                else:
                    self.assertEqual(row["proposed_student_delivery"], "ARFD (System)")

        held = self.report([
            Transaction("TPPY", "P", "800", 50, -50),
        ], refund_hold=True)
        self.assertEqual(held["third_party_review_required_ind"], "Y")
        self.assertEqual(held["proposed_student_delivery"], "Refund Hold - Student")
        self.assertEqual(held["review_status"], "HOLD")

    def test_homp_review_uses_surviving_charge_age_even_in_settled_history(self) -> None:
        run_date = date(2099, 8, 31)
        for term, age, reversal, flag in (
            ("209980", 90, 0, True),
            ("209955", 0, 0, True),
            ("209955", 32, 0, True),
            ("209955", 33, 0, False),
            ("209955", -1, 0, False),
            ("209955", None, 0, False),
            ("209980", None, 0, True),
            ("209980", 0, -100, False),
            ("209955", 32, -50, True),
        ):
            with self.subTest(term=term, age=age, reversal=reversal):
                effective = (run_date - timedelta(days=age)).isoformat() if age is not None else None
                row = self.report([
                    Transaction("HOMP", "C", "889", 100, 0, term=term, effective_date=effective),
                    Transaction("HOMP", "C", "889", reversal, 0, term=term),
                    Transaction("PAID", "P", "000", 100 + reversal, 0, term=term),
                    Transaction("FREE", "P", "000", 50, -50),
                ])
                self.assert_split(row, "0.00", "50.00")
                self.assertEqual("Mines Park Charge - Review" in (row["review_reasons"] or ""), flag)
                self.assertEqual(row["proposed_student_delivery"], "ARFD (System)")
                if flag:
                    self.assertEqual(row["review_status"], "MANUAL_REVIEW")

    def test_reversed_recent_homp_does_not_rejuvenate_old_surviving_charge(self) -> None:
        row = self.report([
            Transaction("HOMP", "C", "889", 100, term="209955", effective_date="2099-06-01"),
            Transaction("HOMP", "C", "889", 100, term="209955", effective_date="2099-08-20"),
            Transaction("HOMP", "C", "889", -100, term="209955", effective_date="2099-08-21"),
            Transaction("PAID", "P", "000", 100, term="209955"),
            Transaction("FREE", "P", "000", 50, -50),
        ])
        self.assert_split(row, "0.00", "50.00")
        self.assertNotIn("Mines Park Charge - Review", row["review_reasons"] or "")

    def test_authorized_plus_and_card_remainders_keep_distinct_student_routes(self) -> None:
        row = self.report([
            Transaction("FDPL", "P", "800", 100, -100, category="FA", title_iv="Y"),
            Transaction("CRMC", "P", "000", 50, -50),
        ], authorizations={"9900": "Y"})
        self.assert_split(row, "0.00", "150.00")
        self.assertEqual(row["proposed_student_delivery"], "CRMC (Transact) 50.00; ARFD (System) 100.00")

    def test_third_party_review_preserves_parent_amount_but_suppresses_delivery(self) -> None:
        row = self.report([
            Transaction("FDPL", "P", "800", 100, -100, category="FA", title_iv="Y"),
            Transaction("TPPY", "P", "800", 50, -50),
        ])
        self.assert_split(row, "100.00", "50.00")
        self.assertEqual(row["proposed_parent_delivery"], "THIRD_PARTY_REVIEW")
        self.assertEqual(row["proposed_student_delivery"], "THIRD_PARTY_REVIEW")

    def test_export_routes_actual_mixed_allocator_deliveries(self) -> None:
        row = self.report([
            Transaction("FDPL", "P", "800", 70, -70, category="FA", title_iv="Y"),
            Transaction("ACHK", "P", "000", 20, -20),
            Transaction("ACHK", "P", "000", 30, -30, effective_date="2099-08-30"),
            Transaction("ACHK", "P", "000", 40, -40, effective_date="2098-08-01"),
            Transaction("ACHK", "P", "000", 50, -50, effective_date=None),
            *(Transaction(code, "P", "000", 10, -10) for code in ("CRAM", "CRDS", "CRMC", "CRVC")),
            Transaction("SCHP", "P", "000", 100, -100, category="FA"),
        ])
        with TemporaryDirectory() as directory:
            output = Path(directory) / "refunds.xlsx"
            export_refund_report(pd.DataFrame([row]), output)
            sheets = pd.read_excel(output, sheet_name=None)
            expected = {"Transact Refunds": 60, "ACH Clearing": 30, "ACH Reviews": 90,
                        "System Refunds": 100, "Parent Refunds": 70}
            self.assertEqual({name: frame.tab_refund_amount.sum() for name, frame in sheets.items()
                              if not frame.empty}, expected)
            self.assertIn("09/15/2099", sheets["ACH Clearing"].iloc[0].tab_delivery)
            for code in ("CRAM", "CRDS", "CRMC", "CRVC"):
                self.assertIn(f"{code} (Transact) 10.00", sheets["Transact Refunds"].iloc[0].tab_delivery)

    def test_closed_cross_fy_history_and_posted_current_refund_are_omitted(self) -> None:
        self.report([
            Transaction(
                "FDPL", "P", "800", 600, 0, term="209810",
                aid_year="9798", category="FA", title_iv="Y",
            ),
            Transaction("LATE", "C", "700", 1320, 0, term="209855"),
            Transaction("PAY0", "P", "000", 720, 0, term="209880"),
            Transaction("FDSL", "P", "800", 100, 0, category="FA", title_iv="Y"),
            Transaction("ARFD", "C", "800", 100, 0),
        ], expected_rows=0)

    def test_bad_inputs_block_split_without_displaying_zero(self) -> None:
        missing_amount = self.report(
            self.worked_example() + [Transaction("NULL", "C", "700", None)]
        )
        self.assertIsNone(missing_amount["student_refund_amount"])
        self.assertIsNone(missing_amount["parent_refund_amount"])
        self.assertIn("MISSING_TRANSACTION_AMOUNT", str(missing_amount["review_reasons"]))

        future = self.report(
            self.worked_example()
            + [Transaction("FUTR", "C", "700", 1, term="210010")]
        )
        self.assertIsNone(future["student_refund_amount"])
        self.assertIn(
            "FUTURE_TERM_ACTIVITY_EXCLUDED_FROM_ALLOCATION",
            str(future["review_reasons"]),
        )

    def test_stored_balance_is_diagnostic_only(self) -> None:
        transactions = self.worked_example()
        transactions[-1] = Transaction("PAYC", "P", "000", 2000, 0)
        row = self.report(transactions)
        self.assert_split(row, "500.00", "1500.00")
        self.assertEqual(row["allocation_review_required_ind"], "Y")
        self.assertIn(
            "STORED_BALANCE_DIFFERS_FROM_RECONSTRUCTED_ALLOCATION",
            str(row["review_reasons"]),
        )

    def test_fdpl_authorization_is_applied_by_aid_year(self) -> None:
        mixed = self.report([
            Transaction("FDPL", "P", "800", 300, -300, term="209880", aid_year="9899", category="FA", title_iv="Y"),
            Transaction("FDPL", "P", "800", 300, -300, aid_year="9900", category="FA", title_iv="Y"),
        ], authorizations={"9899": "Y", "9900": "N"})
        self.assert_split(mixed, "300.00", "300.00")
        self.assertEqual(mixed["plus_to_student_status"], "MIXED")
        self.assertEqual(mixed["proposed_parent_delivery"], "RFDP")

        missing = self.report([
            Transaction("FDPL", "P", "800", 100, -100, category="FA", title_iv="Y")
        ], authorizations={})
        self.assertIsNone(missing["parent_refund_amount"])
        self.assertIn("PLUS_AUTH_RECORD_MISSING", str(missing["review_reasons"]))

    def test_ach_boundaries_hold_and_standard_delivery(self) -> None:
        waiting = self.report([
            Transaction("ACHK", "P", "000", 50, -50, effective_date="2099-08-16")
        ])
        self.assertEqual(waiting["proposed_student_delivery"], "ACHK Clearing Wait until 09/01/2099")
        self.assertEqual(waiting["review_status"], "WAIT_ACH_CLEARING")

        old = self.report([
            Transaction("ACHK", "P", "000", 50, -50, effective_date="2099-03-03")
        ])
        self.assertEqual(old["proposed_student_delivery"], "AFRD (Transact) - May Be Too Old")

        check = self.report([Transaction("PAY0", "P", "000", 50, -50)], active_ed=False)
        self.assertEqual(check["proposed_student_delivery"], "RFND (CHECK)")
        hold = self.report([Transaction("PAY0", "P", "000", 50, -50)], refund_hold=True)
        self.assertEqual(hold["proposed_student_delivery"], "Refund Hold - Student")

    def test_chck_waits_through_day_15_without_changing_delivery_route(self) -> None:
        ready = self.report([
            Transaction("CHCK", "P", "000", 50, -50, effective_date="2099-08-15")
        ])
        self.assertEqual(ready["proposed_student_delivery"], "ARFD (System)")
        self.assertEqual(ready["review_status"], "READY_FOR_STAFF_REVIEW")
        self.assertNotIn("CHCK_CLEARING_WAIT", ready["review_reasons"] or "")

        waiting = self.report([
            Transaction("CHCK", "P", "000", 30, -30, effective_date="2099-08-16"),
            Transaction("CHCK", "P", "000", 50, -50, effective_date="2099-08-20"),
        ])
        self.assertEqual(waiting["proposed_student_delivery"], "ARFD (System)")
        self.assertEqual(waiting["review_status"], "WAIT_CHECK_CLEARING")
        self.assertIn(
            "CHCK_CLEARING_WAIT_UNTIL_09/05/2099_AMOUNT_80.00",
            waiting["review_reasons"],
        )

        check = self.report([
            Transaction("CHCK", "P", "000", 50, -50, effective_date="2099-08-16")
        ], active_ed=False)
        self.assertEqual(check["proposed_student_delivery"], "RFND (CHECK)")
        self.assertEqual(check["review_status"], "WAIT_CHECK_CLEARING")

        held = self.report([
            Transaction("CHCK", "P", "000", 50, -50, effective_date="2099-08-16")
        ], refund_hold=True)
        self.assertEqual(held["proposed_student_delivery"], "Refund Hold - Student")
        self.assertEqual(held["review_status"], "HOLD")
        self.assertIn("CHCK_CLEARING_WAIT_UNTIL_09/01/2099", held["review_reasons"])

        invalid_date = self.report([
            Transaction("CHCK", "P", "000", 50, -50, effective_date=None)
        ])
        self.assertEqual(invalid_date["proposed_student_delivery"], "ARFD (System)")
        self.assertEqual(invalid_date["review_status"], "MANUAL_REVIEW")
        self.assertIn("CHCK_EFFECTIVE_DATE_MISSING_OR_FUTURE", invalid_date["review_reasons"])

    def test_rh_hold_overrides_mines_park_and_parent_delivery(self) -> None:
        row = self.report([
            Transaction("HOMP", "C", "889", 100),
            Transaction("FDPL", "P", "800", 200, -100, category="FA", title_iv="Y"),
        ], authorizations={"9900": "N"}, refund_hold=True)
        self.assert_split(row, "100.00", "0.00")
        self.assertEqual(row["proposed_student_delivery"], "NONE")
        self.assertEqual(row["proposed_parent_delivery"], "Refund Hold - Parent")
        self.assertEqual(row["review_status"], "HOLD")
        self.assertIn("REFUND_HOLD_RH", row["review_reasons"])
        self.assertIn("Mines Park Charge - Review", row["review_reasons"])

    def test_activity_audit_columns_retain_time_of_day(self) -> None:
        row = self.report([
            Transaction(
                "PAY0", "P", "000", 50, -50,
                activity_date="2099-08-01T23:07:09",
            )
        ])
        self.assertEqual(
            row["last_ar_activity_date"],
            datetime(2099, 8, 1, 23, 7, 9),
        )


class RefundTermTests(unittest.TestCase):
    def test_mines_term_and_fiscal_year_rules(self) -> None:
        self.assertEqual(derive_target_term(date(2026, 5, 15)), "202610")
        self.assertEqual(derive_target_term(date(2026, 5, 16)), "202655")
        self.assertEqual(derive_target_term(date(2026, 7, 16)), "202680")
        self.assertEqual(fiscal_year_start("202680"), 2026)
        for term in ("202710", "202750", "202755", "202760"):
            self.assertEqual(fiscal_year_start(term), 2026)
        self.assertEqual(previous_term("202680", "202660"), "202660")


class RefundExtractTests(unittest.TestCase):
    def test_template_is_rendered_only_from_validated_values(self) -> None:
        settings = ExtractSettings("202680", 20, Path("unused"))
        rendered = render_extract_sql(
            "WITH s AS (__REFUND_SCOPE_SQL__) SELECT __BATCH_COUNT__, __BATCH_INDEX__",
            settings,
            3,
        )
        self.assertIn("CAST('202680' AS varchar(6)) AS target_term_override", rendered)
        self.assertTrue(rendered.endswith("SELECT 20, 3"))
        with self.assertRaises(ValueError):
            ExtractSettings("202680", 1, Path("unused"), cwid="bad'value")

    def test_manual_exports_are_generated_from_api_scope_and_templates(self) -> None:
        root = Path(__file__).resolve().parents[2] / "query" / "AR" / "refunds"
        for name in ("transactions", "context"):
            template = (root / f"refund_{name}_extract.sql").read_text(encoding="utf-8")
            self.assertEqual(
                (root / f"refund_{name}_manual.sql").read_text(encoding="utf-8"),
                render_manual_extract_sql(template),
            )

    def test_row_count_guard_detects_api_truncation(self) -> None:
        frame = pd.DataFrame({"extract_row_count": [3, 3], "pidm": [1, 2]})
        with self.assertRaisesRegex(RuntimeError, "truncated"):
            _validate_complete_result(frame, "transactions", 0)

    def test_completed_batches_are_cached_and_resumable(self) -> None:
        class FakeClient:
            def __init__(self) -> None:
                self.calls = 0

            def run_sql(self, sql: str) -> pd.DataFrame:
                self.calls += 1
                if "t.tbraccd_amount AS amount" in sql:
                    return pd.DataFrame({"extract_row_count": [1], "pidm": [self.calls]})
                return pd.DataFrame({"extract_row_count": [1], "pidm": [self.calls]})

        with TemporaryDirectory() as directory:
            root = Path(directory)
            transaction_template = root / "transactions.sql"
            context_template = root / "context.sql"
            template = (
                "WITH batch_scope AS (__REFUND_SCOPE_SQL__) "
                "SELECT COUNT(*) OVER () AS extract_row_count, "
                "__BATCH_COUNT__ AS total, __BATCH_INDEX__ AS batch"
            )
            transaction_template.write_text(
                template + ", t.tbraccd_amount AS amount", encoding="utf-8"
            )
            context_template.write_text(template, encoding="utf-8")
            settings = ExtractSettings("202680", 2, root / "cache", run_date=date(2026, 9, 7))
            client = FakeClient()
            first = extract_refund_data(
                client,
                settings,
                transaction_template_path=transaction_template,
                context_template_path=context_template,
            )
            self.assertEqual(client.calls, 4)
            self.assertEqual(tuple(len(frame) for frame in first), (2, 2))

            resumed_client = FakeClient()
            resumed = extract_refund_data(
                resumed_client,
                replace(settings, resume=True),
                transaction_template_path=transaction_template,
                context_template_path=context_template,
            )
            self.assertEqual(resumed_client.calls, 0)
            self.assertEqual(tuple(len(frame) for frame in resumed), (2, 2))
            self.assertEqual(tuple(len(frame) for frame in read_refund_extracts(settings)), (2, 2))
            next_day = replace(settings, resume=True, run_date=date(2026, 9, 8))
            with self.assertRaisesRegex(ValueError, "manifest does not match"):
                read_refund_extracts(next_day)
            with self.assertRaisesRegex(ValueError, "manifest does not match"):
                extract_refund_data(
                    resumed_client, next_day,
                    transaction_template_path=transaction_template,
                    context_template_path=context_template,
                )
            self.assertEqual(resumed_client.calls, 0)

    def test_export_keeps_report_column_order(self) -> None:
        with TemporaryDirectory() as directory:
            output = Path(directory) / "refunds.xlsx"
            export_refund_report(pd.DataFrame(columns=REPORT_COLUMNS), output)
            workbook = load_workbook(output)
            self.assertEqual(workbook.sheetnames, list(REFUND_SHEETS))
            for sheet in workbook:
                headers = [cell.value for cell in sheet[1]]
                self.assertEqual(headers, WORKBOOK_COLUMNS)
                self.assertEqual([h for h in headers if not h.startswith("tab_")], REPORT_COLUMNS)
                self.assertFalse(sheet.tables)
                self.assertEqual(sheet.freeze_panes, "A2")
            workbook.close()

    def test_export_partitions_methods_and_preserves_review_accounts(self) -> None:
        def account(cwid, student=0, parent=0, delivery="NONE", **changes):
            row = dict.fromkeys(REPORT_COLUMNS)
            row.update(cwid=cwid, first_name="Synthetic", last_name="Example",
                       student_refund_amount=Decimal(str(student)), parent_refund_amount=Decimal(str(parent)),
                       total_refund_amount=Decimal(str(student)) + Decimal(str(parent)),
                       proposed_student_delivery=delivery,
                       proposed_parent_delivery="RFDP" if parent else "NONE",
                       third_party_review_required_ind="N", review_status="READY_FOR_STAFF_REVIEW",
                       last_ar_activity_date=datetime(2099, 8, 31, 12, 30), fdpl_row_count=1)
            row.update(changes)
            return row

        rows = [
            account("TEST-MIX", 150, 75,
                    "AFRD (Transact) 20.00; CRVC (Transact) 30.00; RFND (CHECK) 100.00"),
            account("TEST-SYS", 90, delivery="ARFD (System)",
                    review_status="WAIT_CHECK_CLEARING",
                    review_reasons="CHCK_CLEARING_WAIT_UNTIL_09/17/2099_AMOUNT_40.00"),
            account("TEST-CHCK", 25, delivery="RFND (CHECK)",
                    review_status="WAIT_CHECK_CLEARING",
                    review_reasons="CHCK_CLEARING_WAIT_UNTIL_09/18/2099_AMOUNT_25.00"),
            account("TEST-WAIT", 100, delivery="ACHK Clearing Wait until 09/16/2099 / 40.00; ARFD (System) 60.00"),
            account("TEST-HOLD", 45, 5, "Refund Hold - Student",
                    proposed_parent_delivery="Refund Hold - Parent", refund_hold_ind="Y",
                    review_status="HOLD",
                    review_reasons=("REFUND_HOLD_RH; Mines Park Charge - Review; "
                                    "CHCK_CLEARING_WAIT_UNTIL_09/19/2099_AMOUNT_20.00")),
            account("TEST-OLD", 15, delivery="AFRD (Transact) - May Be Too Old"),
            account("TEST-DATE", 10, delivery="ACHK Date Review"),
            account("TEST-THIRD", 10, 20, "THIRD_PARTY_REVIEW",
                    third_party_review_required_ind="Y", proposed_parent_delivery="THIRD_PARTY_REVIEW"),
            account("TEST-HOMP", 70, delivery="ARFD (System)", review_reasons="Mines Park Charge - Review"),
            account("TEST-UNKNOWN", 80, delivery="NEW DELIVERY CODE"),
            account("TEST-BLOCKED", total_refund_amount=Decimal(25), student_refund_amount=None, parent_refund_amount=None),
            account("TEST-MISMATCH", 50, delivery="AFRD (Transact) 20.00; RFND (CHECK) 20.00"),
        ]
        expected = {
            "Transact Refunds": {"TEST-MIX": 50},
            "Check Refunds": {"TEST-MIX": 100, "TEST-CHCK": 25},
            "Parent Refunds": {"TEST-MIX": 75},
            "System Refunds": {"TEST-SYS": 90, "TEST-WAIT": 60},
            "Refund Holds": {"TEST-HOLD": 50},
            "ACH Clearing": {"TEST-WAIT": 40},
            "ACH Reviews": {"TEST-OLD": 15, "TEST-DATE": 10},
            "Third Party Reviews": {"TEST-THIRD": 30},
            "Mines Park Reviews": {"TEST-HOMP": 70},
            "Manual Reviews": {"TEST-UNKNOWN": 80, "TEST-BLOCKED": 25, "TEST-MISMATCH": 50},
        }
        report = pd.DataFrame(rows)
        original = report.copy(deep=True)
        with TemporaryDirectory() as directory:
            output = Path(directory) / "refunds.xlsx"
            export_refund_report(report, output)
            sheets = pd.read_excel(output, sheet_name=None)
            for name, amounts in expected.items():
                self.assertEqual(dict(zip(sheets[name].cwid, sheets[name].tab_refund_amount)), amounts)
            self.assertEqual(sum(frame.tab_refund_amount.sum() for frame in sheets.values()),
                             sum(row["total_refund_amount"] for row in rows))
            self.assertEqual(sheets["Transact Refunds"].iloc[0].total_refund_amount, 225)
            self.assertEqual(sheets["System Refunds"].iloc[0].review_status,
                             "WAIT_CHECK_CLEARING")
            self.assertEqual(sheets["System Refunds"].iloc[0].tab_review_note,
                             "CHCK clearing wait: 40.00 becomes eligible on 09/17/2099.")
            self.assertEqual(sheets["Check Refunds"].loc[
                sheets["Check Refunds"].cwid == "TEST-CHCK", "tab_review_note"
            ].iloc[0], "CHCK clearing wait: 25.00 becomes eligible on 09/18/2099.")
            self.assertNotIn("TEST-HOLD", set(sheets["Mines Park Reviews"].cwid))
            self.assertEqual(
                sheets["Refund Holds"].iloc[0].tab_review_note,
                "RH account hold: do not issue any refund. "
                "CHCK clearing wait: 20.00 becomes eligible on 09/19/2099.",
            )
            self.assertTrue(sheets["Manual Reviews"].tab_review_note.notna().all())
            workbook = load_workbook(output)
            for sheet in workbook:
                self.assertFalse(sheet.tables)
                self.assertIn("$", sheet.cell(2, WORKBOOK_COLUMNS.index("tab_refund_amount") + 1).number_format)
                self.assertEqual(sheet.cell(2, WORKBOOK_COLUMNS.index("last_ar_activity_date") + 1).number_format,
                                 "mm/dd/yyyy h:mm AM/PM")
            workbook.close()
        pd.testing.assert_frame_equal(report, original)

    def test_manual_download_reader_accepts_xlsx_and_checks_term(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "transactions.xlsx"
            pd.DataFrame({
                "Extract Row Count": [1],
                "Extract Target Term": ["202680"],
                "PIDM": [1],
            }).to_excel(path, index=False)
            result = read_refund_download(
                path,
                label="Transaction",
                expected_target_term="202680",
            )
            self.assertEqual(result.columns.tolist(), ["pidm"])
            with self.assertRaisesRegex(ValueError, "expects 202710"):
                read_refund_download(
                    path,
                    label="Transaction",
                    expected_target_term="202710",
                )


if __name__ == "__main__":
    unittest.main()
