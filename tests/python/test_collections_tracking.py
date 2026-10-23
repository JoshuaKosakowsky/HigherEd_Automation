"""Exercise actual collections SQL and its existing application export path."""
import re
import sqlite3
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from unittest.mock import MagicMock, patch

import pandas as pd
from openpyxl import load_workbook

from app.gui.models import WorkflowContext, WorkflowMode
from app.gui.services.insights import _export_query, run_insights_connection
from shared.insights.config import InsightsConfigurationError
from shared.insights.query_catalog import get_query


class CollectionsTrackingTests(unittest.TestCase):
    def setUp(self):
        self.db = sqlite3.connect(":memory:")
        self.addCleanup(self.db.close)
        self.db.executescript("""
            ATTACH DATABASE ':memory:' AS taismgr;
            ATTACH DATABASE ':memory:' AS saturn;
            CREATE TABLE saturn.spriden (spriden_pidm INTEGER, spriden_id TEXT,
                spriden_first_name TEXT, spriden_last_name TEXT, spriden_change_ind TEXT);
            CREATE TABLE saturn.sprhold (sprhold_pidm INTEGER, sprhold_hldd_code TEXT,
                sprhold_from_date TEXT, sprhold_to_date TEXT);
            CREATE TABLE taismgr.tbraccd (tbraccd_pidm INTEGER,
                tbraccd_detail_code TEXT, tbraccd_amount NUMERIC,
                tbraccd_feed_date TEXT, tbraccd_term_code TEXT);
            CREATE TABLE taismgr.tbbdetc (tbbdetc_detail_code TEXT, tbbdetc_type_ind TEXT);
            INSERT INTO taismgr.tbbdetc VALUES ('CHARGE', 'C'), ('PAYMENT', 'P'), ('BAD', 'X');
        """)
        self.query = re.sub(
            r"CAST\((h\.\w+date) AS date\)", r"date(\1)",
            get_query("collections_tracking").render_sql(),
        ).replace("CURRENT_DATE", "'2099-10-02'")

    def account(self, pidm, start="2099-01-01", end="2099-12-31", code="CO"):
        self.db.execute("INSERT INTO saturn.spriden VALUES (?, ?, 'Synthetic', 'Student', NULL)",
                        (pidm, f"TEST-{pidm}"))
        self.db.execute("INSERT INTO saturn.sprhold VALUES (?, ?, ?, ?)",
                        (pidm, code, start, end))

    def transaction(self, pidm, amount, code="CHARGE", feed="2099-10-01", term="209910"):
        self.db.execute("INSERT INTO taismgr.tbraccd VALUES (?, ?, ?, ?, ?)",
                        (pidm, code, amount, feed, term))

    def test_hold_end_boundary_without_start_or_enrollment_restriction(self):
        self.account(1, end="2099-10-02 23:59:59")
        self.account(2, end="2099-10-01")
        self.account(3, end=None)
        self.account(4, code="PP")
        self.account(5, start="2099-11-01", code=" co ")
        self.account(6, start=None)
        rows = self.db.execute(self.query).fetchall()
        self.assertEqual([row[0] for row in rows], ["TEST-5", "TEST-6"])
        self.assertEqual([row[5] for row in rows], ["CO", "CO"])
        self.assertTrue(all(row[-1] == 2 for row in rows))

    def test_latest_qualifying_hold_nulls_ties_and_duplicates(self):
        self.account(1, start=None)
        self.transaction(1, 100)
        self.db.executescript("""
            INSERT INTO saturn.sprhold VALUES
                (1, 'CO', '2099-03-01', '2099-11-01'),
                (1, 'CO', '2099-03-01', '2099-12-01'),
                (1, 'CO', '2099-03-01', '2099-12-01'),
                (1, 'CO', '2099-09-01', '2099-10-01'),
                (1, 'PP', '2099-10-01', '2099-12-31');
            INSERT INTO saturn.spriden VALUES (1, 'OLD', 'Old', 'Name', 'N');
        """)
        self.assertEqual(self.db.execute(self.query).fetchall(), [
            ('TEST-1', 'Synthetic', 'Student', 100, '2099-10-01',
             'CO', '2099-03-01', '2099-12-01', 1),
        ])

    def test_all_term_net_balance_reversals_zero_credit_and_latest_feed(self):
        for pidm in range(1, 4):
            self.account(pidm)
            self.transaction(pidm, 100)
        self.transaction(1, -25)
        self.transaction(1, 20, "PAYMENT", "2099-10-02 14:30:00", "209980")
        self.transaction(1, -5, "PAYMENT", None)
        self.transaction(2, 100, "PAYMENT")
        self.transaction(3, 150, "PAYMENT")
        rows = self.db.execute(self.query).fetchall()
        self.assertEqual([row[3] for row in rows], [60, 0, -50])
        self.assertEqual(rows[0][4], "2099-10-02 14:30:00")

    def test_missing_or_unknown_transactions_keep_account_with_blank_balance(self):
        for pidm in range(1, 5):
            self.account(pidm)
        self.transaction(2, None)
        self.transaction(3, 50, "MISSING")
        self.transaction(4, 100)
        self.transaction(4, 50, "BAD")
        rows = self.db.execute(self.query).fetchall()
        self.assertEqual(len(rows), 4)
        self.assertTrue(all(row[3] is None for row in rows))
        self.assertIsNone(rows[0][4])

    def test_empty_results_keep_expected_headers(self):
        cursor = self.db.execute(self.query)
        self.assertEqual(cursor.fetchall(), [])
        self.assertEqual([col[0] for col in cursor.description], [
            "CWID", "First Name", "Last Name", "Current Account Balance",
            "Latest Feed Date", "Hold Code", "From Date", "To Date", "Report Row Count",
        ])

    def test_app_export_dates_completeness_and_administrator_access(self):
        self.account(1, end="9999-12-31")
        self.transaction(1, -50, feed="2099-10-02 14:30:00")
        frame = pd.read_sql_query(self.query, self.db)
        query = get_query("collections_tracking")
        self.assertIsNone(query.term_variable)
        for administrator, truncated, empty in (
            (True, False, False), (True, True, False),
            (False, False, False), (True, False, True),
        ):
            with self.subTest(administrator=administrator, truncated=truncated, empty=empty):
                client = MagicMock()
                client.__enter__.return_value = client
                result_frame = frame.iloc[:0].copy() if empty else frame.copy()
                if truncated:
                    result_frame['Report Row Count'] = 2
                client.run_sql_file.return_value = result_frame
                policy = MagicMock()
                policy.is_administrator.return_value = administrator
                with tempfile.TemporaryDirectory() as directory, patch(
                    'app.gui.services.insights.load_department_profiles',
                    return_value={'TEST': MagicMock()},
                ), patch('app.gui.services.insights.load_access_configuration', return_value=policy), patch(
                    'app.gui.services.insights.get_shared_gui_access_path',
                    return_value=Path(directory) / 'access.json',
                ), patch('app.gui.services.insights.build_authenticated_client',
                         return_value=(client, 'cached')) as build:
                    output = Path(directory) / 'collections.xlsx'
                    result = run_insights_connection(WorkflowContext(
                        'insights_query_export', {'action': 'query_export',
                         'query_id': query.query_id, 'output_path': str(output)}, WorkflowMode.TEST,
                    ))
                    success = administrator and not truncated
                    self.assertEqual(result.success, success, result.message)
                    self.assertEqual(output.exists(), success)
                    if not administrator:
                        build.assert_not_called()
                    if success:
                        client.run_sql_file.assert_called_once_with(query.sql_path)
                        workbook = load_workbook(output)
                        try:
                            sheet = workbook.active
                            self.assertEqual(sheet.max_column, 8)
                            self.assertEqual(sheet.max_row, 1 if empty else 2)
                            if not empty:
                                self.assertEqual(sheet['D2'].value, -50)
                                self.assertEqual(sheet['E2'].value, datetime(2099, 10, 2, 14, 30))
                                self.assertEqual(sheet['G2'].value, datetime(2099, 1, 1))
                                self.assertTrue(sheet['H2'].is_date)
                                self.assertEqual(sheet['H2'].value, datetime(9999, 12, 31))
                        finally:
                            workbook.close()

    def test_calendar_dates_preserve_far_future_dates_and_blanks_in_excel(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / 'hold-dates.xlsx'
            _export_query(pd.DataFrame({
                'From Date': ['1900-01-01', None, date(2099, 1, 1), pd.NaT],
                'To Date': ['9999-12-31', '9999-12-31T00:00:00Z',
                            datetime(9999, 12, 31), None],
            }), output)
            workbook = load_workbook(output)
            try:
                sheet = workbook.active
                for row in (2, 3, 4):
                    self.assertEqual(sheet.cell(row, 2).value, datetime(9999, 12, 31))
                    self.assertTrue(sheet.cell(row, 2).is_date)
                self.assertEqual(sheet['A2'].value, datetime(1900, 1, 1))
                self.assertIsNone(sheet['A3'].value)
                self.assertIsNone(sheet['A5'].value)
                self.assertIsNone(sheet['B5'].value)
            finally:
                workbook.close()

    def test_invalid_hold_dates_still_reject_export(self):
        for value in ('not-a-date', '9999-02-30', 123):
            with self.subTest(value=value), tempfile.TemporaryDirectory() as directory:
                output = Path(directory) / 'invalid.xlsx'
                with self.assertRaisesRegex(InsightsConfigurationError, 'unreadable To Date'):
                    _export_query(pd.DataFrame({'To Date': [value]}), output)
                self.assertFalse(output.exists())


if __name__ == '__main__':
    unittest.main()
