"""Execute the shared report with synthetic accounts; never use live systems."""
from __future__ import annotations

import re
import sqlite3
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import MagicMock, patch

import pandas as pd
from openpyxl import load_workbook

from app.gui.models import WorkflowContext, WorkflowMode
from app.gui.services.insights import run_insights_connection
from data_processing.refunds.terms import previous_term
from shared.banner.term import get_banner_term
from shared.insights.query_catalog import get_query


QUERY = get_query("term_balances_no_au")
COLUMNS = ["CWID", "Last Name", "First Name", "Term", "Amount", "Amount w/ LF",
           "Account Balance", "Cross-Term Credit", "PP Hold", "Review Notes"]


class TermBalancesNoAUQueryTests(unittest.TestCase):
    def setUp(self) -> None:
        self.db = sqlite3.connect(":memory:")
        self.db.row_factory = sqlite3.Row
        self.addCleanup(self.db.close)
        self.db.create_function("CONCAT", -1, lambda *parts: "".join(str(part) for part in parts))
        self.db.create_function("MAKE_DATE", 3, lambda y, m, d: date(y, m, d).isoformat())
        self.db.executescript("""
            ATTACH DATABASE ':memory:' AS taismgr;
            ATTACH DATABASE ':memory:' AS saturn;
            CREATE TABLE taismgr.tbraccd (tbraccd_pidm INTEGER,
                tbraccd_term_code TEXT, tbraccd_detail_code TEXT, tbraccd_amount NUMERIC);
            CREATE TABLE taismgr.tbbdetc (tbbdetc_detail_code TEXT, tbbdetc_type_ind TEXT);
            INSERT INTO taismgr.tbbdetc VALUES
                ('CHARGE', 'C'), ('PAYMENT', 'P'), ('BAD', 'X'), ('NORMALIZED', ' p ');
            CREATE TABLE saturn.sgbstdn (sgbstdn_pidm INTEGER);
            CREATE TABLE saturn.spriden (spriden_pidm INTEGER, spriden_id TEXT,
                spriden_last_name TEXT, spriden_first_name TEXT, spriden_change_ind TEXT);
            CREATE TABLE saturn.sprhold (sprhold_pidm INTEGER, sprhold_hldd_code TEXT,
                sprhold_from_date TEXT, sprhold_to_date TEXT);
        """)

    def account(self, pidm: int, *, student: bool = True, identity: bool = True) -> None:
        if student:
            self.db.execute("INSERT INTO saturn.sgbstdn VALUES (?)", (pidm,))
        if identity:
            self.db.execute("INSERT INTO saturn.spriden VALUES (?, ?, 'Student', 'Synthetic', NULL)",
                            (pidm, f"TEST-{pidm}"))

    def transaction(self, pidm: int, amount: float | None, term: str | None = "209980",
                    code: str = "CHARGE") -> None:
        self.db.execute("INSERT INTO taismgr.tbraccd VALUES (?, ?, ?, ?)", (pidm, term, code, amount))

    def hold(self, pidm: int, code: str, end: str | None,
             start: str | None = "2099-01-01") -> None:
        self.db.execute("INSERT INTO saturn.sprhold VALUES (?, ?, ?, ?)", (pidm, code, start, end))

    def results(self, run_date: str = "2099-10-02") -> list[dict]:
        sql = QUERY.render_sql()
        # Run the actual query, adapting only PostgreSQL date syntax and clock.
        sql = sql.replace("EXTRACT(YEAR FROM report_date)", "CAST(strftime('%Y', report_date) AS integer)")
        sql = re.sub(r"CAST\((h\.sprhold_(?:from|to)_date) AS date\)", r"date(\1)", sql)
        sql = sql.replace("CURRENT_DATE", f"'{run_date}'")
        cursor = self.db.execute(sql)
        self.assertEqual([column[0] for column in cursor.description], COLUMNS + ["Report Row Count"])
        rows = [dict(row) for row in cursor.fetchall()]
        self.assertTrue(all(row["Report Row Count"] == len(rows) for row in rows))
        return rows

    def test_separate_term_balances_and_late_fee_projection(self) -> None:
        self.account(1)
        self.transaction(1, 30000)
        self.transaction(1, 15000, code="PAYMENT")
        self.transaction(1, 1000, "209955")
        self.transaction(1, 500, "209910")
        rows = self.results()
        self.assertEqual([(r["Term"], r["Amount"], r["Amount w/ LF"]) for r in rows],
                         [("209980", 15000, 15225), ("209955", 1000, 1015), ("209910", 500, 507.5)])
        self.assertTrue(all(r["Account Balance"] == 16500 and r["Cross-Term Credit"] == "No"
                            and r["Review Notes"] is None for r in rows))

    def test_cross_term_credit_keeps_zero_and_credit_accounts_for_review(self) -> None:
        for pidm, payment in ((1, 15000), (2, 16000)):
            self.account(pidm)
            self.transaction(pidm, 30000)
            self.transaction(pidm, 15000, code="PAYMENT")
            self.transaction(pidm, payment, "209955", "PAYMENT")
        rows = self.results()
        self.assertEqual([r["Account Balance"] for r in rows], [0, -1000])
        self.assertTrue(all(r["Amount"] == 15000 and r["Amount w/ LF"] is None
                            and r["Cross-Term Credit"] == "Yes" for r in rows))

    def test_partial_credit_blanks_every_positive_term_fee_without_allocating(self) -> None:
        self.account(1)
        self.transaction(1, 1000, "209955")
        self.transaction(1, 500)
        self.transaction(1, 100, "209910", "PAYMENT")
        rows = self.results()
        self.assertEqual([r["Amount"] for r in rows], [500, 1000])
        self.assertTrue(all(r["Amount w/ LF"] is None and r["Account Balance"] == 1400 for r in rows))

    def test_older_and_future_credits_are_account_context(self) -> None:
        for pidm, term in ((1, "209880"), (2, "210010")):
            self.account(pidm)
            self.transaction(pidm, 1000)
            self.transaction(pidm, 300, term, "PAYMENT")
        rows = self.results()
        self.assertEqual(len(rows), 2)
        self.assertTrue(all(r["Account Balance"] == 700 and r["Cross-Term Credit"] == "Yes"
                            and r["Amount w/ LF"] is None for r in rows))

    def test_positive_history_does_not_expand_output_or_create_credit(self) -> None:
        self.account(1)
        for term in ("209980", "209955", "209910", "209880", "210010"):
            self.transaction(1, 100, term)
        rows = self.results()
        self.assertEqual(len(rows), 3)
        self.assertTrue(all(r["Account Balance"] == 500 and r["Amount w/ LF"] == 101.5 for r in rows))

    def test_reversals_normalized_types_and_same_term_payments(self) -> None:
        self.account(1)
        for amount, code in ((1000, "CHARGE"), (-100, "CHARGE"),
                             (500, "PAYMENT"), (-50, "PAYMENT"), (50, "NORMALIZED")):
            self.transaction(1, amount, code=code)
        row = self.results()[0]
        self.assertEqual((row["Amount"], row["Amount w/ LF"], row["Cross-Term Credit"]), (400, 406, "No"))

    def test_zero_negative_terms_and_accounts_without_selected_activity(self) -> None:
        for pidm in range(1, 4):
            self.account(pidm)
        self.transaction(1, 100)
        self.transaction(1, 100, code="PAYMENT")
        self.transaction(2, 100, code="PAYMENT")
        self.transaction(3, 100, "209880")
        self.assertEqual(self.results(), [])

    def test_au_uses_only_end_date_and_ignores_start(self) -> None:
        for pidm in range(1, 6):
            self.account(pidm)
            self.transaction(pidm, 100)
        self.hold(1, " au ", "2099-10-03", "2099-12-01")
        self.hold(2, "AU", "2099-10-02 23:59:59")
        self.hold(3, "AU", "2099-10-01")
        self.hold(4, "AU", None)
        self.hold(5, "AU", "9999-12-31", None)
        self.assertEqual([r["CWID"] for r in self.results()], ["TEST-2", "TEST-3", "TEST-4"])

    def test_pp_accounts_retained_without_hold_fanout(self) -> None:
        for pidm in range(1, 4):
            self.account(pidm)
            self.transaction(pidm, 100)
        for _ in range(2):
            self.hold(1, " pp ", "2099-10-03", None)
        self.hold(2, "PP", "2099-10-02")
        self.hold(3, "PP", "2099-12-01", "2099-11-01")
        rows = self.results()
        self.assertEqual([r["PP Hold"] for r in rows], ["PP", None, None])

    def test_student_history_not_current_enrollment_defines_student(self) -> None:
        self.account(1)
        self.account(2, student=False)
        self.transaction(1, 100)
        self.transaction(2, 100)
        self.db.execute("INSERT INTO saturn.sgbstdn VALUES (1)")
        self.db.execute("INSERT INTO saturn.spriden VALUES (1, 'OLD', 'Old', 'Name', 'N')")
        self.hold(1, "CO", "9999-12-31")
        self.assertEqual([r["CWID"] for r in self.results()], ["TEST-1"])

    def test_unknown_selected_amounts_and_types_remain_review_rows(self) -> None:
        for pidm, amount, code in ((1, None, "CHARGE"), (2, 100, "BAD"), (3, 100, "UNMAPPED")):
            self.account(pidm)
            self.transaction(pidm, amount, code=code)
        rows = self.results()
        self.assertEqual(len(rows), 3)
        self.assertTrue(all(r["Amount"] is None and r["Amount w/ LF"] is None
                            and r["Account Balance"] is None and r["Cross-Term Credit"] == "Unknown"
                            and "review balances" in r["Review Notes"] for r in rows))

    def test_unknown_history_and_missing_terms_prevent_fee_projection(self) -> None:
        self.account(1)
        self.account(2)
        self.transaction(1, 100)
        self.transaction(1, None, "209880")
        self.transaction(2, 100)
        self.transaction(2, 25, None)
        rows = self.results()
        self.assertEqual([r["Account Balance"] for r in rows], [None, 125])
        self.assertTrue(all(r["Amount w/ LF"] is None and r["Cross-Term Credit"] == "Unknown" for r in rows))

    def test_missing_and_duplicate_identity_do_not_hide_or_duplicate_balances(self) -> None:
        self.account(1, identity=False)
        self.account(2)
        self.db.execute("INSERT INTO saturn.spriden VALUES (2, 'ALT', 'Student', 'Synthetic', NULL)")
        self.transaction(1, 100)
        self.transaction(2, 200)
        rows = self.results()
        self.assertEqual(len(rows), 2)
        self.assertEqual(sum(r["Amount"] for r in rows), 300)
        self.assertTrue(all(r["Amount w/ LF"] is None and "identity" in r["Review Notes"] for r in rows))

    def test_duplicate_metadata_cannot_double_money(self) -> None:
        self.account(1)
        self.transaction(1, 100)
        self.db.execute("INSERT INTO taismgr.tbbdetc VALUES ('CHARGE', 'C')")
        row = self.results()[0]
        self.assertIsNone(row["Amount"])
        self.assertIsNone(row["Amount w/ LF"])
        self.assertIn("duplicate detail type", row["Review Notes"])

    def test_term_boundaries_match_shared_institution_configuration(self) -> None:
        self.account(1)
        for year in (2098, 2099):
            for suffix in ("10", "55", "80"):
                self.transaction(1, 100, f"{year}{suffix}")
        for run_date in ("2099-01-01", "2099-05-15", "2099-05-16", "2099-07-15",
                         "2099-07-16", "2099-12-31"):
            with self.subTest(run_date=run_date):
                current = get_banner_term(date.fromisoformat(run_date)).code
                prior = previous_term(current)
                self.assertEqual([r["Term"] for r in self.results(run_date)],
                                 [current, prior, previous_term(prior)])

    def test_account_rounding_uses_all_raw_activity_and_fee_rounds_to_cents(self) -> None:
        self.account(1)
        # Avoid a binary floating-point half-cent tie in SQLite; PostgreSQL
        # executes this report with Banner's numeric monetary columns.
        self.transaction(1, 1.01)
        self.transaction(1, 0.004, "209955")
        self.transaction(1, 0.004, "209910")
        row = self.results()[0]
        self.assertEqual((row["Amount"], row["Amount w/ LF"], row["Account Balance"]), (1.01, 1.03, 1.02))


class TermBalancesNoAUExportTests(unittest.TestCase):
    def test_existing_export_preserves_blank_fees_and_checks_access_and_count(self) -> None:
        data = [["TEST-1", "Student", "Synthetic", "209980", 15000, None, 0, "Yes", "PP",
                 "Cross-term credit; review payment application.", 1]]
        cases = ((True, "complete", True), (True, "empty", True),
                 (True, "capped", False), (True, "missing-count", False), (False, "complete", False))
        for administrator, kind, success in cases:
            with self.subTest(administrator=administrator, kind=kind), tempfile.TemporaryDirectory() as directory:
                frame = pd.DataFrame([] if kind == "empty" else data, columns=COLUMNS + ["Report Row Count"])
                if kind == "capped":
                    frame["Report Row Count"] = 2
                elif kind == "missing-count":
                    frame = frame.drop(columns="Report Row Count")
                client = MagicMock()
                client.__enter__.return_value = client
                client.run_sql_file.return_value = frame
                output = Path(directory) / "review.xlsx"
                with patch("app.gui.services.insights.load_access_configuration") as policy, patch(
                    "app.gui.services.insights.build_authenticated_client", return_value=(client, "synthetic")
                ) as build:
                    policy.return_value.is_administrator.return_value = administrator
                    result = run_insights_connection(WorkflowContext("insights_query_export", {
                        "action": "query_export", "query_id": QUERY.query_id, "output_path": str(output),
                    }, WorkflowMode.TEST))
                self.assertEqual(result.success, success, result.message)
                self.assertEqual(output.exists(), success)
                if not administrator:
                    build.assert_not_called()
                if success:
                    client.run_sql_file.assert_called_once_with(QUERY.sql_path)
                    workbook = load_workbook(output)
                    try:
                        sheet = workbook.active
                        self.assertEqual([cell.value for cell in sheet[1]], COLUMNS)
                        self.assertEqual(sheet.max_row, 1 if kind == "empty" else 2)
                        if kind != "empty":
                            self.assertIsNone(sheet["F2"].value)
                            self.assertEqual(sheet["E2"].value, 15000)
                            self.assertEqual(sheet["G2"].value, 0)
                            self.assertEqual(sheet["I2"].value, "PP")
                    finally:
                        workbook.close()


if __name__ == "__main__":
    unittest.main()
