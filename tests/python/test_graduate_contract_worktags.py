"""Synthetic student worktag parsing and workbook enrichment checks."""

from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from data_processing.graduate_contract_recon.export import export_reconciliation
from data_processing.graduate_contract_recon.sources import SourceData, read_source
from data_processing.graduate_contract_recon.worktags import prepare_worktags
import test_graduate_contract_recon as recon_fixtures
from test_graduate_contract_recon import wd_row, bn_row


def worktag_source(rows):
    return SourceData("Worktags", "synthetic.xlsx", ["Line Memo", "Worktags"],
                      list(enumerate(rows, 2)), {"worktag_memo": 0, "worktags": 1})


class StudentWorktagTests(unittest.TestCase):
    def test_extracts_all_types_terms_and_deduplicates_without_names(self):
        result = prepare_worktags(worktag_source([
            ["10000001 SU26 Synthetic Student", "Cost Center: 999999 Other\nGrant: 012345 Synthetic\nProject: 654321 Other"],
            ["GCSA 10000001 SU26 Synthetic Student", "Grant: 012345 Synthetic"],
            ["10000001 FA26 Synthetic Student", "Operating Account: 123456 Other\nFund: 11"],
        ]))
        self.assertEqual(result.rows, [
            ["10000001", "SU26", "Summer 2026", "Grant", "012345"],
            ["10000001", "SU26", "Summer 2026", "Project", "654321"],
            ["10000001", "FA26", "Fall 2026", "Operating Account", "123456"]])
        self.assertEqual(result.by_cwid["10000001"],
            "Grant: 012345; Operating Account: 123456; Project: 654321")
        self.assertEqual(result.warning_count, 0)

    def test_missing_ambiguous_ids_and_wrong_tag_lengths_are_not_guessed(self):
        result = prepare_worktags(worktag_source([
            ["SU26 Synthetic", "Grant: 123456"],
            ["10000001 10000002 SP26 Synthetic", "Project 654321"],
            ["10000003 FA26 Synthetic", "Grant: 1234567\nOperating Account: 12345"],
        ]))
        self.assertEqual(result.by_cwid, {})
        self.assertEqual(result.warning_count, 3)
        self.assertEqual(result.rows[1][2:4], ["Spring 2026", "Project"])
        with self.assertRaisesRegex(ValueError, "no data rows"):
            prepare_worktags(worktag_source([]))

    def test_headers_detected_after_title_with_reordered_columns(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "source.xlsx"
            book = Workbook()
            book.active.append(["Find Journal Lines"])
            book.active.append(["Worktags", "Unused", "Line Memo"])
            book.active.append(["Grant: 123456", "Discard me", "10000001 SU26 Synthetic"])
            book.save(path); book.close()
            result = prepare_worktags(read_source(path, "Worktags"))
            self.assertEqual(result.rows, [["10000001", "SU26", "Summer 2026", "Grant", "123456"]])

    def test_export_enriches_cwid_without_duplicating_financial_rows(self):
        builder = recon_fixtures.GraduateContractReconTests(); builder.setUp()
        review = builder.build([wd_row()], [bn_row(amount=30), bn_row(amount=30, account="10000002")])
        tags = prepare_worktags(worktag_source([
            ["10000001 SU26 Synthetic", "Grant: 012345\nProject: 654321"],
            ["10000001 FA26 Synthetic", "Grant: 012345"],
        ]))
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "review.xlsx"
            export_reconciliation(review, path, app_version="synthetic", banner_source="manual", worktags=tags)
            book = load_workbook(path)
            try:
                self.assertEqual(len(book.sheetnames), 7)
                self.assertEqual(book["Worktags"].max_column, 5)
                self.assertEqual(book["1305 CWID Recon"].max_row, 3)
                self.assertEqual(book["1305 CWID Recon"]["H1"].value, "Student Worktag")
                self.assertEqual(book["1305 CWID Recon"]["H2"].value, "Grant: 012345; Project: 654321")
                self.assertEqual(book["Worktags"]["E2"].value, "012345")
                self.assertEqual(book["Worktags"]["E2"].number_format, "@")
                self.assertEqual(book["1305 CWID Recon"]["H3"].value, "No matching worktag")
                self.assertEqual(book["1305 CWID Recon"]["D2"].value, 30)
                self.assertTrue(all(not sheet.tables for sheet in book))
            finally:
                book.close()


if __name__ == "__main__":
    unittest.main()
